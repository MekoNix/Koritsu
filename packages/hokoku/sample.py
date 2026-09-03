"""
sample — чужой готовый отчёт как источник оформления.

Пользователь приносит образец: работу с прошлого курса, методичку, кафедральный пример.
Оттуда берётся вид — поля страницы, основной шрифт, заголовки, нумерация разделов, вид
подписей к рисункам и таблицам, — и накладывается на наш документ (`apply_style`).

**Пример — не шаблон.** Шаблон это документ *с тегами* `{{…}}`, его заполняют; пример —
законченный текст *без тегов*, из него берут оформление. Вход у них разный намеренно:
подсунув одно вместо другого в общую дверь, человек получил бы пустой отчёт («шаблон без
тегов») или потерянное оформление и не понял бы, почему. Поэтому `style_from_sample`
отказывается работать с документом, в котором есть теги, а имя функции называет вход.

**Пример — недоверенный DOCX.** Его прислал пользователь, ровно как шаблон, и проходит
он ту же проверку: `walker.open_document` → `safety.validate_docx` (макросы VBA,
zip-бомба, zip-slip, XXE) и `strip_external_refs`. Отдельного пути чтения здесь нет и
быть не должно: вторая дверь в тот же дом — вторая дырка. Из примера в наш документ
не переносится ни один байт содержимого, только числа и имена шрифтов, разобранные
этим модулем, — поэтому картинка-бомба или внешняя ссылка из примера не может доехать
до собранного отчёта даже в принципе.

**Что берётся** (`StyleProfile` — плоские поля, их видно и можно показать человеку):
размер листа и поля; шрифт, кегль, интервал и абзацный отступ основного текста;
на каждый уровень заголовка — шрифт, кегль, жирность, выключка и нумерован ли он;
слово и тире подписей («Рисунок 1 — …» против «Рис. 1. …»).

**Что взять нельзя** — и это не забывчивость, а решение; всё замеченное такое лежит
в `StyleProfile.notes` строками для человека:
  • содержимое — текст, картинки, титульный лист, колонтитулы: это чужая работа;
  • тема оформления (`theme1.xml`) целиком: шрифты из неё *разрешаются* в конкретные
    имена, но связь со схемой темы теряется, и переключение темы в Word наш документ
    больше не перекрасит;
  • произвольные пользовательские стили примера: их некому применять — рендер знает
    только `Heading N`, `Code`, `Caption`, `Quote` (`template.STYLE_SPECS`);
  • нумерация подписей по главам («Рисунок 1.2»): наши поля SEQ считают сквозным
    номером, а глава в них не участвует;
  • подпись над рисунком: `render` ставит подпись рисунка под ним, таблицы — над.

**Строение** (какие разделы и в каком порядке) — `outline_from_sample`, отдельная
функция и отдельный список. Это не оформление, и выдавать одно за другое нельзя:
`kadai` сочиняет структуру вслепую и списком заголовков воспользуется, а перепутав
его с оформлением, получил бы чужие разделы в готовой работе.
"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field

from docx.oxml.ns import qn

from . import docx_ops as ops
from .images import EMU_PER_CM
from .model import HokokuError
from .template import (BodyText, PageSetup, _ALIGN, _apply_font, apply_body_text,
                       apply_page_setup, ensure_styles)

MAX_LEVEL = 6                     # глубже заголовки не нумеруются и в оглавление не идут

_ALIGN_NAME = {0: "left", 1: "center", 2: "right", 3: "both"}   # WD_ALIGN_PARAGRAPH

# Слово подписи → что это. Признаём и сокращения, и латиницу: образец бывает переводной.
_CAPTION_WORDS = {"рисунок": "figure", "рис": "figure", "figure": "figure", "fig": "figure",
                  "таблица": "table", "табл": "table", "table": "table"}
_CAPTION_RE = re.compile(
    r"^\s*(?P<word>[A-Za-zА-Яа-яЁё]+)\.?\s*(?P<num>\d+(?:[.\-]\d+)*)\s*"
    r"(?P<sep>[—–-]|\.|:)?\s*(?P<rest>.*)$")
# «1.2 Название» — номер раздела, набранный руками (нумерация Word выглядит иначе:
# в тексте абзаца её нет вовсе)
_MANUAL_NUM_RE = re.compile(r"^\s*(?P<num>\d+(?:\.\d+)*)\.?\s+(?P<rest>\S.*)$")


class SampleError(HokokuError):
    """Из этого документа оформление взять нельзя (текст — пользовательский)."""


# ── профиль оформления ────────────────────────────────────────────────────────

@dataclass(frozen=True)
class HeadingLook:
    """Как выглядит заголовок одного уровня. None — «в примере не сказано, наследуется»."""
    level: int
    font: str | None = None
    size_pt: float | None = None
    bold: bool | None = None
    align: str | None = None
    numbered: bool = False


@dataclass(frozen=True)
class StyleProfile:
    """Оформление, снятое с примера. Плоские поля, а не куски чужого XML: то, что нельзя
    назвать словом, нельзя ни показать человеку, ни объяснить, почему оно не применилось."""
    page: PageSetup
    body: BodyText
    headings: tuple[HeadingLook, ...] = ()
    captions: dict = field(default_factory=dict)   # figure/table + *_align, как в styles.yaml
    numbered_headings: bool = False
    notes: tuple[str, ...] = ()

    def style_overrides(self) -> dict:
        """Перегрузки для `render(style=…)`: вид подписей живёт там, а не в стиле Caption.

        Стиль задаёт кегль и отступы подписи, а слова «Рисунок 1 —» собирает рендер по
        `styles.yaml`. Поэтому оформление примера доезжает до отчёта двумя путями, и
        забыть второй нельзя: подписи выйдут нашими, а не кафедральными.
        """
        out = {"captions": dict(self.captions)} if self.captions else {}
        sizes = {h.level: h.size_pt for h in self.headings if h.size_pt}
        if sizes:
            out["headings"] = sizes
        return out


# ── чтение примера ────────────────────────────────────────────────────────────

def _open_sample(sample, allow_tags: bool):
    """Пример — недоверенный DOCX: та же дверь, что у шаблона, со всеми её проверками."""
    from .walker import marked_paragraphs, open_document
    doc = open_document(sample)
    if not allow_tags and marked_paragraphs(doc, ("{{",)):
        raise SampleError(
            "в документе есть теги {{…}} — это шаблон, а не пример готового отчёта. "
            "Шаблон заполняют (render), из примера берут оформление (style_from_sample); "
            "перепутав их, вы получите либо пустой отчёт, либо своё оформление вместо чужого")
    return doc


def _theme_fonts(doc) -> dict:
    """Шрифты темы: в настоящем документе Word `Heading 1` ссылается не на гарнитуру,
    а на `majorHAnsi`. Без разрешения темы кегль мы бы сняли, а шрифт вернули пустым —
    ровно на тех кафедральных документах, ради которых всё это и написано."""
    out: dict = {}
    try:
        parts = [p for p in doc.part.package.iter_parts()
                 if str(p.partname).endswith(".xml") and "/theme/" in str(p.partname)]
    except Exception:                                          # noqa: BLE001
        return out
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for part in parts:
        try:
            from lxml import etree
            root = etree.fromstring(part.blob)
        except Exception:                                      # noqa: BLE001
            continue
        for kind, tag in (("major", "majorFont"), ("minor", "minorFont")):
            el = root.find(f".//{A}fontScheme/{A}{tag}/{A}latin")
            if el is not None and el.get("typeface"):
                out.setdefault(kind, el.get("typeface"))
    return out


def _chain(style):
    """Стиль и его базовые — по цепочке `basedOn`; циклы в чужом файле не редкость."""
    seen = set()
    while style is not None and id(style.element) not in seen:
        seen.add(id(style.element))
        yield style
        try:
            style = style.base_style
        except (KeyError, AttributeError):
            return


def _doc_defaults(doc):
    """<w:docDefaults> — то, от чего наследуется даже Normal."""
    try:
        return doc.styles.element.find(qn("w:docDefaults"))
    except AttributeError:
        return None


def _rfonts_name(rpr, themes: dict) -> str | None:
    if rpr is None:
        return None
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        return None
    name = rf.get(qn("w:ascii")) or rf.get(qn("w:hAnsi"))
    if name:
        return name
    theme = rf.get(qn("w:asciiTheme")) or rf.get(qn("w:hAnsiTheme")) or ""
    return themes.get("major" if theme.startswith("major") else "minor") if theme else None


def _style_font(doc, style, themes: dict) -> str | None:
    for st in _chain(style):
        name = _rfonts_name(st.element.find(qn("w:rPr")), themes)
        if name:
            return name
    dd = _doc_defaults(doc)
    if dd is not None:
        rpr = dd.find(qn("w:rPrDefault"))
        return _rfonts_name(rpr.find(qn("w:rPr")) if rpr is not None else None, themes)
    return None


def _style_prop(style, getter):
    for st in _chain(style):
        try:
            value = getter(st)
        except (KeyError, AttributeError, ValueError):
            value = None
        if value is not None:
            return value
    return None


def _para_font(doc, p, themes: dict) -> str | None:
    for r in p.runs:
        name = _rfonts_name(r._r.find(qn("w:rPr")), themes)
        if name:
            return name
    return _style_font(doc, p.style, themes)


def _para_size(p) -> float | None:
    for r in p.runs:
        if r.font.size is not None:
            return round(r.font.size.pt, 1)
    size = _style_prop(p.style, lambda s: s.font.size)
    return round(size.pt, 1) if size is not None else None


def _para_bold(p) -> bool | None:
    vals = [r.font.bold for r in p.runs if r.text.strip()]
    if vals and all(v is not None for v in vals):
        return all(vals)
    return _style_prop(p.style, lambda s: s.font.bold)


def _para_align(p) -> str | None:
    align = p.paragraph_format.alignment
    if align is None:
        align = _style_prop(p.style, lambda s: s.paragraph_format.alignment)
    return _ALIGN_NAME.get(int(align)) if align is not None else None


def _outline_level(doc, p) -> int | None:
    """Уровень заголовка: сначала сам абзац, потом цепочка стилей, потом имя стиля.
    По имени «Heading N» одному судить нельзя — в русском Word стиль называется
    «Заголовок 1», а у переделанного кафедрального образца как угодно."""
    ppr = p._p.find(qn("w:pPr"))
    if ppr is not None:
        el = ppr.find(qn("w:outlineLvl"))
        if el is not None:
            lvl = int(el.get(qn("w:val"))) + 1
            return lvl if 1 <= lvl <= MAX_LEVEL else None
    for st in _chain(p.style):
        spr = st.element.find(qn("w:pPr"))
        el = spr.find(qn("w:outlineLvl")) if spr is not None else None
        if el is not None:
            lvl = int(el.get(qn("w:val"))) + 1
            return lvl if 1 <= lvl <= MAX_LEVEL else None
    m = re.fullmatch(r"Heading (\d)", (p.style.name or "") if p.style is not None else "")
    return int(m.group(1)) if m and int(m.group(1)) <= MAX_LEVEL else None


def _numbered(p) -> bool:
    """Нумерован ли заголовок — полем Word или руками. Второе для нас так же важно:
    сняв только первое, мы отдали бы `kadai` ненумерованные разделы там, где в образце
    они пронумерованы, — и отчёт не приняли бы."""
    ppr = p._p.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:numPr")) is not None:
        return True
    for st in _chain(p.style):
        spr = st.element.find(qn("w:pPr"))
        if spr is not None and spr.find(qn("w:numPr")) is not None:
            return True
    return bool(_MANUAL_NUM_RE.match(p.text))


def _para_pf(p, name: str):
    """Свойство абзацного форматирования: у самого абзаца, иначе по цепочке стилей.
    Проверяем `is not None`, а не истинность: явный нулевой отступ — это решение автора
    примера, и подменять его отступом стиля значит вернуть оформление, которого нет."""
    value = getattr(p.paragraph_format, name)
    if value is not None:
        return value
    return _style_prop(p.style, lambda st: getattr(st.paragraph_format, name))


def _caption_match(text: str):
    """(вид, разбор) для «Рисунок 1 — …» / «Табл. 2. …», иначе None. Слово обязано быть
    из списка: иначе подписью считался бы любой абзац, начинающийся с «В 2020 году»."""
    if len(text) > 200:
        return None
    m = _CAPTION_RE.match(text.strip())
    if not m:
        return None
    kind = _CAPTION_WORDS.get(m.group("word").lower().rstrip("."))
    return (kind, m) if kind else None


def _mode(values, default=None):
    """Самое частое непустое значение: оформление примера — это то, чего в нём больше,
    а не то, что стоит в первом попавшемся абзаце (первый обычно титульный)."""
    vals = [v for v in values if v is not None]
    return Counter(vals).most_common(1)[0][0] if vals else default


# ── оформление ────────────────────────────────────────────────────────────────

def style_from_sample(sample, *, allow_tags: bool = False) -> StyleProfile:
    """Готовый отчёт-образец → оформление. Бросает `SampleError`, брать нечего.

    `sample` — путь / bytes / file-like / Document; недоверенный DOCX, проверяется
    как шаблон. `allow_tags=True` — читать оформление и из документа с тегами
    (осознанный случай: у службы один и тот же файл и образец, и шаблон).

    Цена ошибки: молча вернуть пустой профиль значит собрать отчёт своим оформлением
    и уверить пользователя, что кафедральное применено, — поэтому «в примере не за что
    зацепиться» это исключение, а мелкие пропажи — строки в `notes`.
    """
    doc = _open_sample(sample, allow_tags)
    themes = _theme_fonts(doc)
    notes: list[str] = []

    section = doc.sections[-1]      # как и page_text_width_cm: у отчёта важна последняя
    page = PageSetup(
        width_cm=round(section.page_width / EMU_PER_CM, 2),
        height_cm=round(section.page_height / EMU_PER_CM, 2),
        left_cm=round(section.left_margin / EMU_PER_CM, 2),
        right_cm=round(section.right_margin / EMU_PER_CM, 2),
        top_cm=round(section.top_margin / EMU_PER_CM, 2),
        bottom_cm=round(section.bottom_margin / EMU_PER_CM, 2))
    if len(doc.sections) > 1:
        notes.append(f"в примере {len(doc.sections)} секций с разными полями — "
                     "взяты поля последней; у нашего документа секция одна")

    paras = list(doc.paragraphs)
    heads: dict[int, list] = {}
    plain = []
    for p in paras:
        if not p.text.strip():
            continue
        lvl = _outline_level(doc, p)
        if lvl:
            heads.setdefault(lvl, []).append(p)
        elif _caption_match(p.text) is None:
            plain.append(p)

    body = _body_text(doc, plain, themes, notes)
    if not heads:
        raise SampleError(
            "в примере нет ни одного заголовка: ни стиля уровня заголовка, ни абзаца "
            "с outlineLvl. Оформление заголовков — главное, ради чего берут образец; "
            "похоже, это не отчёт (или заголовки в нём набраны жирным вручную, а такой "
            "вид повторить нечем)")
    headings = tuple(_heading_look(lvl, ps, doc, themes)
                     for lvl, ps in sorted(heads.items()))
    captions = _captions(paras, notes)
    numbered = any(h.numbered for h in headings)
    if numbered and not all(h.numbered for h in headings):
        notes.append("нумерованы не все уровни заголовков — наша нумерация сквозная "
                     "по уровням от первого")
    _note_unused_styles(doc, notes)
    return StyleProfile(page=page, body=body, headings=headings, captions=captions,
                        numbered_headings=numbered, notes=tuple(notes))


def _body_text(doc, plain, themes: dict, notes: list) -> BodyText:
    """Основной текст — по самим абзацам, а не по стилю Normal: в кафедральных
    документах Normal сплошь и рядом остаётся заводским (Calibri 11), а весь текст
    набран прямым форматированием, и профиль вышел бы не тот, что человек видит."""
    default = BodyText()
    if not plain:
        notes.append("в примере нет обычных абзацев — основной текст взят по умолчанию")
        return default
    font = _mode([_para_font(doc, p, themes) for p in plain], default.font)
    size = _mode([_para_size(p) for p in plain], default.size_pt)
    line = _mode([_para_pf(p, "line_spacing") for p in plain], default.line_spacing)
    first = _mode([_para_pf(p, "first_line_indent") for p in plain])
    align = _mode([_para_align(p) for p in plain], default.align)
    return BodyText(font=font, size_pt=float(size),
                    line_spacing=float(line) if line else default.line_spacing,
                    first_line_cm=round(first / EMU_PER_CM, 2) if first is not None else 0.0,
                    align=align if align in _ALIGN else default.align)


def _heading_look(level: int, ps: list, doc, themes: dict) -> HeadingLook:
    return HeadingLook(
        level=level,
        font=_mode([_para_font(doc, p, themes) for p in ps]),
        size_pt=_mode([_para_size(p) for p in ps]),
        bold=_mode([_para_bold(p) for p in ps]),
        align=_mode([_para_align(p) for p in ps]),
        numbered=sum(_numbered(p) for p in ps) * 2 >= len(ps))


def _captions(paras, notes: list) -> dict:
    """Вид подписей: слово, разделитель и выключка — из самих подписей примера."""
    seen: dict[str, list] = {"figure": [], "table": []}
    for p in paras:
        hit = _caption_match(p.text)
        if hit is not None:
            seen[hit[0]].append((hit[1], p))
    out: dict = {}
    word_of = {"figure": "рисункам", "table": "таблицам"}
    for kind, hits in seen.items():
        if not hits:
            notes.append(f"в примере нет ни одной подписи к {word_of[kind]} — "
                         "вид подписи останется нашим (styles.yaml)")
            continue
        word = _mode([m.group("word") for m, _ in hits])
        sep = _mode([m.group("sep") for m, _ in hits]) or "—"
        align = _mode([_para_align(p) for _, p in hits])
        # «Рис. 1» и «Рисунок 1» — разные подписи: точка после сокращения входит в вид
        dot = "." if any(m.group(0).lstrip().startswith(m.group("word") + ".")
                         for m, _ in hits) else ""
        tail = f" {sep} " if sep in ("—", "–", "-") else f"{sep} "
        out[kind] = f"{word}{dot} {{n}}{tail}{{caption}}"
        if align:
            out[f"{kind}_align"] = align
        if any("." in m.group("num") or "-" in m.group("num") for m, _ in hits):
            notes.append(f"подписи «{word} 1.2» нумерованы по главам — наши поля SEQ "
                         "считают сквозным номером, номер главы в них не подставить")
    return out


def _note_unused_styles(doc, notes: list) -> None:
    """Стили, которыми в примере *набраны абзацы*, а мы их не применим. Считать по
    styles.xml нельзя: заводской шаблон Word объявляет полторы сотни стилей, из которых
    не использован ни один, и человек получил бы список ни о чём."""
    from .template import STYLE_SPECS
    known = set(STYLE_SPECS) | {"Normal", "Header", "Footer", "List Paragraph",
                                "Title", "Subtitle", "No Spacing"}
    used = {p.style.name for p in doc.paragraphs
            if p.text.strip() and p.style is not None and p.style.name}
    extra = sorted(n for n in used if n not in known and not re.fullmatch(r"Heading \d", n))
    if extra:
        notes.append(f"абзацы примера набраны стилями, которых рендер не знает: "
                     f"{', '.join(extra[:5])} — оформление берётся только из "
                     "Heading N, Caption, Quote, Code")


def apply_style(doc, profile: StyleProfile) -> None:
    """Наложить оформление примера на наш документ (обычно `template.blank_document()`).

    Меняет документ на месте: поля секций, стиль Normal, стили `Heading N` и нумерацию
    разделов. Стили, которых нет, создаются тем же механизмом, что у документа без
    шаблона (`template.ensure_styles`) — второй набор запасных стилей развёлся бы
    с первым на первой же правке.

    Подписи сюда не входят намеренно: их слова живут в `styles.yaml`, а не в стиле
    Caption, и едут отдельно — `render(style=profile.style_overrides())`.
    """
    apply_page_setup(doc, profile.page)
    apply_body_text(doc, profile.body)
    ensure_styles(doc)
    for look in profile.headings:
        name = f"Heading {look.level}"
        if not ops.has_style(doc, name):
            continue
        style = doc.styles[name]
        _apply_font(style, font=look.font, size_pt=look.size_pt, bold=look.bold)
        if look.align in _ALIGN:
            ops.set_child(style.element.get_or_add_pPr(), "w:jc", val=_ALIGN[look.align])
    if profile.numbered_headings:
        apply_heading_numbering(doc, max(h.level for h in profile.headings))


def apply_heading_numbering(doc, levels: int = 3) -> bool:
    """Нумерация разделов («1.», «1.1.») силами Word. False — не вышло (нет numbering.xml)."""
    # уровень нумерации и уровень заголовка обязаны совпадать: пропусти мы Heading 2
    # (в шаблоне его может не быть), «1.1» встало бы на заголовки третьего уровня
    levels = min(levels, MAX_LEVEL)
    names = [f"Heading {n}" for n in range(1, levels + 1)]
    if not all(ops.has_style(doc, n) for n in names):
        return False
    num_id = ops.Numbering(doc).new_heading_numbering([ops.style_ids(doc)[n] for n in names])
    if num_id is None:
        return False
    for i, name in enumerate(names):
        ppr = doc.styles[name].element.get_or_add_pPr()
        numpr = ops.set_child(ppr, "w:numPr")
        ops.set_child(numpr, "w:ilvl", val=i)
        ops.set_child(numpr, "w:numId", val=num_id)
    return True


def document_from_sample(sample, *, page_numbers: bool = True, allow_tags: bool = False):
    """Пустой документ, оформленный по примеру. → (Document, StyleProfile).

    Профиль возвращается вместе с документом не для красоты: без
    `render(style=profile.style_overrides())` подписи к рисункам выйдут нашими,
    и половина работы пропадёт незаметно.
    """
    from .template import blank_document
    profile = style_from_sample(sample, allow_tags=allow_tags)
    doc = blank_document(page=profile.page, body=profile.body, page_numbers=page_numbers)
    apply_style(doc, profile)
    return doc, profile


# ── строение ──────────────────────────────────────────────────────────────────

def outline_from_sample(sample, *, allow_tags: bool = True) -> list[dict]:
    """Разделы примера по порядку: [{level, text, number}]. Это строение, не оформление.

    `number` — номер, набранный в самом тексте («2.1»), или None: нумерация полем Word
    в тексте абзаца не хранится, и выдумывать её здесь нельзя. `text` — заголовок без
    этого номера, то есть то, что можно положить в свой отчёт как имя раздела.

    Заголовки берутся из тела документа по порядку; в таблицах и колонтитулах их не ищем
    (там их не бывает, а найденное было бы шапкой листа, а не разделом работы).
    """
    doc = _open_sample(sample, allow_tags)
    out = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        lvl = _outline_level(doc, p)
        if not lvl:
            continue
        m = _MANUAL_NUM_RE.match(text)
        out.append({"level": lvl,
                    "text": m.group("rest").strip() if m else text,
                    "number": m.group("num") if m else None})
    return out
