"""
render — подстановка значений в шаблон.

Для каждого абзаца (тело, таблицы, колонтитулы, текстовые поля):
  1. inline-значения (str / Text) вписываются внутрь runs на место тега —
     форматирование соседних слов и самого тега сохраняется;
  2. блочные значения (Markdown / Code / Image / Table / Blocks) — тег
     вырезается, блоки вставляются абзацами после абзаца с тегом;
     если абзац после вырезания опустел — он удаляется;
  3. тег без значения — вырезается, ключ попадает в RenderResult.unfilled.
Ошибка в значении (битая картинка, нет файла) — HokokuError и рендер прерывается;
с on_error="skip" тег пропускается, беда пишется в RenderResult.errors, остальное собирается.
"""
from __future__ import annotations

import copy
import io
import os
import subprocess
import unicodedata

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import docx_ops as ops
from . import markdown as md
from .images import count_pages, drawio_to_png, fit, natural_width_cm, to_raster
from .model import (Blocks, Code, Diagram, Formula, Image, Markdown, PageBreak, RenderResult, Table, Toc,
                    HokokuError, Text)
from .styles import get_style
from .safety import DocxValidationError, safe_join
from .tags import find_tags, norm_key
from .walker import ParaLoc, iter_paragraphs, marked_paragraphs, open_document

DXA_PER_CM = 567


class _Ctx:
    def __init__(self, doc, result: RenderResult, style: dict, images_dir: str | None,
                 strict_paths: bool, on_error: str = "raise",
                 drawio_timeout: float | None = None):
        self.doc = doc
        self.result = result
        self.style = style
        self.numbering = ops.Numbering(doc)
        self.images_dir = images_dir
        self.strict_paths = strict_paths
        self.on_error = on_error
        self.drawio_timeout = drawio_timeout
        self.page_w = ops.page_text_width_cm(doc)
        self.page_h = ops.page_text_height_cm(doc)
        self.ref_fields: list = []          # (w:t, имя) — кэш номеров ставим в конце
        self.current_key: str | None = None
        self.used_keys: set = set()         # какие значения нашли свой тег


def render(template, values: dict, output=None, *,
           images_dir: str | None = None, style: dict | None = None,
           strict_paths: bool = False, figure_caption: str | None = None,
           table_caption: str | None = None, on_error: str = "raise",
           drawio_timeout: float | None = None) -> RenderResult:
    """
    template  — путь / bytes / Document (переданный Document не меняется — рендерится копия);
    values    — {ключ: str | int | float | bool | Text | Markdown | Code | Image | Table | Blocks};
                значения без тега в шаблоне возвращаются в RenderResult.unknown_keys;
    output    — путь, file-like или None (тогда результат — RenderResult.data, bytes).
    images_dir — каталог файлов для `![…](имя)` в Markdown (пути наружу запрещены);
    strict_paths=True — и `Image(source=путь)` обязан лежать внутри images_dir.
    style — перегрузки styles.yaml, например {"captions": {"figure": "Рисунок {n} – {caption}"}}.
    on_error — "raise" (по умолчанию: первая же битая картинка прерывает рендер) или "skip"
    (собрать что можно; каждая беда — записью {key, message} в RenderResult.errors).
    drawio_timeout — секунды на один запуск drawio CLI (None — умолчание images.drawio_to_png,
    120 с); серверу нужен свой, короче.
    """
    if on_error not in ("raise", "skip"):
        raise ValueError('on_error: "raise" или "skip"')
    st = get_style(style)
    if figure_caption:
        st["captions"]["figure"] = figure_caption
    if table_caption:
        st["captions"]["table"] = table_caption
    if hasattr(template, "paragraphs") and hasattr(template, "part"):
        buf = io.BytesIO()
        template.save(buf)
        template = buf.getvalue()
    doc = open_document(template)
    try:
        # запасной стиль заголовка создаётся глубоко внутри docx_ops.set_style, куда
        # словарь оформления не доходит; без этой передачи render(style={"headings": …})
        # молча не действовал бы на шаблон, в котором стилей Heading N нет.
        # Пометка на документе безопасна: переданный Document мы не рендерим, а копируем
        # через bytes (выше), поэтому два вызова с разным style не мешают друг другу
        doc._hokoku_heading_sizes = st["headings"]
    except AttributeError:                                # экзотический объект документа
        pass
    values = {norm_key(str(k)): v for k, v in values.items()}
    result = RenderResult(output=output if isinstance(output, str) else None)
    ctx = _Ctx(doc, result, st, images_dir, strict_paths, on_error, drawio_timeout)
    hot = marked_paragraphs(doc)
    for loc in iter_paragraphs(doc):
        if loc.paragraph._p not in hot:
            continue
        _process_paragraph(ctx, loc, values)
        _process_static_refs(ctx, loc)
    result.unfilled = sorted(set(result.unfilled))
    # значение, для которого в шаблоне нет тега: опечатка в ключе (в том числе у модели)
    # раньше просто исчезала — ни в unfilled (там только объявленные теги), ни в errors
    result.unknown_keys = sorted(set(values) - ctx.used_keys)
    if st["captions"]["fields"] and (result.figures or result.tables or result.formulas):
        ops.set_update_fields(doc)
    for t_elem, name in ctx.ref_fields:                    # кэш номеров для REF-полей
        plain = name[len("_Ref_"):] if name.startswith("_Ref_") else name
        n = result.refs.get(plain)
        if n is None and plain not in result.unresolved_refs:
            result.unresolved_refs.append(plain)           # имя как в тексте, без «_Ref_»
        t_elem.text = str(n) if n is not None else "?"
    if output is None:
        buf = io.BytesIO()
        doc.save(buf)
        result.data = buf.getvalue()
    elif isinstance(output, str):
        os.makedirs(os.path.dirname(os.path.abspath(output)) or ".", exist_ok=True)
        doc.save(output)
    else:
        doc.save(output)
    return result


# ── абзац с тегами ────────────────────────────────────────────────────────────

def _span_map(runs) -> list:
    """Карта «позиция в склеенном тексте абзаца → run»."""
    spans, pos = [], 0
    for r in runs:
        t = r.text or ""
        spans.append((pos, pos + len(t), r))
        pos += len(t)
    return spans


REF_TEXT_RE = md.REF_RE            # один разбор ссылок на все места (см. markdown.split_refs)


def _process_static_refs(ctx: _Ctx, loc: ParaLoc):
    """`{ref:имя}`, написанная в самом шаблоне (не пришедшая со значением), — тоже поле REF.
    Раньше такая ссылка оставалась в документе текстом и не попадала в unresolved_refs."""
    para = loc.paragraph
    if para._p.getparent() is None:                    # абзац удалён при подстановке
        return
    text = loc.text()
    if "{ref:" not in text:
        return
    runs = loc.runs()
    spans = _span_map(runs)
    for m in reversed(list(REF_TEXT_RE.finditer(text))):
        rpr = None
        for s, e, r in spans:
            if s <= m.start() < e:
                found = r._r.find(qn("w:rPr"))
                rpr = copy.deepcopy(found) if found is not None else None
                break
        _inline_replace(spans, m.start(), m.end(), "")
        _inline_insert_spans(ctx, para, spans, m.start(), md.parse_inline(m.group(0)), rpr)
        spans = _span_map(runs)


def _process_paragraph(ctx: _Ctx, loc: ParaLoc, values: dict):
    para = loc.paragraph
    text = loc.text()
    if "{{" not in text:
        return
    matches = find_tags(text)
    if not matches:
        return

    runs = loc.runs()
    spans = _span_map(runs)

    def run_at(i: int):
        for s, e, r in spans:
            if s <= i < e:
                return r
        return runs[0] if runs else None

    # Есть ли в абзаце текст помимо тегов? Тогда первый обычный абзац блочного
    # значения вписываем на место тега, чтобы не оставлять «Цель: .».
    other_text = find_tags(text) and _re_sub_tags(text).strip() != ""

    blocks: list = []
    base_rpr = None
    for m in reversed(matches):
        key = norm_key(m.group("key"))
        v = values.get(key)
        if key in values:
            ctx.used_keys.add(key)
        if isinstance(v, bool):
            v = "да" if v else "нет"
        elif isinstance(v, (int, float)):
            v = str(v)
        tag_run = run_at(m.start())
        if base_rpr is None and tag_run is not None:
            rpr = tag_run._r.find(qn("w:rPr"))
            base_rpr = copy.deepcopy(rpr) if rpr is not None else None
        inline_spans = None
        ctx.current_key = key
        if v is not None and _empty_value(v):
            if ctx.on_error != "skip":
                raise HokokuError(f"значение пустое: {key!r}")
            ctx.result.errors.append({"key": key, "message": "значение пустое"})
            v = None
            _inline_replace(spans, m.start(), m.end(), "")     # тег убираем, в unfilled не пишем
            spans = _span_map(runs)
            continue
        if isinstance(v, Text) and "{ref:" in v.text:
            v = Markdown(v.text.replace("\n", "  \n"))
        elif isinstance(v, str) and "{ref:" in v:
            v = Markdown(v)
        if v is None:
            ctx.result.unfilled.append(key)
            replacement = ""
        elif isinstance(v, str):
            replacement = v
        elif isinstance(v, Text):
            replacement = v.text
        else:
            replacement = ""
            if other_text:
                inline_spans, v = _split_leading_paragraph(ctx, v)
            if v is not None:
                blocks.insert(0, (key, v))
        end = m.end()
        if replacement == "" and v is not None and not inline_spans:
            # блочное значение вырезано из строки: одинокий знак препинания сразу за ним
            # («Цель: {{цель}}.») тоже убираем, чтобы не осталось «Цель: .»
            tail = text[end:end + 2]
            if tail[:1] in ".,;!" and (len(tail) < 2 or tail[1] in " \t"):
                end += 1
        _inline_replace(spans, m.start(), end, replacement)
        if inline_spans:
            _inline_insert_spans(ctx, para, spans, m.start(), inline_spans, base_rpr)
        spans = _span_map(runs)                       # карта после замены

    if not blocks:
        return
    p_elem = para._p
    ppr = p_elem.find(qn("w:pPr"))
    ref = p_elem
    for key, v in blocks:
        ctx.current_key = key
        try:
            ref = _emit_value(ctx, v, ref, para, ppr, base_rpr, loc)
        except Exception as e:                      # noqa: BLE001 — skip обязан собрать остальное
            if ctx.on_error != "skip" or isinstance(e, (KeyboardInterrupt, SystemExit)):
                raise
            ctx.result.errors.append({"key": key, "message": f"{type(e).__name__}: {e}"
                                      if not isinstance(e, HokokuError) else str(e)})
    vals = [v for _, v in blocks]
    if not _has_content(p_elem):
        parent = p_elem.getparent()
        siblings = [c for c in parent if c.tag == qn("w:p")]
        if not (parent.tag == qn("w:tc") and len(siblings) == 1):
            prev = p_elem.getprevious()
            parent.remove(p_elem)
            # «Блок-схема алгоритма:» перед картинкой — прилипает к ней
            if ctx.style["image"]["keep_intro"] and prev is not None and prev.tag == qn("w:p") \
                    and _starts_with_picture(vals) and _para_text(prev).rstrip().endswith(":"):
                ops.keep_with_next(prev)
    elif _starts_with_picture(vals):
        ops.keep_with_next(p_elem)


def _para_text(p_elem) -> str:
    return "".join(t.text or "" for t in p_elem.iter(qn("w:t")))


def _starts_with_picture(blocks) -> bool:
    if not blocks:
        return False
    b = blocks[0]
    if isinstance(b, (Image, Diagram)):
        return True
    if isinstance(b, (Formula, Toc)):
        return False
    if isinstance(b, Blocks) and b.items:
        return _starts_with_picture([b.items[0]])
    if isinstance(b, Markdown):
        parsed = md.parse(b.text)
        return bool(parsed) and isinstance(parsed[0], md.ImgBlock)
    if isinstance(b, _ParsedMarkdown):
        return bool(b.blocks) and isinstance(b.blocks[0], md.ImgBlock)
    return False


def _re_sub_tags(text: str) -> str:
    from .tags import TAG_RE
    return TAG_RE.sub("", text)


def _split_leading_paragraph(ctx, v):
    """Если значение начинается с обычного абзаца — вернуть (его spans, остаток)."""
    if isinstance(v, Markdown):
        blocks = md.parse(v.text)
        if blocks and isinstance(blocks[0], md.Para) and blocks[0].kind == "p":
            rest = blocks[1:]
            return blocks[0].spans, (_ParsedMarkdown(rest, v.images_dir) if rest else None)
        return None, v
    if isinstance(v, Blocks) and v.items:
        first = v.items[0]
        if isinstance(first, (str, Text)):
            t = first if isinstance(first, str) else first.text
            rest = v.items[1:]
            return md.split_refs(t), (Blocks(rest) if rest else None)
        if isinstance(first, Markdown):
            sp, rest_md = _split_leading_paragraph(ctx, first)
            if sp is not None:
                rest = ([rest_md] if rest_md else []) + v.items[1:]
                return sp, (Blocks(rest) if rest else None)
    return None, v


class _ParsedMarkdown:
    """Уже разобранные блоки markdown (остаток после вписанного первого абзаца)."""
    def __init__(self, blocks, images_dir):
        self.blocks = blocks
        self.images_dir = images_dir


def _inline_insert_spans(ctx, para, spans, at: int, md_spans, base_rpr):
    """Вставить runs с форматированием на позицию at в абзаце (внутри runs)."""
    # ищем run, в котором стоит позиция at, и делим его
    for s0, e0, r in spans:
        if s0 <= at <= e0:
            old = r.text or ""
            head, tail = old[:at - s0], old[at - s0:]
            r.text = head
            new_p = []
            tmp = ops.OxmlElement("w:p")
            _add_spans(ctx, tmp, md_spans, base_rpr, para.part)
            anchor = r._r
            for el in list(tmp):
                anchor.addnext(el)
                anchor = el
            if tail:
                anchor.addnext(ops.make_run(tail, base_rpr))
            return
    tmp = ops.OxmlElement("w:p")
    _add_spans(ctx, tmp, md_spans, base_rpr, para.part)
    for el in list(tmp):
        para._p.append(el)


def _inline_replace(spans, start: int, end: int, text: str):
    """Заменить [start, end) в склеенном тексте runs, трогая только затронутые runs."""
    first = last = None
    for i, (s, e, _) in enumerate(spans):
        if first is None and s <= start < e:
            first = i
        if s < end <= e:
            last = i
            break
    if first is None:
        return
    if last is None:
        last = len(spans) - 1
    s0, _, r0 = spans[first]
    if first == last:
        old = r0.text or ""
        r0.text = old[:start - s0] + text + old[end - s0:]
        return
    r0.text = (r0.text or "")[:start - s0] + text
    for j in range(first + 1, last):
        spans[j][2].text = ""
    sl, _, rl = spans[last]
    rl.text = (rl.text or "")[end - sl:]


def _has_content(p_elem) -> bool:
    if "".join(t.text or "" for t in p_elem.iter(qn("w:t"))).strip():
        return True
    for tag in ("w:drawing", "w:pict", "w:br", "w:tab", "w:sym", "w:fldSimple", "w:fldChar"):
        if p_elem.find(".//" + qn(tag)) is not None:
            return True
    return False


# ── блочные значения ──────────────────────────────────────────────────────────

def _emit_value(ctx: _Ctx, v, ref, para, ppr, base_rpr, loc: ParaLoc):
    if isinstance(v, str):
        v = Text(v)
    if isinstance(v, Text):
        p = ops.new_paragraph_after(ref, ppr)
        # `{ref:имя}` внутри Blocks уходила в документ буквальной строкой и мимо
        # unresolved_refs — молча. Здесь она такое же поле REF, как в обычном тексте;
        # остальную разметку не разбираем, Text есть Text (см. markdown.split_refs)
        _add_spans(ctx, p, md.split_refs(v.text), base_rpr, para.part)
        return p
    if isinstance(v, Markdown):
        return _emit_markdown(ctx, md.parse(v.text), ref, para, ppr, base_rpr, loc,
                              v.images_dir or ctx.images_dir)
    if isinstance(v, _ParsedMarkdown):
        return _emit_markdown(ctx, v.blocks, ref, para, ppr, base_rpr, loc,
                              v.images_dir or ctx.images_dir)
    if isinstance(v, Code):
        return _emit_code(ctx, v, ref, ppr)
    if isinstance(v, Image):
        return _emit_image(ctx, v, ref, para, ppr, loc)
    if isinstance(v, Diagram):
        pages = [v.page] if v.page is not None else list(range(1, count_pages(v.xml) + 1))
        kw = {} if ctx.drawio_timeout is None else {"timeout": ctx.drawio_timeout}
        try:
            sheets = [drawio_to_png(v.xml, p, **kw) for p in pages]
        except subprocess.TimeoutExpired:
            raise HokokuError("drawio не уложился в таймаут")
        except (ValueError, OSError) as e:
            raise HokokuError(str(e))
        # все страницы mxfile — листы одного рисунка: «Рисунок N (лист k из m)»
        return _emit_image(ctx, Image(sheets[0], caption=v.caption, width_cm=v.width_cm,
                                      align=v.align, ref=v.ref),
                           ref, para, ppr, loc, sheets=sheets)
    if isinstance(v, Table):
        rows = [[md.parse_inline(str(c)) for c in row] for row in v.rows]
        return _emit_table(ctx, rows, v.header, v.caption, v.align, ref, para, ppr, base_rpr, loc,
                           v.col_widths_cm, v.ref)
    if isinstance(v, PageBreak):
        return ops.add_page_break(ref)
    if isinstance(v, Formula):
        return _emit_formula(ctx, v.latex, v.numbered, v.ref, ref, ppr)
    if isinstance(v, Toc):
        p = ops.add_toc(ctx.doc, ref, v.levels, v.title, ppr)
        if v.title and "{ref:" in v.title:
            # заголовок оглавления — отдельный абзац перед полем TOC, а не содержимое
            # поля: поле REF в нём законно и обновление оглавления его не трогает.
            # Сам абзац add_toc не возвращает (возвращает абзац с TOC), берём соседа слева
            title_p = p.getprevious()
            if title_p is not None:
                _register_ref_fields(ctx, title_p)
        return p
    if isinstance(v, Blocks):
        for item in v.items:
            try:
                ref = _emit_value(ctx, item, ref, para, ppr, base_rpr, loc)
            except HokokuError as e:
                if ctx.on_error != "skip":
                    raise
                # хвост Blocks не должен пропадать из-за одного битого элемента
                ctx.result.errors.append({"key": ctx.current_key, "message": str(e)})
        return ref
    raise HokokuError(f"неподдерживаемый тип значения: {type(v).__name__}")


def _register_ref_fields(ctx: _Ctx, elem):
    """Найти в поддереве поля REF и запомнить их: кэш номеров ставится в конце рендера,
    когда все номера известны. Ячейки таблицы-значения идут мимо _add_spans (их пишет
    ops.add_table напрямую), поэтому обход именно по поддереву, а не по абзацу."""
    for f in elem.iter(qn("w:fldSimple")):
        instr = f.get(qn("w:instr")) or ""
        if instr.strip().startswith("REF "):
            name = instr.split()[1]
            t = f.find(".//" + qn("w:t"))
            if t is not None and (t, name) not in ctx.ref_fields:
                ctx.ref_fields.append((t, name))


def _add_spans(ctx, p, spans, base_rpr, part, **extra):
    """add_spans + регистрация REF-полей для подстановки кэша номеров в конце."""
    ops.add_spans(ctx.doc, p, spans, base_rpr, part, **extra)
    _register_ref_fields(ctx, p)


def _emit_markdown(ctx, blocks, ref, para, ppr, base_rpr, loc, images_dir):
    doc = ctx.doc
    part = para.part
    lists: dict[int, tuple[str, int | None]] = {}   # уровень → (вид, numId); сброс между списками
    counters: dict[int, int] = {}                    # для запасного варианта без numbering.xml
    for b in blocks:
        if not isinstance(b, md.Para) or b.kind not in ("ul", "ol"):
            lists = {}
        if isinstance(b, md.Para):
            p = ops.new_paragraph_after(ref, ppr)
            extra = {}
            if b.kind.startswith("h"):
                extra = ops.style_heading(doc, p, int(b.kind[1]), base_rpr)
                if extra:
                    extra["size_pt"] = ctx.style["headings"].get(int(b.kind[1]), extra["size_pt"])
                _add_spans(ctx, p, b.spans, base_rpr if not extra else None, part, **extra)
            elif b.kind == "quote":
                ops.style_quote(doc, p, b.level)
                _add_spans(ctx, p, b.spans, base_rpr, part, italic=True)
            elif b.kind in ("ul", "ol"):
                cur = lists.get(b.level)
                if cur is None or cur[0] != b.kind:
                    cur = (b.kind, ctx.numbering.new_list(b.kind == "ol", b.ordered_start))
                    lists[b.level] = cur
                    counters[b.level] = b.ordered_start
                    for deeper in [l for l in lists if l > b.level]:
                        del lists[deeper]
                else:
                    counters[b.level] = counters.get(b.level, 0) + 1
                ops.set_list_item(doc, p, cur[1], b.level, b.kind == "ol", counters[b.level])
                _add_spans(ctx, p, b.spans, base_rpr, part)
            else:
                _add_spans(ctx, p, b.spans, base_rpr, part)
            ref = p
        elif isinstance(b, md.CodeBlock):
            ref = _emit_code(ctx, Code(b.text, b.lang), ref, ppr)
        elif isinstance(b, md.Hr):
            p = ops.new_paragraph_after(ref, None)
            ops.style_hr(p)
            ref = p
        elif isinstance(b, md.ImgBlock):
            src = _resolve_image(b.src, images_dir)
            ref = _emit_image(ctx, Image(src, caption=b.caption or None, width_cm=b.width_cm,
                                         align=b.align or "center"), ref, para, ppr, loc)
        elif isinstance(b, md.TableBlock):
            ref = _emit_table(ctx, b.rows, b.header, None, b.align, ref, para, ppr, base_rpr, loc)
        elif isinstance(b, md.MathBlock):
            ref = _emit_formula(ctx, b.latex, True, None, ref, ppr)
    return ref


def _emit_formula(ctx: _Ctx, latex: str, numbered: bool, ref_name, ref, ppr):
    if not latex.strip():
        # пустая формула рисовалась пустым местом с номером «(1)»: номер потрачен,
        # в отчёте пусто, и ни unfilled, ни errors об этом не говорили
        raise HokokuError("формула пустая")
    cap = ctx.style["captions"]
    n = 0
    bookmark = None
    name = None
    if numbered:
        ctx.result.formulas += 1
        n = ctx.result.formulas
        name = ref_name or ctx.current_key
        if name:
            ctx.result.refs[name] = n
            bookmark = f"_Ref_{name}"
    try:
        return ops.add_formula(ctx.doc, ref, latex, numbered=numbered, n=n,
                               seq_name=cap.get("seq_formula", "Формула") if cap["fields"] else None,
                               bookmark=bookmark, page_w_cm=ctx.page_w, ppr_template=ppr)
    except Exception as e:                                        # noqa: BLE001
        # номер забираем обратно: при on_error="skip" битая формула иначе съедала его,
        # и следующая получала «3», хотя в документе она вторая — Word перенумерует поля
        # SEQ при обновлении, и {ref:} укажет не на ту формулу
        if numbered:
            ctx.result.formulas -= 1
            if name:
                ctx.result.refs.pop(name, None)
        raise HokokuError(f"формула не разобрана: {latex!r}: {e}")


def _emit_code(ctx: _Ctx, c: Code, ref, ppr):
    cs = ctx.style["code"]
    return ops.add_code_lines(
        ctx.doc, ref, c.text, ppr, lang=c.lang, font=cs["font"], size_pt=cs["size_pt"],
        highlight=cs["highlight"] if c.highlight is None else c.highlight,
        line_numbers=cs["line_numbers"] if c.line_numbers is None else c.line_numbers,
        style_name=cs["style"])


def _resolve_image(src: str, images_dir: str | None) -> str:
    if src.startswith("file://"):
        src = src[7:]
    if os.path.isabs(src):
        if images_dir:
            return safe_join(images_dir, os.path.basename(src))
        return src
    if not images_dir:
        raise HokokuError(f"картинка {src!r}: не задан images_dir")
    try:
        return safe_join(images_dir, src)
    except DocxValidationError as e:
        raise HokokuError(str(e))


def _empty_value(v) -> bool:
    """Значение, из которого нечего вставить. Раньше такое молча съедало тег: ни в `unfilled`
    (там только теги вовсе без значения), ни в `errors` — «модель ничего не вернула» выглядело
    ровно как «тег заполнен» (решение владельца 2026-08-29: считать ошибкой)."""
    if isinstance(v, str):
        return not v.strip()
    if isinstance(v, (Text, Markdown)):
        return not v.text.strip()
    if isinstance(v, Code):
        return not v.text.strip()
    if isinstance(v, Table):
        return not v.rows or all(not any(str(c).strip() for c in row) for row in v.rows)
    if isinstance(v, Blocks):
        return not v.items
    return False


def _numbered(caption, loc: ParaLoc) -> bool:
    """Нумеровать и подписывать? `caption=False` — картинка/таблица без подписи (логотип,
    декоративная врезка); в колонтитулах не нумеруем никогда — иначе логотип в шапке
    становится «Рисунок 1» и сдвигает нумерацию всего отчёта."""
    return caption is not False and loc.where not in ("header", "footer")


def _grid_span(tc) -> int:
    gs = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan"))
    try:
        return max(1, int(gs.get(qn("w:val"))))
    except (AttributeError, TypeError, ValueError):
        return 1


def _grid_width_cm(tc) -> float:
    """Ширина ячейки по w:tblGrid — столбцы, которые она занимает (с учётом w:gridSpan)."""
    tr = tc.getparent()
    tbl = tr.getparent() if tr is not None else None
    grid = tbl.find(qn("w:tblGrid")) if tbl is not None else None
    if grid is None:
        return 0.0
    cells = tr.findall(qn("w:tc"))
    if tc not in cells:
        return 0.0
    start = sum(_grid_span(c) for c in cells[:cells.index(tc)])
    total = 0
    for gc in grid.findall(qn("w:gridCol"))[start:start + _grid_span(tc)]:
        try:
            total += int(gc.get(qn("w:w")))
        except (TypeError, ValueError):
            return 0.0
    return total / DXA_PER_CM


def _cell_width_cm(loc: ParaLoc, ctx: _Ctx) -> float:
    tc = loc.paragraph._p.getparent()
    while tc is not None and tc.tag != qn("w:tc"):
        tc = tc.getparent()
    if tc is None:
        return ctx.page_w / 2
    tcw = tc.find(qn("w:tcPr") + "/" + qn("w:tcW"))
    if tcw is not None and tcw.get(qn("w:type"), "dxa") == "dxa":
        try:
            return max(2.0, int(tcw.get(qn("w:w"))) / DXA_PER_CM - 0.5)
        except (TypeError, ValueError):
            pass
    # ячейка без w:tcW (типовой титульник из Word): ширина — из сетки таблицы,
    # а если и сетки нет — полоса набора, делённая на число ячеек строки
    w = _grid_width_cm(tc)
    if not w:
        w = ctx.page_w / max(1, len(tc.getparent().findall(qn("w:tc"))))
    return max(2.0, w - 0.5)


def _emit_image(ctx: _Ctx, img: Image, ref, para, ppr, loc: ParaLoc, sheets: list | None = None):
    """sheets — готовые листы одного рисунка (страницы схемы); иначе картинка одна."""
    if ctx.strict_paths and isinstance(img.source, str):
        if not ctx.images_dir:
            raise HokokuError("strict_paths: не задан images_dir")
        try:
            img = Image(safe_join(ctx.images_dir, img.source), img.caption, img.width_cm,
                        img.align, img.ref)
        except DocxValidationError as e:
            raise HokokuError(str(e))
    try:
        pieces = [to_raster(s) for s in sheets] if sheets else [to_raster(img.read())]
    except ValueError as e:
        raise HokokuError(str(e))
    ist = ctx.style["image"]
    cap = ctx.style["captions"]
    max_w = _cell_width_cm(loc, ctx) if loc.in_table else ctx.page_w
    max_h = ctx.page_h - 2.0
    want_w = img.width_cm
    if want_w is None and ist["default_width"] != "natural":
        want_w = float(ist["default_width"])
    try:
        sizes = [fit(p, max_w, max_h, want_w) for p in pieces]
    except Exception as e:
        raise HokokuError(f"картинка не читается: {e}")
    def put(piece, w, h, at):
        try:
            return ops.add_picture(ctx.doc, at, para._parent, piece, w, h, img.align)
        except HokokuError:
            raise
        except Exception as e:                      # python-docx не знает формат (WEBP и т.п.)
            raise HokokuError(f"картинка не вставлена: {type(e).__name__}: {e}")

    if not _numbered(img.caption, loc):
        for piece, (w, h) in zip(pieces, sizes):
            ref = put(piece, w, h, ref)
        return ref
    ctx.result.figures += 1
    n = ctx.result.figures
    name = img.ref or ctx.current_key
    if name:
        ctx.result.refs[name] = n
    for i, (piece, (w, h)) in enumerate(zip(pieces, sizes)):
        ref = put(piece, w, h, ref)
        suffix = cap["sheet"].format(k=i + 1, total=len(pieces)) if len(pieces) > 1 else ""
        ref = ops.add_caption(ctx.doc, ref, cap["figure"], n, img.caption, align=img.align,
                              seq_name=cap["seq_figure"] if cap["fields"] else None,
                              bookmark=(f"_Ref_{name}" if name and i == 0 else None),
                              suffix=suffix, repeat=(i > 0))
        _register_ref_fields(ctx, ref)        # «ср. {ref:схема}» в самой подписи
    return ref


def _emit_table(ctx: _Ctx, rows, header, caption, align, ref, para, ppr, base_rpr, loc,
                col_widths_cm=None, ref_name=None):
    if not rows:
        return ref
    cap = ctx.style["captions"]
    # тег в пункте списка: таблица и подпись над ней встают под своим пунктом,
    # а не у левого поля
    indent = ops.left_indent_dxa(ctx.doc, para._p)
    if _numbered(caption, loc):
        ctx.result.tables += 1
        n = ctx.result.tables
        name = ref_name or ctx.current_key
        if name:
            ctx.result.refs[name] = n
        ref = ops.add_caption(ctx.doc, ref, cap["table"], n, caption, align=cap["table_align"],
                              seq_name=cap["seq_table"] if cap["fields"] else None,
                              bookmark=(f"_Ref_{name}" if name else None), indent_dxa=indent)
        _register_ref_fields(ctx, ref)        # «см. {ref:рис}» в самой подписи
        ops.keep_with_next(ref)
    max_w = max(2.0, (_cell_width_cm(loc, ctx) if loc.in_table else ctx.page_w) - indent / DXA_PER_CM)
    ts = ctx.style["table"]
    tbl = ops.add_table(ctx.doc, ref, rows, header, para.part, align, base_rpr, col_widths_cm, max_w,
                        header_fill=ts["header_fill"], header_center=bool(ts["header_center"]),
                        min_col_cm=float(ts["min_col_cm"]), indent_dxa=indent)
    _register_ref_fields(ctx, tbl)          # «см. {ref:рис}» в ячейке значения-таблицы
    # пустой абзац после таблицы, иначе Word склеивает соседние таблицы
    p = ops.new_paragraph_after(tbl, None)
    ops.set_spacing(p, before=0, after=0)
    return p
