"""
docx_ops — низкоуровневые операции над OOXML, которых нет в python-docx:
абзацы после заданного, runs с форматированием, гиперссылки, настоящие
списки (w:numPr), стили с запасным вариантом, код, картинки, таблицы.
"""
from __future__ import annotations

import copy
import io
import re

from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt
from docx.text.paragraph import Paragraph

from kyotsu.text import text_width
from .images import EMU_PER_CM
from .markdown import Span, split_refs

CODE_FONT = "Courier New"
CODE_SIZE_PT = 10
HEADING_SIZES = {1: 16, 2: 14, 3: 13, 4: 12, 5: 12, 6: 12}
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


# ── абзацы ────────────────────────────────────────────────────────────────────

def new_paragraph_after(ref_elem, ppr_template=None):
    """Новый <w:p> сразу после ref_elem (абзац или таблица)."""
    p = OxmlElement("w:p")
    if ppr_template is not None:
        ppr = copy.deepcopy(ppr_template)
        for tag in ("w:numPr", "w:pStyle", "w:outlineLvl", "w:keepNext", "w:pageBreakBefore"):
            for old in ppr.findall(qn(tag)):
                ppr.remove(old)
        p.append(ppr)
    ref_elem.addnext(p)
    return p


def get_ppr(p_elem):
    ppr = p_elem.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p_elem.insert(0, ppr)
    return ppr


# Порядок дочерних элементов по схеме OOXML — Word отказывается открывать файл
# («нечитаемое содержимое»), если, например, w:jc стоит раньше w:spacing.
_ORDER = {
    "w:pPr": ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
              "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs",
              "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
              "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
              "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
              "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
              "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
              "w:sectPr", "w:pPrChange"],
    "w:rPr": ["w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps",
              "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint",
              "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden", "w:color", "w:spacing",
              "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
              "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
              "w:eastAsianLayout", "w:specVanish", "w:oMath"],
    "w:tblPr": ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
                "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
                "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"],
    "w:lvl": ["w:start", "w:numFmt", "w:lvlRestart", "w:pStyle", "w:isLgl", "w:suff", "w:lvlText",
              "w:lvlPicBulletId", "w:legacy", "w:lvlJc", "w:pPr", "w:rPr"],
    "w:abstractNum": ["w:nsid", "w:multiLevelType", "w:tmpl", "w:name", "w:styleLink",
                      "w:numStyleLink", "w:lvl"],
    "w:numPr": ["w:ilvl", "w:numId", "w:numberingChange", "w:ins"],
    "w:num": ["w:abstractNumId", "w:lvlOverride"],
}
_ORDER_Q = {qn(k): [qn(t) for t in v] for k, v in _ORDER.items()}


def insert_ordered(parent, el):
    """Вставить el в parent на место, предписанное схемой (или в конец)."""
    order = _ORDER_Q.get(parent.tag)
    if order is None or el.tag not in order:
        parent.append(el)
        return el
    rank = order.index(el.tag)
    for child in parent:
        if child.tag in order and order.index(child.tag) > rank:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def set_child(parent, tag: str, **attrs):
    """Заменить/создать дочерний элемент tag с атрибутами (w:*) в правильном месте."""
    for old in parent.findall(qn(tag)):
        parent.remove(old)
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return insert_ordered(parent, el)


def style_ids(doc) -> dict:
    """Словарь «имя стиля → id», один раз на документ. python-docx ищет стиль по имени
    XPath'ом по всему styles.xml, а спрашиваем мы его на каждый заголовок, пункт списка
    и таблицу — на отчёте это была четверть времени рендера."""
    ids = getattr(doc, "_hokoku_style_ids", None)
    if ids is None:
        ids = {s.name: s.style_id for s in doc.styles if s.name}
        try:
            doc._hokoku_style_ids = ids
        except AttributeError:                       # экзотический объект документа — без кэша
            pass
    return ids


def has_style(doc, name: str) -> bool:
    return name in style_ids(doc)


def set_style(doc, p_elem, name: str) -> bool:
    """Стиль абзаца по имени. Чего в документе нет — создаём, если это наш запасной
    (`template.STYLE_SPECS`: Heading N, Code, Caption, Quote): один механизм на «шаблон
    без стилей» и «документ без шаблона». Остальные имена (List Paragraph, TOC Heading,
    Hyperlink) сочинять нельзя — их вид задаёт шаблон, и подделка была бы хуже прямого
    форматирования, которым эти места обходятся сегодня."""
    style_id = style_ids(doc).get(name)
    if style_id is None:
        from .template import ensure_style           # поздний импорт: template знает docx_ops
        style_id = ensure_style(doc, name)
    if style_id is None:
        return False
    set_child(get_ppr(p_elem), "w:pStyle", val=style_id)
    return True


def set_alignment(p_elem, align: str):
    set_child(get_ppr(p_elem), "w:jc", val=align)


def set_spacing(p_elem, before: int | None = None, after: int | None = None, line: int | None = None):
    ppr = get_ppr(p_elem)
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = insert_ordered(ppr, OxmlElement("w:spacing"))
    if before is not None:
        sp.set(qn("w:before"), str(before))
    if after is not None:
        sp.set(qn("w:after"), str(after))
    if line is not None:
        sp.set(qn("w:line"), str(line))
        sp.set(qn("w:lineRule"), "auto")


def keep_with_next(p_elem):
    set_child(get_ppr(p_elem), "w:keepNext")


# ── runs ──────────────────────────────────────────────────────────────────────

def make_run(text: str, base_rpr=None, *, bold=False, italic=False, strike=False,
             code=False, size_pt: float | None = None, color: str | None = None,
             underline=False, font: str | None = None):
    r = OxmlElement("w:r")
    rpr = copy.deepcopy(base_rpr) if base_rpr is not None else OxmlElement("w:rPr")
    if code or font:
        for old in rpr.findall(qn("w:rFonts")):
            rpr.remove(old)
        rf = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), font or CODE_FONT)
        insert_ordered(rpr, rf)
    if bold:
        set_child(rpr, "w:b")
    if italic:
        set_child(rpr, "w:i")
    if strike:
        set_child(rpr, "w:strike")
    if underline:
        set_child(rpr, "w:u", val="single")
    if color:
        set_child(rpr, "w:color", val=color)
    if size_pt is not None:
        set_child(rpr, "w:sz", val=int(size_pt * 2))
        set_child(rpr, "w:szCs", val=int(size_pt * 2))
    if len(rpr):
        r.append(rpr)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i:
            r.append(OxmlElement("w:br"))
        if line:
            t = OxmlElement("w:t")
            t.text = line
            if line != line.strip() or "  " in line:
                t.set(XML_SPACE, "preserve")
            r.append(t)
    return r


def add_spans(doc, p_elem, spans: list[Span], base_rpr=None, part=None, **extra):
    """Runs по списку Span; ссылки — настоящие w:hyperlink (part нужен для rId)."""
    for s in spans:
        if getattr(s, "math", False):
            from .omml import omml_element
            p_elem.append(omml_element(s.text))
            continue
        if getattr(s, "ref", None):
            p_elem.append(ref_field(s.ref, s.text))
            continue
        if s.link and part is not None:
            rid = part.relate_to(s.link, RT.HYPERLINK, is_external=True)
            h = OxmlElement("w:hyperlink")
            h.set(qn("r:id"), rid)
            styled = has_style(doc, "Hyperlink")
            r = make_run(s.text, base_rpr, bold=s.bold or extra.get("bold", False),
                         italic=s.italic or extra.get("italic", False), strike=s.strike, code=s.code,
                         underline=not styled, color=None if styled else "0563C1")
            if styled:
                rpr = r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr")
                    r.insert(0, rpr)
                rs = OxmlElement("w:rStyle")
                rs.set(qn("w:val"), style_ids(doc)["Hyperlink"])
                insert_ordered(rpr, rs)
            h.append(r)
            p_elem.append(h)
        else:
            p_elem.append(make_run(s.text, base_rpr, bold=s.bold or extra.get("bold", False),
                                   italic=s.italic or extra.get("italic", False),
                                   strike=s.strike, code=s.code, size_pt=extra.get("size_pt")))


# ── списки ────────────────────────────────────────────────────────────────────

class Numbering:
    """Создаёт новый num для каждого списка — нумерация начинается заново.
    Часть numbering.xml создаётся лениво (в шаблоне её может не быть); если создать
    не удалось — new_list возвращает None, и список рисуется маркером-текстом."""

    def __init__(self, doc):
        self.doc = doc
        self._numbering = None
        self._failed = False

    def _element(self):
        if self._numbering is not None or self._failed:
            return self._numbering
        self._numbering = numbering_element(self.doc)
        self._failed = self._numbering is None
        return self._numbering

    def _next_id(self, tag: str, attr: str) -> int:
        ids = [int(e.get(qn(attr))) for e in self._numbering.findall(qn(tag))]
        return (max(ids) + 1) if ids else 1

    _abs_cache: dict

    def new_list(self, ordered: bool, start: int = 1) -> int | None:
        """Новый w:num (нумерация с start); abstractNum переиспользуется по виду списка."""
        if self._element() is None:
            return None
        cache = self.__dict__.setdefault("_abs_cache", {})
        abs_id = cache.get(ordered)
        if abs_id is None:
            abs_id = cache[ordered] = self._new_abstract(ordered)
        num_id = self._next_id("w:num", "w:numId")
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        set_child(num, "w:abstractNumId", val=abs_id)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), "0")
        set_child(ov, "w:startOverride", val=start)
        num.append(ov)
        self._numbering.append(num)
        return num_id

    def new_heading_numbering(self, style_ids: list[str], fmt: str = "%1.") -> int | None:
        """Многоуровневая нумерация заголовков («1.», «1.1.», …), привязанная к стилям.

        Нумерует Word, а не мы: вписать «1.1» текстом в заголовок значит получить отчёт,
        в котором вставка раздела перенумеровывает всё вручную. Привязка двусторонняя —
        `w:pStyle` в уровне и `w:numPr` в самом стиле: Word довольствуется первой,
        LibreOffice без второй нумерацию не показывает.
        """
        if self._element() is None or not style_ids:
            return None
        abs_id = self._next_id("w:abstractNum", "w:abstractNumId")
        absn = OxmlElement("w:abstractNum")
        absn.set(qn("w:abstractNumId"), str(abs_id))
        set_child(absn, "w:multiLevelType", val="multilevel")
        for lvl, sid in enumerate(style_ids):
            l = OxmlElement("w:lvl")
            l.set(qn("w:ilvl"), str(lvl))
            set_child(l, "w:start", val=1)
            set_child(l, "w:numFmt", val="decimal")
            set_child(l, "w:pStyle", val=sid)
            set_child(l, "w:lvlText", val=".".join(f"%{i + 1}" for i in range(lvl + 1)) + fmt[-1])
            set_child(l, "w:lvlJc", val="left")
            ppr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), "0")
            ind.set(qn("w:hanging"), "0")
            ppr.append(ind)
            insert_ordered(l, ppr)
            absn.append(l)
        nums = self._numbering.findall(qn("w:num"))
        if nums:
            nums[0].addprevious(absn)
        else:
            self._numbering.append(absn)
        num_id = self._next_id("w:num", "w:numId")
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        set_child(num, "w:abstractNumId", val=abs_id)
        self._numbering.append(num)
        return num_id

    def _new_abstract(self, ordered: bool) -> int:
        abs_id = self._next_id("w:abstractNum", "w:abstractNumId")
        absn = OxmlElement("w:abstractNum")
        absn.set(qn("w:abstractNumId"), str(abs_id))
        set_child(absn, "w:multiLevelType", val="hybridMultilevel")
        for lvl in range(4):
            l = OxmlElement("w:lvl")
            l.set(qn("w:ilvl"), str(lvl))
            set_child(l, "w:start", val=1)
            if ordered:
                fmt = ["decimal", "lowerLetter", "lowerRoman", "decimal"][lvl]
                set_child(l, "w:numFmt", val=fmt)
                set_child(l, "w:lvlText", val=f"%{lvl + 1}.")
            else:
                set_child(l, "w:numFmt", val="bullet")
                set_child(l, "w:lvlText", val=["•", "–", "▪", "•"][lvl])
            set_child(l, "w:lvlJc", val="left")
            ppr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(720 + 360 * lvl))
            ind.set(qn("w:hanging"), "360")
            ppr.append(ind)
            insert_ordered(l, ppr)
            absn.append(l)
        # abstractNum должны идти перед num
        nums = self._numbering.findall(qn("w:num"))
        if nums:
            nums[0].addprevious(absn)
        else:
            self._numbering.append(absn)
        return abs_id


def numbering_element(doc):
    """Корень numbering.xml; части в шаблоне может не быть — тогда создаём.
    None — не вышло: списки рисуются маркером-текстом, нумерация заголовков не ставится."""
    try:
        return doc.part.numbering_part.element
    except NotImplementedError:
        try:
            return _create_numbering_part(doc).element
        except Exception:                                      # noqa: BLE001
            return None


def _create_numbering_part(doc):
    from docx.opc.constants import CONTENT_TYPE as CT
    from docx.opc.packuri import PackURI
    from docx.oxml import parse_xml
    from docx.parts.numbering import NumberingPart
    xml = ('<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>')
    part = NumberingPart(PackURI("/word/numbering.xml"), CT.WML_NUMBERING, parse_xml(xml), doc.part.package)
    doc.part.relate_to(part, RT.NUMBERING)
    return part


def _ind_left(ind) -> int:
    try:
        return int(ind.get(qn("w:left")))
    except (AttributeError, TypeError, ValueError):
        return 0



def _num_ind_dxa(doc, numpr) -> int:
    """Отступ уровня списка из numbering.xml по w:numPr."""
    num_id = numpr.find(qn("w:numId"))
    if num_id is None:
        return 0
    ilvl = numpr.find(qn("w:ilvl"))
    lvl = ilvl.get(qn("w:val")) if ilvl is not None else "0"
    try:
        numbering = doc.part.numbering_part.element
    except (NotImplementedError, KeyError):
        return 0
    num = next((n for n in numbering.findall(qn("w:num"))
                if n.get(qn("w:numId")) == num_id.get(qn("w:val"))), None)
    abs_id = num.find(qn("w:abstractNumId")) if num is not None else None
    if abs_id is None:
        return 0
    absn = next((a for a in numbering.findall(qn("w:abstractNum"))
                 if a.get(qn("w:abstractNumId")) == abs_id.get(qn("w:val"))), None)
    node = next((l for l in absn.findall(qn("w:lvl")) if l.get(qn("w:ilvl")) == lvl), None) \
        if absn is not None else None
    return _ind_left(node.find(qn("w:pPr") + "/" + qn("w:ind"))) if node is not None else 0


def left_indent_dxa(doc, p_elem) -> int:
    """Отступ абзаца слева в dxa: прямой w:ind, иначе уровень списка (w:numPr → numbering.xml),
    иначе стиль абзаца. Word держит отступ пункта списка в трёх разных местах, а таблица
    и её подпись, вставленные на место тега, иначе убегают к левому полю из-под своего пункта."""
    ppr = p_elem.find(qn("w:pPr"))
    if ppr is None:
        return 0
    for src in (ppr, _style_ppr(doc, ppr)):
        if src is None:
            continue
        left = _ind_left(src.find(qn("w:ind")))
        if left:
            return left
        numpr = src.find(qn("w:numPr"))
        if numpr is not None:
            left = _num_ind_dxa(doc, numpr)
            if left:
                return left
    return 0


def _style_ppr(doc, ppr):
    """w:pPr стиля абзаца (без basedOn — глубже Word и сам редко ходит за отступом)."""
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        return None
    sid = pstyle.get(qn("w:val"))
    el = next((s for s in doc.styles.element.findall(qn("w:style")) if s.get(qn("w:styleId")) == sid), None)
    return el.find(qn("w:pPr")) if el is not None else None



def set_list_item(doc, p_elem, num_id: int | None, level: int, ordered: bool = False, index: int = 1):
    """Пункт списка. num_id=None — запасной вариант без numbering.xml: маркер текстом + отступ."""
    set_style(doc, p_elem, "List Paragraph")
    ppr = get_ppr(p_elem)
    if num_id is None:
        set_child(ppr, "w:ind", left=720 + 360 * level, hanging=360)
        marker = f"{index}." if ordered else ["•", "–", "▪", "•"][level]
        p_elem.append(make_run(marker + "\t", None))
        return
    for old in ppr.findall(qn("w:numPr")):
        ppr.remove(old)
    numpr = insert_ordered(ppr, OxmlElement("w:numPr"))
    set_child(numpr, "w:ilvl", val=level)
    set_child(numpr, "w:numId", val=num_id)
    for old in ppr.findall(qn("w:ind")):
        ppr.remove(old)


# ── специальные абзацы ────────────────────────────────────────────────────────

def style_heading(doc, p_elem, level: int, base_rpr=None):
    if set_style(doc, p_elem, f"Heading {level}"):
        return {}
    set_spacing(p_elem, before=240, after=120)
    keep_with_next(p_elem)
    set_child(get_ppr(p_elem), "w:outlineLvl", val=level - 1)
    return {"bold": True, "size_pt": HEADING_SIZES[level]}


def style_quote(doc, p_elem, level: int = 0):
    """Цитата; level — вложенность «> >»: каждый уровень добавляет отступ к отступу стиля."""
    if set_style(doc, p_elem, "Quote"):
        if level:
            base = doc.styles["Quote"].paragraph_format.left_indent
            set_child(get_ppr(p_elem), "w:ind", left=int(base.twips if base else 0) + 720 * level)
        return
    ppr = get_ppr(p_elem)
    set_child(ppr, "w:ind", left=720 * (level + 1))
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", "12"), ("space", "8"), ("color", "BBBBBB")):
        left.set(qn(f"w:{k}"), v)
    bdr.append(left)
    insert_ordered(ppr, bdr)


def style_hr(p_elem):
    ppr = get_ppr(p_elem)
    bdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "888888")):
        bot.set(qn(f"w:{k}"), v)
    bdr.append(bot)
    insert_ordered(ppr, bdr)


def _highlight_lines(text: str, lang: str, style_name: str):
    """[[(chunk, color_hex|None, bold, italic), …] по строкам] через pygments; None — без подсветки."""
    try:
        from pygments import lex
        from pygments.lexers import get_lexer_by_name
        from pygments.styles import get_style_by_name
    except ImportError:
        return None
    try:
        lexer = get_lexer_by_name(lang)
        style = get_style_by_name(style_name)
    except Exception:                                      # noqa: BLE001 — неизвестный язык/стиль
        return None
    lines, cur = [], []
    for tok, val in lex(text, lexer):
        st = style.style_for_token(tok)
        color = st["color"] or None
        parts = val.split("\n")
        for i, part in enumerate(parts):
            if i:
                lines.append(cur)
                cur = []
            if part:
                cur.append((part, color, st["bold"], st["italic"]))
    lines.append(cur)
    if lines and not lines[-1]:
        lines.pop()
    return lines


def add_code_lines(doc, ref_elem, text: str, ppr_template=None, *, lang: str = "",
                   font: str = CODE_FONT, size_pt: float = CODE_SIZE_PT,
                   highlight: bool = False, line_numbers: bool = False, style_name: str = "default"):
    """Листинг: по абзацу на строку (переносы страниц между строками), моноширинный шрифт,
    опционально подсветка pygments и номера строк."""
    lines = text.split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    if not lines:
        lines = [""]
    colored = _highlight_lines(text, lang, style_name) if (highlight and lang) else None
    if colored is not None and len(colored) != len(lines):
        colored = None
    width = len(str(len(lines)))
    last = ref_elem
    for i, line in enumerate(lines):
        p = new_paragraph_after(last, ppr_template)
        if not set_style(doc, p, "Code"):
            set_spacing(p, before=0, after=0)
            set_child(get_ppr(p), "w:jc", val="left")
            for old in get_ppr(p).findall(qn("w:ind")):
                get_ppr(p).remove(old)
        if line_numbers:
            p.append(make_run(f"{i + 1:>{width}}  ", None, font=font, size_pt=size_pt, color="808080"))
        if colored is None:
            p.append(make_run(line, None, font=font, size_pt=size_pt))
        else:
            for chunk, color, bold, italic in colored[i]:
                p.append(make_run(chunk, None, font=font, size_pt=size_pt, color=color, bold=bold, italic=italic))
        last = p
    return last


def add_page_break(ref_elem):
    p = new_paragraph_after(ref_elem)
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    return p


def add_picture(doc, ref_elem, parent_proxy, data: bytes, w_cm: float, h_cm: float, align: str = "center"):
    """Абзац с картинкой после ref_elem. parent_proxy — объект с .part (абзац тега)."""
    p = new_paragraph_after(ref_elem)
    set_alignment(p, align)
    keep_with_next(p)
    para = Paragraph(p, parent_proxy)
    run = para.add_run()
    run.add_picture(io.BytesIO(data), width=Emu(int(w_cm * EMU_PER_CM)), height=Emu(int(h_cm * EMU_PER_CM)))
    return p


def _next_bookmark_id(doc) -> int:
    """Номер закладки свой у каждого документа: общий на процесс счётчик делал вывод
    невоспроизводимым (два рендера подряд давали разные w:id) и мог разорвать пару
    bookmarkStart/bookmarkEnd, если два рендера идут одновременно."""
    n = getattr(doc, "_hokoku_bookmark_id", 100)
    try:
        doc._hokoku_bookmark_id = n + 1
    except AttributeError:
        pass
    return n


def _field(instr: str, cached: str, **run_kw):
    """Простое поле Word с кэшированным значением (Word обновит при F9/печати).
    run_kw — оформление кэша: номер в жирном заголовке оглавления должен быть жирным."""
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), instr)
    f.append(make_run(cached, None, **run_kw))
    return f


def page_number_field(cached: str = "1"):
    """Поле PAGE — номер страницы. Текстом номер писать нельзя: он не пересчитается."""
    return _field(" PAGE ", cached)


def _plain_with_refs(p, text: str, **run_kw):
    """Простой текст в абзац: «ср. {ref:схема}» → поле REF вместо литерала. Разметку
    здесь не разбираем (см. `markdown.split_refs`), run_kw — оформление обычных кусков."""
    for s in split_refs(text):
        p.append(ref_field(s.ref, s.text, **run_kw) if s.ref else make_run(s.text, None, **run_kw))


def add_caption(doc, ref_elem, fmt: str, n: int, caption: str | None, *, align: str = "center",
                ppr_template=None, seq_name: str | None = None, bookmark: str | None = None,
                suffix: str = "", repeat: bool = False, indent_dxa: int = 0):
    r"""
    Подпись «Рисунок N — текст». seq_name → номер полем SEQ (repeat=True: `\c` — тот же
    номер, для листов одной картинки); bookmark → закладка вокруг номера для ссылок REF.
    indent_dxa — отступ слева: подпись стоит над своей таблицей (тег в пункте списка).
    """
    p = new_paragraph_after(ref_elem, ppr_template)
    if not set_style(doc, p, "Caption"):
        set_spacing(p, before=60, after=200)
    set_alignment(p, align)
    if indent_dxa:
        set_child(get_ppr(p), "w:ind", left=indent_dxa)
    if not caption:
        fmt = re.sub(r"\s*[—–-]\s*\{caption\}", "", fmt)
    before, _, after = fmt.partition("{n}")
    after = after.replace("{caption}", caption or "") + suffix
    if before:
        _plain_with_refs(p, before)
    bm_id = _next_bookmark_id(doc) if bookmark else 0
    if bookmark:
        bs = OxmlElement("w:bookmarkStart")
        bs.set(qn("w:id"), str(bm_id))
        bs.set(qn("w:name"), bookmark)
        p.append(bs)
    if seq_name:
        p.append(_field(f" SEQ {seq_name} \\* ARABIC" + (" \\c" if repeat else "") + " ", str(n)))
    else:
        p.append(make_run(str(n), None))
    if bookmark:
        be = OxmlElement("w:bookmarkEnd")
        be.set(qn("w:id"), str(bm_id))
        p.append(be)
    if after:
        _plain_with_refs(p, after)
    return p


def add_formula(doc, ref_elem, latex: str, *, numbered: bool, n: int, seq_name: str | None,
                bookmark: str | None, page_w_cm: float, ppr_template=None):
    """Абзац с формулой по центру; при numbered — «(n)» у правого поля через табуляцию."""
    from .omml import omml_element
    p = new_paragraph_after(ref_elem, ppr_template)
    ppr = get_ppr(p)
    for old in ppr.findall(qn("w:ind")):
        ppr.remove(old)
    set_spacing(p, before=120, after=120)
    if numbered:
        # табуляции: центр — для формулы, правый край — для номера
        tabs = OxmlElement("w:tabs")
        for kind, pos in (("center", page_w_cm / 2), ("right", page_w_cm)):
            t = OxmlElement("w:tab")
            t.set(qn("w:val"), kind)
            t.set(qn("w:pos"), str(int(pos * DXA_PER_CM)))
            tabs.append(t)
        insert_ordered(ppr, tabs)
        set_alignment(p, "left")
        p.append(make_run("\t", None))
        p.append(omml_element(latex))
        p.append(make_run("\t(", None))
        bm_id = _next_bookmark_id(doc) if bookmark else 0
        if bookmark:
            bs = OxmlElement("w:bookmarkStart")
            bs.set(qn("w:id"), str(bm_id))
            bs.set(qn("w:name"), bookmark)
            p.append(bs)
        p.append(_field(f" SEQ {seq_name} \\* ARABIC ", str(n)) if seq_name else make_run(str(n), None))
        if bookmark:
            be = OxmlElement("w:bookmarkEnd")
            be.set(qn("w:id"), str(bm_id))
            p.append(be)
        p.append(make_run(")", None))
    else:
        set_alignment(p, "center")
        p.append(omml_element(latex))
    return p


def add_toc(doc, ref_elem, levels: int, title: str | None, ppr_template=None):
    """Поле TOC; содержимое заполнит Word при открытии (settings: updateFields)."""
    last = ref_elem
    if title:
        p = new_paragraph_after(last, None)
        if not set_style(doc, p, "TOC Heading"):
            style_heading(doc, p, 1)
        # заголовок оглавления — отдельный абзац перед полем TOC, а не его содержимое:
        # поле REF в нём законно и переживает обновление оглавления
        _plain_with_refs(p, title, bold=True)
        last = p
    p = new_paragraph_after(last, ppr_template)
    r = OxmlElement("w:r")
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    fc.set(qn("w:dirty"), "true")
    r.append(fc)
    p.append(r)
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(XML_SPACE, "preserve")
    it.text = f' TOC \\o "1-{levels}" \\h \\z \\u '
    r.append(it)
    p.append(r)
    r = OxmlElement("w:r")
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    r.append(fc)
    p.append(r)
    # результат поля оставляем пустым: LibreOffice поле TOC не разворачивает и печатает
    # кэш как обычный текст — служебная подсказка уезжала в сданный PDF
    r = OxmlElement("w:r")
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r.append(fc)
    p.append(r)
    set_update_fields(doc)
    return p


def set_update_fields(doc):
    """settings.xml: w:updateFields — Word предложит обновить поля при открытии."""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)


def ref_field(bookmark: str, cached: str, **run_kw):
    """Поле REF на закладку подписи: `{ref:имя}` в тексте → номер рисунка/таблицы."""
    return _field(f" REF {bookmark} \\h ", cached, **run_kw)


def page_text_width_cm(doc) -> float:
    s = doc.sections[-1]
    return max(5.0, (s.page_width - s.left_margin - s.right_margin) / EMU_PER_CM)


def page_text_height_cm(doc) -> float:
    s = doc.sections[-1]
    return max(5.0, (s.page_height - s.top_margin - s.bottom_margin) / EMU_PER_CM)


# ── таблицы ───────────────────────────────────────────────────────────────────

DXA_PER_CM = 567
HEADER_SHADE = "E7E6E6"


def _doc_font_pt(doc) -> float:
    pt = getattr(doc, "_hokoku_font_pt", None)
    if pt is None:
        pt = 12.0
        try:
            sz = doc.styles["Normal"].font.size
            if sz:
                pt = sz.pt
        except KeyError:
            pass
        try:
            doc._hokoku_font_pt = pt
        except AttributeError:
            pass
    return pt


def _auto_col_widths(rows, ncols: int, max_w_cm: float, font_pt: float = 12.0,
                     min_col_cm: float = 1.2) -> list[float]:
    """Ширины колонок по самому длинному содержимому (+ отступы), не шире max_w.
    Если в полосу не влезает, ужимается только избыток над полом (самое длинное слово —
    переносить его всё равно некуда): пропорциональное ужатие продавливало короткие
    колонки, хотя рвался в итоге длинный текст, которому перенос не страшен."""
    font_px = font_pt * 96 / 72
    space = text_width(" ", font_px, bold=True)

    def cm(px: float) -> float:
        return px * 1.2 / 37.8 + 0.8                       # px→cm при 96 dpi + поля ячейки

    need, floor = [], []
    for ci in range(ncols):
        w = fl = min_col_cm
        for row in rows:
            if ci < len(row):
                # ширина всей ячейки складывается из ширин слов — заодно самое длинное слово
                words = [text_width(x, font_px, bold=True)
                         for x in "".join(sp.text for sp in row[ci]).split()]
                if words:
                    w = max(w, cm(sum(words) + space * (len(words) - 1)))
                    fl = max(fl, cm(max(words)))
        need.append(min(w, max_w_cm))
        floor.append(min(fl, need[-1]))
    total = sum(need)
    if total > max_w_cm:
        room, surplus = max_w_cm - sum(floor), total - sum(floor)
        if room > 0 and surplus > 0:
            need = [f + (w - f) * room / surplus for w, f in zip(need, floor)]
        else:                                 # одни только неразрывные слова — ужать пропорционально
            need = [w * max_w_cm / total for w in need]
    return [round(w, 2) for w in need]


def add_table(doc, ref_elem, rows: list[list[list[Span]]], header: bool, part,
              align: list[str] | None = None, base_rpr=None, col_widths_cm: list[float] | None = None,
              max_w_cm: float | None = None, header_fill: str = HEADER_SHADE, header_center: bool = True,
              min_col_cm: float = 1.2, indent_dxa: int = 0):
    """Таблица после ref_elem: стиль Table Grid если есть, иначе рамки вручную."""
    ncols = max(len(r) for r in rows) if rows else 1
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl_elem = tbl._tbl
    if not col_widths_cm:
        col_widths_cm = _auto_col_widths(rows, ncols, max_w_cm or 16.0, _doc_font_pt(doc), min_col_cm)
    cols = [float(w) for w in col_widths_cm[:ncols]]
    if len(cols) < ncols:                     # ширин задали меньше, чем колонок — добрать остатком
        rest = max(0.0, (max_w_cm or sum(cols)) - sum(cols))
        cols += [max(min_col_cm, rest / (ncols - len(cols)))] * (ncols - len(cols))
    tbl.autofit = False
    set_child(tbl_elem.tblPr, "w:tblW", w=int(sum(cols) * DXA_PER_CM), type="dxa")
    if indent_dxa:
        set_child(tbl_elem.tblPr, "w:tblInd", w=indent_dxa, type="dxa")
    # LibreOffice раскладывает колонки по w:tblGrid, а не по w:tcW: без этой строки посчитанные
    # ширины видны в Word, но в PDF колонки выходят равными (сетку python-docx делит поровну)
    grid = tbl_elem.find(qn("w:tblGrid"))
    if grid is not None:
        for gc, wcm in zip(grid.findall(qn("w:gridCol")), cols):
            gc.set(qn("w:w"), str(int(wcm * DXA_PER_CM)))
    # tbl.cell(r, c) и row.cells каждый раз заново строят всю матрицу ячеек (учёт gridSpan
    # и vMerge), поэтому наполнение таблицы росло квадратично: 400 строк — 42 с, 98 % времени
    # внутри python-docx. Матрица нужна одна на таблицу; объединённых ячеек мы не делаем,
    # поэтому строгий порядок «строка за строкой» верен.
    cells = tbl._cells
    for ci, wcm in enumerate(cols):
        w = Emu(int(wcm * EMU_PER_CM))
        for ri in range(len(rows)):
            cells[ri * ncols + ci].width = w
    grid_id = style_ids(doc).get("Table Grid")
    if grid_id:
        # напрямую в w:tblStyle: сеттер python-docx на каждую таблицу заново ищет стиль
        # по имени и «стиль по умолчанию» XPath'ом — это была пятая часть времени рендера
        set_child(tbl_elem.tblPr, "w:tblStyle", val=grid_id)
    else:
        tblpr = tbl_elem.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "000000")):
                b.set(qn(f"w:{k}"), v)
            borders.append(b)
        insert_ordered(tblpr, borders)
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = cells[ri * ncols + ci]
            p = cell.paragraphs[0]._p
            spans = row[ci] if ci < len(row) else []
            if align and ci < len(align):
                set_alignment(p, align[ci])
            elif header and ri == 0 and header_center:
                set_alignment(p, "center")
            set_spacing(p, before=0, after=0)
            if header and ri == 0 and header_fill:
                tcpr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), header_fill)
                tcpr.append(shd)
            add_spans(doc, p, spans, base_rpr, part, bold=(header and ri == 0))
    # переносим из конца документа на место
    ref_elem.addnext(tbl_elem)
    return tbl_elem
