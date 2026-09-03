"""
Фикстуры для materials: всё собирается на месте, сеть и внешние файлы не нужны.
PDF строится через pymupdf, DOCX — через python-docx, картинки — через PIL.
"""
import io
import os
import stat
import zipfile

import docx
import pymupdf
import pytest
from docx.oxml import parse_xml
from docx.shared import Inches
from PIL import Image as PIL

from materials import Store


@pytest.fixture
def png():
    """Фабрика картинок: png(ширина, высота, цвет) → байты."""
    def _mk(w=160, h=90, color="orange"):
        buf = io.BytesIO()
        PIL.new("RGB", (w, h), color).save(buf, format="PNG")
        return buf.getvalue()
    return _mk


@pytest.fixture
def store(tmp_path):
    return Store(str(tmp_path / "материалы"))


def build_pdf(pages, title="", images=(), table_page=None):
    """
    Собрать PDF: pages — текст страниц, images — [(номер страницы, байты png)],
    table_page — номер страницы, куда нарисовать настоящую таблицу линиями.
    """
    doc = pymupdf.open()
    if title:
        doc.set_metadata({"title": title})
    for number, text in enumerate(pages, start=1):
        page = doc.new_page()
        page.insert_text((72, 100), text, fontsize=14)
        for where, data in images:
            if where == number:
                page.insert_image(pymupdf.Rect(72, 150, 232, 240), stream=data)
        if table_page == number:
            cols, rows = [70, 200, 330], [300, 330, 360]
            for x in cols:
                page.draw_line((x, rows[0]), (x, rows[-1]))
            for y in rows:
                page.draw_line((cols[0], y), (cols[-1], y))
            for ri, row in enumerate([["God", "Summa"], ["2025", "250"]]):
                for ci, val in enumerate(row):
                    page.insert_text((cols[ci] + 6, rows[ri] + 20), val, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture
def pdf(png):
    """Обычный PDF: три страницы текста, картинка на второй, таблица на третьей."""
    return build_pdf(["Vvedenie v temu raboty", "Hod raboty i izmereniya", "Rezultaty i vyvody"],
                     title="Методичка по химии", images=[(2, png())], table_page=3)


@pytest.fixture
def scanned_pdf(png):
    """PDF без текстового слоя: одна страница-картинка, букв в разметке нет."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(50, 50, 450, 350), stream=png(400, 300, "white"))
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture
def no_tesseract(monkeypatch):
    """Система, где tesseract не установлен: пустой PATH — which ничего не найдёт."""
    monkeypatch.setenv("PATH", "")
    return None


@pytest.fixture
def fake_tesseract(tmp_path, monkeypatch):
    """
    Подставной tesseract на PATH: проверяем именно проводку (аргументы, чтение
    stdout), а не качество распознавания. Возвращает текст, который «распознал».
    """
    recognized = "Раствор нагрели до 80 °C\nвторая строка"
    exe = tmp_path / "bin" / "tesseract"
    exe.parent.mkdir(parents=True, exist_ok=True)
    exe.write_text(
        "#!/bin/sh\n"
        'if [ "$1" = "--list-langs" ]; then\n'
        '  echo "List of available languages (2):"\n'
        "  echo rus\n  echo eng\n  exit 0\nfi\n"
        # Ожидаем вызов вида: tesseract <файл> stdout -l rus+eng
        '[ "$2" = "stdout" ] || exit 2\n'
        '[ "$3" = "-l" ] || exit 3\n'
        f'printf "%s\\n" "{recognized.splitlines()[0]}" "{recognized.splitlines()[1]}"\n',
        encoding="utf-8")
    os.chmod(exe, os.stat(exe).st_mode | stat.S_IEXEC | stat.S_IXGRP | stat.S_IXOTH)
    monkeypatch.setenv("PATH", str(exe.parent))
    return recognized


# ── Word ──────────────────────────────────────────────────────────────────────
# Настоящий .docx собирается здесь же через python-docx: файл-образец в репозитории
# нельзя ни прочитать глазами, ни поправить, а разбор Word проверять надо по
# настоящему пакету, а не по подделке из пары XML-строк.

MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def build_docx(paragraphs=(), title="", images=(), table=None, formula=None):
    """
    Собрать .docx: paragraphs — абзацы подряд, images — [(номер абзаца, байты png)],
    table — [[строка], …] после абзацев, formula — (номер абзаца, OMML-строка).
    OMML пишем руками: hokoku его умеет строить, но импортировать чужой пакет
    в materials нельзя ни в коде, ни в тестах.
    """
    doc = docx.Document()
    if title:
        doc.core_properties.title = title
    for number, text in enumerate(paragraphs, start=1):
        p = doc.add_paragraph(text)
        if formula and formula[0] == number:
            p._p.append(parse_xml(formula[1]))
        for where, data in images:
            if where == number:
                doc.add_picture(io.BytesIO(data), width=Inches(2))
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for ri, row in enumerate(table):
            for ci, val in enumerate(row):
                t.cell(ri, ci).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def fraction_omml(num="a+b", den="2"):
    """Дробь в родной записи Word — та самая, что python-docx не видит."""
    return (f'<m:oMath xmlns:m="{MATH_NS}"><m:f>'
            f"<m:num><m:r><m:t>{num}</m:t></m:r></m:num>"
            f"<m:den><m:r><m:t>{den}</m:t></m:r></m:den></m:f></m:oMath>")


def repack(data, add=None, drop=()):
    """Тот же zip, но с добавленными/выброшенными членами: так делают .docm и
    порченые файлы, которых честным путём не собрать."""
    add = add or {}
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            if info.filename in drop or info.filename in add:
                continue                    # заменяемый член пишем один раз, ниже
            dst.writestr(info.filename, src.read(info.filename))
        for name, blob in add.items():
            dst.writestr(name, blob)
    return out.getvalue()


@pytest.fixture
def docx_file(png):
    """Обычная методичка: три абзаца, формула во втором, картинка после второго,
    таблица в конце."""
    return build_docx(["Vvedenie v temu raboty", "Formula: ", "Vyvody po rabote"],
                      title="Методичка по химии", images=[(2, png())],
                      table=[["God", "Summa"], ["2025", "250"]],
                      formula=(2, fraction_omml()))
