"""
Прогон целиком: значения → `hokoku.validate` → `hokoku.build_report`.

Второго способа собрать отчёт не заводится и второго валидатора тоже: тесты
здесь стерегут именно это — что в `out/` лёг файл, собранный тем же
`build_report`, что и сборка из командной строки, и что замечания посчитаны тем
же `validate`, который проверял значения модели по одному.
"""
from __future__ import annotations

import json
import os

from docx import Document

import hokoku
import orchestrator

from .conftest import markdown_value, report_json, script

ЦЕЛЬ = {"type": "markdown", "text": "Цель — сравнить алгоритмы."}
ВВЕДЕНИЕ = {"type": "markdown", "text": "Постановка задачи."}
ТАБЛИЦА = {"type": "table", "rows": [["алгоритм", "время"], ["быстрая", "0,3 с"]]}


def test_задание_собирается_из_состояния_проекта(project):
    project.set_value("цель", ЦЕЛЬ, source="manual")
    job = orchestrator.job_of(project, outputs=["docx"])
    assert job["template"]["artifact"] == project.template_artifact()
    # sha256 шаблона кладём всегда: собрать по подменённому шаблону — значит
    # выдать не тот документ за тот, и заметят это, только открыв файл.
    assert job["template"]["sha256"] == project.manifest().template_sha256
    assert job["values"]["цель"]["text"] == ЦЕЛЬ["text"]
    # Номер контракта значений сверяется с `hokoku`, а не вписан числом: контракт
    # поднимают в `wire`, и тест, знающий число наизусть, падает на чужой правке,
    # ничего про службу не проверив.
    assert job["wire_version"] == hokoku.WIRE_VERSION
    assert job["template"]["manifest_version"] == 1


def test_номер_манифеста_доезжает_до_задания_сборки(project):
    """`template.manifest_version` уходит в журнал сборки и отвечает на вопрос
    «по какому манифесту собран этот файл». Пока счётчик стоял на единице, он
    отвечал на него неверно — и по журналу это было не видно."""
    m = project.manifest()
    m.tags["цель"].prompt = "Сформулируй цель работы"
    project.save_manifest(m)
    project.set_value("цель", ЦЕЛЬ, source="manual")
    assert orchestrator.job_of(project)["template"]["manifest_version"] == 2


def test_сборка_кладёт_документ_и_считает_замечания(project):
    project.set_value("цель", ЦЕЛЬ, source="manual")
    project.set_value("введение", ВВЕДЕНИЕ, source="manual")
    project.set_value("таблица", ТАБЛИЦА, source="manual")

    итог = orchestrator.build(project)

    assert итог["ok"] is True
    имя = итог["report"]["outputs"]["docx"]["file"]
    путь = os.path.join(итог["workdir"], имя)
    assert os.path.isfile(путь)
    тексты = [p.text for p in Document(путь).paragraphs]
    assert any("Цель — сравнить алгоритмы." in t for t in тексты)
    assert итог["report"]["unfilled"] == []
    assert [p for p in итог["problems"] if p.level == "error"] == []


def test_полная_проверка_видит_то_чего_не_видит_проверка_одного_тега(project):
    """Посреди прогона «обязательный тег без значения» — это «ещё не дошли»;
    перед сборкой — это ответ на вопрос «годен ли отчёт»."""
    project.set_value("цель", ЦЕЛЬ, source="manual")
    беды = orchestrator.check(project)
    коды = {(p.key, p.code) for p in беды}
    assert ("введение", "missing_required") in коды
    assert ("таблица", "missing_required") in коды
    assert ("цель", "missing_required") not in коды


def test_модель_заполнила_отчёт_и_он_собрался(project, endpoint):
    """Сквозной путь без сети: раскладка → поток → версии → документ."""
    ep, _ = endpoint(script(report_json(цель=ЦЕЛЬ, введение=ВВЕДЕНИЕ, таблица=ТАБЛИЦА),
                            pieces=25))
    прогон = orchestrator.fill_report(project, endpoint=ep)
    assert прогон.ok

    итог = orchestrator.build(project)
    assert итог["ok"] is True
    assert [p for p in итог["problems"] if p.level == "error"] == []
    путь = os.path.join(итог["workdir"], итог["report"]["outputs"]["docx"]["file"])
    тексты = "\n".join(p.text for p in Document(путь).paragraphs)
    assert "Постановка задачи." in тексты
    # Каждое значение документа объяснимо: кто поставил и в каком прогоне.
    for key in ("цель", "введение", "таблица"):
        версия = project.versions(key)[-1]
        assert (версия.source, версия.run) == ("agent", прогон.run.id)


def test_сборка_не_срывается_из_за_негодного_значения(project):
    """Отчёт с пометками полезнее человеку, чем отказ: он его открывает и правит."""
    project.set_value("цель", ЦЕЛЬ, source="manual")
    project.set_value("введение", markdown_value("см. {ref:нетакого}"), source="manual")
    project.set_value("таблица", ТАБЛИЦА, source="manual")
    итог = orchestrator.build(project)
    assert итог["ok"] is True
    assert any(p.code == "unresolved_ref" for p in итог["problems"])


def test_картинка_из_материалов_доезжает_до_документа(project, tmp_path):
    """Стык, который сошёлся сам: идентификатор материала проходит `ARTIFACT_RE`,
    а `Store.blob` — это ровно сигнатура `resolve_artifact`."""
    import io

    from PIL import Image as PIL
    buf = io.BytesIO()
    PIL.new("RGB", (60, 40), "lightblue").save(buf, format="PNG")
    material = project.store().add(buf.getvalue(), name="снимок.png", do_ocr=False)

    m = project.manifest()
    m.tags["таблица"].type = "image"
    project.save_manifest(m)
    project.set_value("цель", ЦЕЛЬ, source="manual")
    project.set_value("введение", ВВЕДЕНИЕ, source="manual")
    project.set_value("таблица", {"type": "image", "artifact": material.id,
                                  "caption": "Снимок"}, source="manual")

    итог = orchestrator.build(project)
    assert итог["ok"] is True
    assert итог["report"]["counts"]["figures"] == 1
    assert json.loads(json.dumps(итог["report"]))["template"]["sha256"]
