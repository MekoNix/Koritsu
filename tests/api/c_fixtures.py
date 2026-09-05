"""
Оснастка пространств: workspace, проекты, корзина.

Отдельным модулем, а не строками в `tests/api/conftest.py`: тот файл общий на
весь набор, и правка его со всех сторон — конфликты слияния на ровном месте.
Фикстуры отсюда обычные, их просто импортируют в свой тест.

**Вошедший — настоящий, подмена — только у зависимости.** Пользователь заводится
через живую регистрацию (`POST /api/auth/register`), потому что внешние ключи в
наших таблицах смотрят на `users.id` и подделка сломалась бы на первой же
вставке. Подменяется одна вещь — `current_user`: подтверждение почты и вход по
cookie проверяет `accounts` у себя, а тесту про пространства нужно только «кто
пришёл».
"""
from __future__ import annotations

import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from api.accounts import User, current_user

# Пароль длиннее предела `accounts` (`PASSWORD_MIN = 10`) и один на все тесты:
# что именно в нём написано, ни одна проверка здесь не спрашивает.
ПАРОЛЬ = "очень-длинный-пароль"


@pytest.fixture
def клиент(app):
    """Клиент службы. Без тестовых маршрутов: беды каркаса проверяет `test_errors`."""
    with TestClient(app, raise_server_exceptions=False) as c:
        yield c


def ник_для(email: str) -> str:
    """Ник по почте: `хозяин@пример.рф` → `хозяин`.

    Регистрация требует ник, а тестам он безразличен — важно, чтобы он был
    свой у каждого. Почта у каждого своя, значит и ник выйдет свой.
    """
    return (email.split("@")[0] or "кто")[:32]


def завести(app, клиент, email: str) -> User:
    """Живая регистрация → объект пользователя из базы.

    Хук личного пространства срабатывает здесь же, внутри регистрации: именно
    это и проверяет `test_workspaces.test_личное_заводится_при_регистрации`.
    """
    ответ = клиент.post("/api/auth/register",
                        json={"email": email, "password": ПАРОЛЬ,
                              "nickname": ник_для(email)})
    assert ответ.status_code == 201, ответ.text
    with app.state.db.session_scope() as s:
        user = s.scalar(select(User).where(User.email == email))
    assert user is not None
    return user


def войти(app, user: User) -> None:
    """Сделать этого пользователя вошедшим для всех следующих запросов.

    Подменяется объект функции `current_user`, а не строка cookie: подписанная
    cookie — дело `accounts`, и повторять её выпуск здесь значило бы завести
    второе описание того, что такое вход.
    """
    app.dependency_overrides[current_user] = lambda: user


def позвать(app, клиент, workspace_id: str, кого: User, role: str = "editor"):
    """Позвать человека в пространство и принять приглашение от его имени.

    Приглашение двухшаговое: `POST …/members` заводит участие в состоянии
    `pending`, и участником человек становится, только ответив «принять».
    Тестам, которым нужен не сам обмен приглашениями, а участник в
    пространстве, это два лишних шага и переключение вошедшего туда-обратно —
    вот они и собраны здесь.

    Возвращает ответ на приглашение (тесты проверяют по нему `201` и тело), а
    вошедшим оставляет того же, кто был до вызова.
    """
    ответ = клиент.post(f"/api/workspaces/{workspace_id}/members",
                        json={"email": кого.email, "role": role})
    if ответ.status_code == 201:
        был = app.dependency_overrides.get(current_user)
        войти(app, кого)
        принято = клиент.post(
            f"/api/workspaces/{workspace_id}/members/{кого.id}/accept")
        assert принято.status_code == 200, принято.text
        if был is None:
            app.dependency_overrides.pop(current_user, None)
        else:
            app.dependency_overrides[current_user] = был
    return ответ


@pytest.fixture
def хозяин(app, клиент):
    """Первый пользователь, уже вошедший. У него есть личное пространство."""
    user = завести(app, клиент, "хозяин@пример.рф")
    войти(app, user)
    return user


@pytest.fixture
def сосед(app, клиент, хозяин):
    """Второй пользователь. Заведён, но не вошедший: вход переключают `войти`."""
    return завести(app, клиент, "сосед@пример.рф")


def войти_по_настоящему(клиент, caplog, email: str = "живой@пример.рф",
                        пароль: str = ПАРОЛЬ):
    """Регистрация, подтверждение и вход — с настоящей cookie-сессией.

    Нужно ровно тем тестам, которые проверяют саму cookie: заслон CSRF
    срабатывает только на запросе с сессией, а `войти` выше подменяет
    зависимость и cookie не заводит вовсе. Токен подтверждения берётся из письма
    в журнале — писем служба не шлёт, `ConsoleMailer` пишет их в `api.mail`.
    """
    import re

    caplog.set_level("INFO")
    ответ = клиент.post("/api/auth/register",
                        json={"email": email, "password": пароль,
                              "nickname": ник_для(email)})
    assert ответ.status_code == 201, ответ.text
    письма = [з.getMessage() for з in caplog.records if з.name == "api.mail"]
    найден = re.search(r"token=([A-Za-z0-9_-]+)", письма[-1])
    assert найден, письма[-1]
    assert клиент.post("/api/auth/confirm",
                       json={"token": найден.group(1)}).status_code == 200
    вход = клиент.post("/api/auth/login",
                       json={"email": email, "password": пароль})
    assert вход.status_code == 200, вход.text
    return вход


def docx_байты(теги=("цель", "выводы")) -> bytes:
    """Настоящий DOCX с тегами `{{ключ}}`: манифест по нему строится честно."""
    doc = Document()
    doc.add_paragraph("Отчёт")
    for key in теги:
        doc.add_paragraph(f"{{{{{key}}}}}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def создать_проект(клиент, workspace_id: str, *, name: str = "проба",
                   шаблон: bytes | None = None) -> dict:
    """Создать проект через API. Без файла — «проект без шаблона»."""
    files = {"template": ("шаблон.docx", шаблон,
                          "application/vnd.openxmlformats-officedocument."
                          "wordprocessingml.document")} if шаблон else None
    ответ = клиент.post("/api/projects", data={"workspace_id": workspace_id,
                                               "name": name}, files=files)
    assert ответ.status_code == 201, ответ.text
    return ответ.json()


def личное_id(клиент) -> str:
    """Идентификатор личного пространства вошедшего."""
    ответ = клиент.get("/api/workspaces/personal")
    assert ответ.status_code == 200, ответ.text
    return ответ.json()["id"]


__all__ = ["клиент", "хозяин", "сосед", "завести", "войти", "позвать",
           "docx_байты",
           "создать_проект", "личное_id", "ПАРОЛЬ", "войти_по_настоящему"]
