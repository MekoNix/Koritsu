import io

import pytest
from docx import Document
from PIL import Image as PIL


@pytest.fixture
def png() -> bytes:
    buf = io.BytesIO()
    PIL.new("RGB", (200, 100), "lightblue").save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def tall_png() -> bytes:
    from PIL import ImageDraw
    im = PIL.new("RGB", (400, 2000), "white")
    dr = ImageDraw.Draw(im)
    for y in range(0, 2000, 250):
        dr.rectangle((50, y + 20, 350, y + 200), outline="black", width=3)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def make_template(build, path):
    """build(doc) наполняет документ; возвращает путь."""
    d = Document()
    build(d)
    d.save(path)
    return str(path)


@pytest.fixture
def template(tmp_path):
    def _mk(build, name="tpl.docx"):
        return make_template(build, tmp_path / name)
    return _mk


def ptext(p) -> str:
    """Текст абзаца включая поля (SEQ/REF) и гиперссылки — python-docx их не отдаёт."""
    from docx.oxml.ns import qn
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def texts(path):
    return [ptext(p) for p in Document(path).paragraphs]
