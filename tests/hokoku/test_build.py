"""
build_report и `python -m hokoku.build`: прогон целиком на поддельном хранилище.

Хранилище — словарь «идентификатор → байты» в памяти: сборка обязана работать, ничего
не зная о диске. Отдельно — проверка, что контракт выражает всё, что hokoku умеет
(случаи лаборатории labs/check).
"""
import hashlib
import io
import json
import os
import pathlib
import subprocess
import sys
import zipfile

import pytest
from docx import Document

from hokoku import (Blocks, Code, Diagram, Formula, Image, Markdown, PageBreak, Table, Text, Toc,
                    WIRE_VERSION, build_report)
from hokoku.pdf import libreoffice_available
from hokoku.wire import value_from_json, value_to_json

from .conftest import make_template

PACKAGES = str(pathlib.Path(__file__).resolve().parents[2] / "packages")
XML = ('<mxfile><diagram name="лист 1"><mxGraphModel><root/></mxGraphModel></diagram>'
       '<diagram name="лист 2"><mxGraphModel><root/></mxGraphModel></diagram></mxfile>')


def _tpl_bytes(build, tmp_path, name="tpl.docx") -> bytes:
    with open(make_template(build, tmp_path / name), "rb") as f:
        return f.read()


@pytest.fixture
def store(tmp_path, png):
    def build(d):
        d.add_paragraph("Цель: {{цель}}.")
        d.add_paragraph("{{схема}}")
        d.add_paragraph("{{замеры}}")
        d.add_paragraph("{{задачи}}")
        d.add_paragraph("см. {ref:схема}")
    return {"tpl_1": _tpl_bytes(build, tmp_path), "af_png": png, "af_xml": XML.encode("utf-8"),
            "af_broken": b"not an image at all"}


@pytest.fixture
def resolve(store):
    def _resolve(art_id):
        return store[art_id]
    return _resolve


@pytest.fixture
def job():
    return {"wire_version": WIRE_VERSION,
            "template": {"artifact": "tpl_1", "manifest_version": 4},
            "values": {
                "цель": {"v": 1, "type": "markdown", "text": "Изучить методы сортировки"},
                "схема": {"v": 1, "type": "image", "artifact": "af_png", "caption": "Схема"},
                "замеры": {"v": 1, "type": "table", "rows": [["n", "t"], ["10", "1,2"]],
                           "caption": "Замеры"}},
            "options": {"name": "otchet"}}


def run(job, resolve, tmp_path, **kw):
    return build_report(job, resolve_artifact=resolve, workdir=str(tmp_path / "out"), **kw)


# ── форма результата ──────────────────────────────────────────────────────────

def test_отчёт_собран(job, resolve, tmp_path):
    res = run(job, resolve, tmp_path)
    assert res["ok"] and res["wire_version"] == WIRE_VERSION
    assert res["outputs"]["docx"]["file"] == "otchet.docx"
    assert res["outputs"]["docx"]["bytes"] > 0 and len(res["outputs"]["docx"]["sha256"]) == 64
    assert res["counts"] == {"figures": 1, "tables": 1, "formulas": 0, "listings": 0}
    assert res["refs"] == {"схема": 1, "замеры": 1}
    assert res["unfilled"] == ["задачи"] and res["unknown_keys"] == []
    assert res["errors"] == [] and res["unresolved_refs"] == []
    assert res["template"]["manifest_version"] == 4
    assert res["timings_ms"]["total"] >= 0
    out = tmp_path / "out" / "otchet.docx"
    assert out.exists() and Document(str(out)).paragraphs


def test_в_ответе_имена_файлов_а_не_пути(job, resolve, tmp_path):
    """Путь на диске в JSON — то же самое, чего мы не пускаем внутрь; каталог назвал
    вызывающий, склеить он умеет сам."""
    res = run(job, resolve, tmp_path)
    assert os.sep not in res["outputs"]["docx"]["file"]
    assert "workdir" not in json.dumps(res) and str(tmp_path) not in json.dumps(res)


def test_наружу_кладётся_только_то_что_просили(job, resolve, tmp_path):
    """`outputs: ["pdf"]` — и DOCX в каталоге не появляется. Раньше он писался туда
    всегда, потому что LibreOffice звали по пути; теперь PDF считается из байтов, и
    файл, которого не просили, класть больше незачем."""
    job["options"]["outputs"] = ["pdf"]
    run(job, resolve, tmp_path)
    assert [p.name for p in (tmp_path / "out").iterdir()
            if p.name.endswith(".docx")] == []


def test_частичный_успех(job, resolve, tmp_path):
    """ok означает ровно одно: DOCX собран. Отчёт с пометками полезнее отказа."""
    job["values"]["битая"] = {"v": 1, "type": "image", "artifact": "af_broken"}
    job["values"]["нет_такого"] = {"v": 1, "type": "image", "artifact": "af_missing"}
    job["values"]["чужое"] = {"v": 1, "type": "нечто"}
    job["values"]["цел"] = {"v": 1, "type": "text", "text": "опечатка в ключе"}
    res = run(job, resolve, tmp_path)
    assert res["ok"]
    stages = {e["key"]: e["stage"] for e in res["errors"]}
    assert stages == {"битая": "wire", "нет_такого": "artifact", "чужое": "wire"}
    assert res["unknown_keys"] == ["цел"]
    warn = next(w for w in res["warnings"] if w["code"] == "unknown_key")
    assert warn["key"] == "цел" and warn["suggest"] == ["цель"]


def test_форма_предупреждения_в_json_не_поменялась(job, resolve, tmp_path):
    """Внутри пакета предупреждение стало `Problem` (общий `kyotsu.Notice`), но
    наружу уезжает `to_dict()` — те же ключи, что и до 2.0.0a4.2.

    Проверяется именно JSON, а не объект: результат `build_report` печатается
    `python -m hokoku.build`, и разбирает его чужой скрипт, который про наши
    dataclass'ы ничего не знает. Заодно — что запись сериализуется: объект,
    просочившийся в результат, свалил бы `json.dumps` уже у вызывающего.
    """
    job["values"]["цел"] = {"v": 1, "type": "text", "text": "опечатка в ключе"}
    res = run(job, resolve, tmp_path)
    warn = next(w for w in res["warnings"] if w["code"] == "unknown_key")
    assert set(warn) == {"module", "level", "code", "key", "message", "suggest"}
    assert warn["module"] == "hokoku" and warn["level"] == "warning"
    # Позиции в исходнике у предупреждений о тегах нет — и она не притворяется
    # пустой: незаданное поле в JSON не попадает вовсе.
    assert "file" not in warn and "line" not in warn
    json.dumps(res, ensure_ascii=False)


def test_пустое_значение_считает_render(job, resolve, tmp_path):
    """Проверка пустоты живёт в одном месте — в render; wire её не дублирует."""
    job["values"]["цель"] = {"v": 1, "type": "markdown", "text": "   "}
    res = run(job, resolve, tmp_path)
    assert res["ok"]
    assert [(e["key"], e["stage"]) for e in res["errors"]] == [("цель", "render")]


def test_битое_значение_роняет_задание_при_raise(job, resolve, tmp_path):
    job["values"]["чужое"] = {"v": 1, "type": "нечто"}
    job["options"]["on_error"] = "raise"
    res = run(job, resolve, tmp_path)
    assert res["ok"] is False and res["error"]["code"] == "render_failed"


def test_пустой_набор_значений_законен(job, resolve, tmp_path):
    job["values"] = {}
    res = run(job, resolve, tmp_path)
    assert res["ok"] and sorted(res["unfilled"]) == ["задачи", "замеры", "схема", "цель"]


def test_ссылка_в_никуда_не_роняет(job, resolve, tmp_path):
    job["values"]["цель"] = {"v": 1, "type": "markdown", "text": "см. {ref:нету}"}
    res = run(job, resolve, tmp_path)
    assert res["ok"] and res["unresolved_refs"] == ["нету"]
    assert [w["code"] for w in res["warnings"]] == ["unresolved_ref"]


def test_схема_листом_и_целиком(job, resolve, tmp_path):
    """page проверяется до сборки: выход за границу — ошибка значения, а не пустой рисунок."""
    job["values"]["схема"] = {"v": 1, "type": "diagram", "artifact": "af_xml", "page": 9}
    res = run(job, resolve, tmp_path)
    assert res["ok"] and res["errors"][0]["stage"] == "wire"
    assert "лист 9" in res["errors"][0]["message"]


# ── режим store_artifact ──────────────────────────────────────────────────────

@pytest.fixture
def store_mode(store):
    """Поддельное хранилище: тот же словарь «идентификатор → байты», что и на чтении.

    Ровно так же ляжет настоящее (`materials.Store.blob(mid) -> bytes` — это и есть
    сигнатура `resolve_artifact`), но `hokoku` про него не знает ничего: он отдал
    байты в функцию и получил взамен строку.
    """
    put = []

    def store_artifact(name, data, kind):
        art_id = f"af_{kind}_{len(put)}"
        store[art_id] = data
        put.append((name, kind, art_id))
        return art_id
    store_artifact.put = put
    return store_artifact


def test_в_ответе_идентификаторы_а_не_файлы(job, resolve, store, store_mode, tmp_path):
    """`hok-bytes-output`: на диске не появляется ничего, документ адресуется
    идентификатором — тем же, по которому его завтра можно вставить в другой отчёт."""
    было = set(tmp_path.iterdir())
    res = build_report(job, resolve_artifact=resolve, store_artifact=store_mode)
    assert res["ok"] and "file" not in res["outputs"]["docx"]
    art = res["outputs"]["docx"]["artifact"]
    assert store_mode.put == [("otchet.docx", "docx", art)]
    assert store[art][:2] == b"PK" and Document(io.BytesIO(store[art])).paragraphs
    assert res["outputs"]["docx"]["bytes"] == len(store[art])
    assert res["outputs"]["docx"]["sha256"] == hashlib.sha256(store[art]).hexdigest()
    assert set(tmp_path.iterdir()) == было               # ни одного файла наружу


def test_остальной_ответ_у_режимов_один(job, resolve, store_mode, tmp_path):
    """Режим — свойство приёмника байтов, а не сборки: всё, кроме `outputs`, совпадает."""
    a = build_report(job, resolve_artifact=resolve, store_artifact=store_mode)
    b = run(job, resolve, tmp_path)
    del a["outputs"], a["timings_ms"], b["outputs"], b["timings_ms"]
    assert a == b


def test_идентификатор_проверяется_той_же_регуляркой(job, resolve, tmp_path):
    """Строка от хранилища завтра приедет обратно во вход. Хранилище, отдавшее путь,
    сделало бы дырой не себя, а нас, — поэтому ARTIFACT_RE и на выходе."""
    for bad in ("out/otchet.docx", "../otchet", "", "af " * 40, 17, None):
        res = build_report(job, resolve_artifact=resolve,
                           store_artifact=lambda n, d, k, v=bad: v)
        assert res["ok"] is False and res["error"]["code"] == "internal"
        assert "идентификатор" in res["error"]["message"]


def test_отказ_хранилища_это_отказ_задания(job, resolve, tmp_path):
    def сломанное(name, data, kind):
        raise OSError("места нет")

    res = build_report(job, resolve_artifact=resolve, store_artifact=сломанное)
    assert res["ok"] is False and res["error"]["code"] == "internal"
    assert "места нет" in res["error"]["message"]


def test_потолок_срабатывает_и_без_каталога(job, resolve, store_mode):
    """Проверки уровня задания режима не знают: они до приёмника."""
    job["options"]["limits"] = {"max_table_rows": 1}
    res = build_report(job, resolve_artifact=resolve, store_artifact=store_mode)
    assert res["ok"] is False and res["error"]["code"] == "limit_exceeded"
    assert store_mode.put == []


@pytest.mark.skipif(not libreoffice_available(), reason="нужен LibreOffice")
def test_pdf_тоже_уезжает_артефактом(job, resolve, store, store_mode, tmp_path):
    job["options"]["outputs"] = ["docx", "pdf"]
    было = set(tmp_path.iterdir())
    res = build_report(job, resolve_artifact=resolve, store_artifact=store_mode)
    assert res["ok"] and [k for _, k, _ in store_mode.put] == ["docx", "pdf"]
    assert store[res["outputs"]["pdf"]["artifact"]][:4] == b"%PDF"
    assert set(tmp_path.iterdir()) == было               # PDF считался во временном каталоге


@pytest.mark.skipif(libreoffice_available(), reason="LibreOffice установлен — беда не воспроизводится")
def test_без_libreoffice_docx_уезжает_и_здесь(job, resolve, store_mode):
    job["options"]["outputs"] = ["docx", "pdf"]
    res = build_report(job, resolve_artifact=resolve, store_artifact=store_mode)
    assert res["ok"] and "pdf" not in res["outputs"] and res["outputs"]["docx"]["artifact"]
    assert [e["stage"] for e in res["errors"]] == ["pdf"]


# ── отказы уровня задания ─────────────────────────────────────────────────────

def test_нужен_ровно_один_приёмник(job, resolve, tmp_path):
    """Ни одного и оба — одинаково отказ: молча выбранный режим это готовый отчёт,
    оставшийся не там, где его будут искать."""
    with pytest.raises(ValueError, match="ровно один.*ни один"):
        build_report(job, resolve_artifact=resolve)
    with pytest.raises(ValueError, match="ровно один.*оба"):
        build_report(job, resolve_artifact=resolve, workdir=str(tmp_path / "out"),
                     store_artifact=lambda n, d, k: "af_1")
    with pytest.raises(ValueError, match="store_artifact"):
        build_report(job, resolve_artifact=resolve, store_artifact="не функция")


def test_коды_отказов(job, resolve, tmp_path):
    def code(patch):
        j = json.loads(json.dumps(job))
        patch(j)
        return run(j, resolve, tmp_path)["error"]["code"]

    assert code(lambda j: j.update(wire_version=9)) == "unknown_wire_version"
    assert code(lambda j: j["options"].update(on_errr="skip")) == "bad_job"
    assert code(lambda j: j["options"].update(on_error="молча")) == "bad_job"
    assert code(lambda j: j["options"].update(style={"нет_такого": {}})) == "bad_job"
    assert code(lambda j: j["template"].update(artifact="нет")) == "bad_template"
    assert code(lambda j: j["template"].update(sha256="00")) == "bad_template"
    assert code(lambda j: j.pop("template")) == "bad_job"
    # наружу всегда JSON, даже когда задание кривое насквозь: иначе шов перестаёт быть швом
    assert code(lambda j: j.update(values=["значение"])) == "bad_job"


def test_шаблон_с_макросом_отклонён(job, resolve, store, tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", "<w:document/>")
        z.writestr("word/vbaProject.bin", "MACRO")
    store["tpl_vba"] = buf.getvalue()
    job["template"] = {"artifact": "tpl_vba"}
    res = run(job, resolve, tmp_path)
    assert res["ok"] is False and res["error"]["code"] == "template_rejected"


def test_потолок_срабатывает_до_сборки(job, resolve, tmp_path):
    job["values"]["замеры"] = {"v": 1, "type": "table", "rows": [["a", "b"]] * 50}
    job["options"]["limits"] = {"max_table_rows": 10}
    res = run(job, resolve, tmp_path)
    assert res["ok"] is False and res["error"]["code"] == "limit_exceeded"
    assert res["error"]["key"] == "замеры" and "max_table_rows" in res["error"]["message"]
    assert not (tmp_path / "out").exists()


def test_потолок_знаков_считает_подписи(job, resolve, tmp_path):
    """Подпись — написанный человеком текст, и в знаки она входит: пока `_chars` её
    не считал, `max_value_chars` недосчитывал, а картинка не весила ни знака."""
    job["values"]["схема"] = {"v": 1, "type": "image", "artifact": "af_png",
                              "caption": "и" * 200}
    job["options"]["limits"] = {"max_value_chars": 100}
    res = run(job, resolve, tmp_path)
    assert res["ok"] is False and res["error"]["code"] == "limit_exceeded"
    assert res["error"]["key"] == "схема" and "max_value_chars" in res["error"]["message"]

    job["values"]["схема"]["caption"] = "коротко"
    assert run(job, resolve, tmp_path)["ok"]


def test_умолчания_потолков_щедрые(job, resolve, tmp_path):
    """Настоящий журнал замеров на 300 строк — обычный отчёт, а не повод для отказа."""
    rows = [["№", "метод", "мс"]] + [[str(i), "пузырёк", "41"] for i in range(300)]
    job["values"]["замеры"] = {"v": 1, "type": "table", "rows": rows}
    assert run(job, resolve, tmp_path)["ok"]


def test_потолок_задания_не_выше_жёсткого(job, resolve, tmp_path):
    job["options"]["limits"] = {"max_table_rows": 10 ** 9}
    job["values"]["замеры"] = {"v": 1, "type": "table", "rows": [["a"]] * 3000}
    res = run(job, resolve, tmp_path)
    assert res["ok"] is False and res["error"]["code"] == "limit_exceeded"


def test_потолок_на_байты_артефакта(job, resolve, tmp_path):
    job["options"]["limits"] = {"max_artifact_bytes": 10}
    res = run(job, resolve, tmp_path)
    assert res["ok"] is False and res["error"]["code"] == "limit_exceeded"


# ── PDF ───────────────────────────────────────────────────────────────────────

@pytest.mark.skipif(libreoffice_available(), reason="LibreOffice установлен — беда не воспроизводится")
def test_без_libreoffice_docx_остаётся(job, resolve, tmp_path):
    job["options"]["outputs"] = ["docx", "pdf"]
    res = run(job, resolve, tmp_path)
    assert res["ok"] and "pdf" not in res["outputs"]
    assert [e["stage"] for e in res["errors"]] == ["pdf"] and res["errors"][0]["key"] is None


@pytest.mark.skipif(not libreoffice_available(), reason="нужен LibreOffice")
def test_pdf_собирается(job, resolve, tmp_path):
    job["options"]["outputs"] = ["docx", "pdf"]
    res = run(job, resolve, tmp_path)
    assert res["ok"] and res["outputs"]["pdf"]["file"] == "otchet.pdf"
    assert (tmp_path / "out" / "otchet.pdf").read_bytes()[:4] == b"%PDF"


# ── CLI ───────────────────────────────────────────────────────────────────────

def _cli(args, tmp_path, stdin=None):
    env = dict(os.environ, PYTHONPATH=PACKAGES)
    return subprocess.run([sys.executable, "-m", "hokoku.build", *args], input=stdin,
                          capture_output=True, text=True, env=env, cwd=str(tmp_path))


def test_cli_собирает_отчёт(job, store, tmp_path):
    arts = tmp_path / "arts"
    arts.mkdir()
    for art_id, data in store.items():
        (arts / art_id).write_bytes(data)
    (tmp_path / "job.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")

    r = _cli(["job.json", "--artifacts", "arts", "--out", "out"], tmp_path)
    assert r.returncode == 0, r.stderr
    res = json.loads(r.stdout)                      # JSON на stdout — разбирается как есть
    assert res["ok"] and res["outputs"]["docx"]["file"] == "otchet.docx"
    assert "собрано" in r.stderr                    # человеческое — на stderr
    assert Document(str(tmp_path / "out" / "otchet.docx")).paragraphs

    # 1 — не собралось, 2 — задание не прочиталось
    bad = dict(job, template={"artifact": "нет_такого"})
    (tmp_path / "bad.json").write_text(json.dumps(bad, ensure_ascii=False), encoding="utf-8")
    r = _cli(["bad.json", "--artifacts", "arts", "--out", "out"], tmp_path)
    assert r.returncode == 1 and json.loads(r.stdout)["error"]["code"] == "bad_template"
    assert _cli(["нет_файла.json", "--artifacts", "arts", "--out", "out"], tmp_path).returncode == 2
    (tmp_path / "кривой.json").write_text("{не json", encoding="utf-8")
    assert _cli(["кривой.json", "--artifacts", "arts", "--out", "out"], tmp_path).returncode == 2

    r = _cli(["--stdin", "--artifacts", "arts", "--out", "out", "--pretty"], tmp_path,
             stdin=json.dumps(job, ensure_ascii=False))
    assert r.returncode == 0 and json.loads(r.stdout)["ok"]


def test_cli_без_out_кладёт_отчёт_артефактом(job, store, tmp_path):
    """Второй режим виден и из скрипта: без `--out` готовый DOCX уезжает в каталог
    артефактов под именем от содержимого, а в JSON стоит идентификатор, не имя файла."""
    arts = tmp_path / "arts"
    arts.mkdir()
    for art_id, data in store.items():
        (arts / art_id).write_bytes(data)
    (tmp_path / "job.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")

    было = {p.name for p in arts.iterdir()}
    r = _cli(["job.json", "--artifacts", "arts"], tmp_path)
    assert r.returncode == 0, r.stderr
    res = json.loads(r.stdout)
    art = res["outputs"]["docx"]["artifact"]
    assert "file" not in res["outputs"]["docx"] and art.startswith("af_")
    assert not (tmp_path / "out").exists()          # каталога с именами не завелось
    assert {p.name for p in arts.iterdir()} - было == {art}
    assert Document(str(arts / art)).paragraphs

    # ни каталога, ни хранилища — «так звать нельзя», а не молчаливый выбор режима
    r = _cli(["job.json"], tmp_path)
    assert r.returncode == 2 and "так звать нельзя" in r.stderr


# ── возможности лаборатории labs/check ────────────────────────────────────────

def lab_values(png, tall_png):
    """По одному значению на случай лаборатории, где случай — про значение.

    Остальные случаи hokoku в JSON значений не выражаются и не должны: где стоит тег —
    свойство шаблона (hok-place-*), unfilled/unknown_keys и защита шаблона — свойства
    результата (hok-unfilled-unknown, hok-safety, hok-extract-tags), а on_error,
    strict_paths, images_dir, оформление подписей и вывод в байтах — опции прогона
    (hok-on-error-skip, hok-strict-paths, hok-md-images-dir, hok-style-gost,
    hok-bytes-output, hok-pdf).
    """
    return {
        "hok-text": Text("строка 1\nстрока 2"),
        "hok-markdown": Markdown("## Заголовок\n\n- пункт\n\n> цитата\n\n---"),
        "hok-md-tasks-links": Markdown("- [x] сделано\n\n[Ссылка](https://example.org)"),
        "hok-code": Code("def f(a):\n    return a\n", "python", line_numbers=True),
        "hok-image": Image(png, caption="Схема установки"),
        "hok-image-nocaption": Image(png, caption=False),
        "hok-image-width": Image(png, caption="Узкая слева", width_cm=6.0, align="left"),
        "hok-image-svg": Image(b'<svg xmlns="http://www.w3.org/2000/svg"><rect/></svg>',
                               caption="Векторная"),
        "hok-image-tall": Image(tall_png, caption="Длинная схема"),
        "hok-table": Table([["№", "Параметр", "Значение"], ["1", "К передачи", "1,25"]],
                           caption="Результаты", align=["left", "center", "right"]),
        "hok-table-widths": Table([["№", "Параметр", "Ед."], ["1", "КНИ", "%"]],
                                  caption="Ширины", col_widths_cm=[1.2, 8.0, 1.5]),
        "hok-table-md": Markdown("| Система | Запись |\n|:---|---:|\n| oct | 31 |"),
        "hok-formula": Formula(r"\frac{\sum_{i=1}^{n} x_i}{n}"),
        "hok-toc": Toc(levels=2, title="Содержание"),
        "hok-blocks": Blocks([Text("Ниже результаты."), Image(png, caption="Установка"),
                              Table([["n", "t"], ["10", "1,2"]], caption="Замеры"), PageBreak(),
                              Code("int main() {\n    return 0;\n}", "cpp")]),
        "hok-refs": Markdown("см. рисунок {ref:схема} и таблицу {ref:замеры}"),
        "hok-ref-in-cell": Table([["что"], ["см {ref:схема}"]], caption="Со ссылкой"),
        "hok-value-not-expanded": Text("тут написано {{другой_тег}} — так и останется"),
        "hok-scalars": Text("да"),
        "hok-diagram-page": Diagram(XML, caption="Только второй лист", page=2),
        "int-flow-in-report": Diagram(XML, caption="Сортировка пузырьком"),
        "int-uml-in-report": Diagram(XML, caption="Классы сортировщиков"),
        "int-multipage-sheets": Diagram(XML, caption="Три функции"),
        "int-diagram-in-cell": Diagram(XML, caption="В ячейке"),
    }


def test_возможности_лаборатории_выражаются_в_json(png, tall_png, store):
    """Каждое значение из labs/check выражается в JSON и возвращается тем же значением.

    Диаграммы лаборатория отдаёт XML строкой (fragmos и uml не знают про хранилище) —
    именно поэтому у diagram есть xml, а не только artifact.
    """
    back = {bytes(data): art_id for art_id, data in store.items()}
    values = lab_values(png, tall_png)
    for cid, value in values.items():
        data = value_to_json(value, artifact_of=lambda b: back.get(bytes(b), "af_png"))
        again = value_from_json(data, resolve_artifact=lambda i: bytes(value.source)
                                if isinstance(value, Image) else store[i])
        assert again == value, cid
    assert len(values) == 24
