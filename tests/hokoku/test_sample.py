"""
Пример готового отчёта → оформление. Проверяем не профиль сам по себе, а то, что
снятое с примера доезжает до собранного файла: профиль, в котором всё правильно, но
который никуда не применяется, — худший вид ошибки, он выглядит как работающий.
"""
import io
import zipfile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from hokoku import (DocxValidationError, Image, Markdown, SampleError, apply_style,
                    document_bytes, document_from_sample, outline_from_sample, render,
                    style_from_sample)
from .conftest import ptext


def _sample(*, captions=("Рис. 1. Схема алгоритма",), headings=True, tags=False) -> bytes:
    """Кафедральный образец: чужие поля, чужой шрифт, свои подписи. Заголовки набраны
    прямым форматированием поверх стиля — так и выглядят настоящие работы."""
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin, s.right_margin = Cm(2.5), Cm(1.0)
    s.top_margin, s.bottom_margin = Cm(2.0), Cm(2.0)
    if headings:
        for text, level, size in (("1 Введение", 1, 18), ("1.1 Постановка задачи", 2, 15),
                                  ("2 Заключение", 1, 18)):
            h = d.add_paragraph(text, style=f"Heading {level}")
            for r in h.runs:
                r.font.name, r.font.size, r.font.bold = "Arial", Pt(size), True
    for text in ("Первый абзац основного текста образца, достаточно длинный.",
                 "Второй абзац основного текста образца, тоже длинный.",
                 "Третий абзац образца — ради устойчивости самого частого значения."):
        p = d.add_paragraph(text)
        p.paragraph_format.first_line_indent = Cm(1.0)
        p.paragraph_format.line_spacing = 1.5
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            r.font.name, r.font.size = "Verdana", Pt(13)
    for cap in captions:
        p = d.add_paragraph(cap, style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if tags:
        d.add_paragraph("{{цель}}")
    return document_bytes(d)


def _with_member(data: bytes, name: str, blob: bytes) -> bytes:
    """Тот же DOCX с подложенным членом архива — так выглядит присланный «образец»
    с макросами: содержимое честное, беда в невидимой части."""
    src = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            out.writestr(info.filename, src.read(info.filename))
        out.writestr(name, blob)
    return buf.getvalue()


# ── что снимается с примера ───────────────────────────────────────────────────

def test_profile_reads_page_body_headings_captions():
    p = style_from_sample(_sample())
    assert (p.page.left_cm, p.page.right_cm) == (2.5, 1.0)
    assert (p.body.font, p.body.size_pt) == ("Verdana", 13.0)
    assert p.body.line_spacing == 1.5 and p.body.first_line_cm == 1.0
    assert [(h.level, h.font, h.size_pt, h.bold) for h in p.headings] == [
        (1, "Arial", 18.0, True), (2, "Arial", 15.0, True)]
    assert p.captions["figure"] == "Рис. {n}. {caption}"
    assert p.captions["figure_align"] == "center"
    assert p.numbered_headings is True                  # «1 Введение» — номер руками


def test_body_text_taken_from_paragraphs_not_from_normal():
    """В кафедральных документах Normal остаётся заводским, а текст набран прямым
    форматированием: сняв стиль Normal, мы вернули бы оформление, которого не видно."""
    p = style_from_sample(_sample())
    assert p.body.font == "Verdana"
    assert Document(io.BytesIO(_sample())).styles["Normal"].font.name != "Verdana"


def test_theme_fonts_are_resolved():
    """У настоящего Word шрифт заголовка задан темой (majorHAnsi), а не именем.
    Без разрешения темы профиль вернул бы None и заголовки остались бы нашими."""
    d = Document()
    d.add_paragraph("Введение", style="Heading 1")       # тема, прямого шрифта нет
    for text in ("Достаточно длинный абзац основного текста примера.",) * 3:
        d.add_paragraph(text)
    assert d.styles["Heading 1"].font.name is None            # имени шрифта в стиле нет
    p = style_from_sample(document_bytes(d))
    assert p.headings[0].font == "Calibri"                     # majorHAnsi штатной темы


def test_notes_name_what_cannot_be_taken():
    p = style_from_sample(_sample(captions=("Рисунок 1.2 — Схема", "Рисунок 1.3 — Ещё")))
    assert any("по главам" in n for n in p.notes)        # «Рисунок 1.2» повторить нечем
    assert p.captions["figure"] == "Рисунок {n} — {caption}"


# ── применение к нашему документу ─────────────────────────────────────────────

def test_sample_style_reaches_the_built_file(tmp_path, png):
    """Главная проверка задачи 2: смотрим сам собранный файл, а не профиль."""
    doc, profile = document_from_sample(_sample())
    doc.add_paragraph("{{тело}}")
    doc.add_paragraph("{{рис}}")
    out = str(tmp_path / "out.docx")
    res = render(document_bytes(doc),
                 {"тело": Markdown("# Раздел\n\nтекст"), "рис": Image(png, caption="Схема")},
                 out, style=profile.style_overrides())
    assert res.errors == [] and res.figures == 1
    o = Document(out)
    assert round(o.sections[0].left_margin.cm, 1) == 2.5
    assert round(o.sections[0].right_margin.cm, 1) == 1.0
    assert o.styles["Normal"].font.name == "Verdana" and o.styles["Normal"].font.size.pt == 13
    assert o.styles["Heading 1"].font.name == "Arial"
    assert o.styles["Heading 1"].font.size.pt == 18
    texts = [(ptext(p), p.style.name) for p in o.paragraphs]
    assert ("Раздел", "Heading 1") in texts
    # подпись едет вторым путём — через render(style=…), а не через стиль Caption
    assert [t for t, n in texts if n == "Caption"] == ["Рис. 1. Схема"]


def test_heading_numbering_is_word_not_typed(tmp_path):
    """Нумерация разделов — полем Word: вписанные руками «1.1» перестают быть верными
    при первой же вставке раздела."""
    doc, profile = document_from_sample(_sample())
    assert profile.numbered_headings
    numpr = doc.styles["Heading 1"].element.find(qn("w:pPr")).find(qn("w:numPr"))
    assert numpr is not None
    num_id = numpr.find(qn("w:numId")).get(qn("w:val"))
    numbering = doc.part.numbering_part.element
    num = [n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == num_id][0]
    abs_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    absn = [a for a in numbering.findall(qn("w:abstractNum"))
            if a.get(qn("w:abstractNumId")) == abs_id][0]
    lvls = absn.findall(qn("w:lvl"))
    assert lvls[0].find(qn("w:pStyle")).get(qn("w:val")) == "Heading1"
    assert lvls[1].find(qn("w:lvlText")).get(qn("w:val")) == "%1.%2."


def test_apply_style_creates_missing_styles_by_the_same_mechanism():
    """Второго набора запасных стилей нет: apply_style зовёт тот же ensure_styles."""
    d = Document()
    root = d.styles.element
    for st in list(root.findall(qn("w:style"))):
        if st.find(qn("w:name")).get(qn("w:val")) not in ("Normal", "Default Paragraph Font"):
            root.remove(st)
    d._hokoku_style_ids = None
    apply_style(d, style_from_sample(_sample()))
    assert d.styles["Heading 1"].font.name == "Arial"
    assert "Code" in [s.name for s in d.styles]


# ── недоверенный вход и внятные отказы ────────────────────────────────────────

def test_sample_with_macros_rejected():
    """Пример приносит пользователь: та же дверь и та же проверка, что у шаблона."""
    data = _with_member(_sample(), "word/vbaProject.bin", b"MACRO")
    with pytest.raises(DocxValidationError):
        style_from_sample(data)
    with pytest.raises(DocxValidationError):
        outline_from_sample(data)


def test_sample_not_a_docx_rejected():
    with pytest.raises(DocxValidationError):
        style_from_sample("это не docx".encode())


def test_sample_without_headings_complains():
    """Пример без нужных стилей обязан жаловаться словами, а не возвращать пустоту:
    молчание здесь означало бы отчёт с нашим оформлением под видом кафедрального."""
    with pytest.raises(SampleError) as e:
        style_from_sample(_sample(headings=False))
    assert "заголовк" in str(e.value)


def test_template_is_not_a_sample():
    """Шаблон с тегами вместо примера — первая же ошибка, которую сделает человек."""
    with pytest.raises(SampleError) as e:
        style_from_sample(_sample(tags=True))
    assert "{{" in str(e.value) and "шаблон" in str(e.value)
    assert style_from_sample(_sample(tags=True), allow_tags=True).body.font == "Verdana"


# ── строение, а не оформление ─────────────────────────────────────────────────

def test_outline_is_separate_and_shaped():
    out = outline_from_sample(_sample())
    assert out == [
        {"level": 1, "text": "Введение", "number": "1"},
        {"level": 2, "text": "Постановка задачи", "number": "1.1"},
        {"level": 1, "text": "Заключение", "number": "2"},
    ]
    # строение не подмешано в оформление: в профиле нет ни одного заголовка примера
    profile = style_from_sample(_sample())
    assert "Введение" not in repr(profile)


def test_outline_without_manual_numbers():
    d = Document()
    d.add_paragraph("Введение", style="Heading 1")
    d.add_paragraph("Обычный абзац, который заголовком не является.")
    assert outline_from_sample(document_bytes(d)) == [
        {"level": 1, "text": "Введение", "number": None}]
