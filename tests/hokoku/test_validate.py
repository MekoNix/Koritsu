"""
Проверка значений против манифеста и шаблона — до сборки документа.

Каждый тест здесь воспроизводит беду, которая до сих пор уезжала в готовый отчёт молча:
ограничения манифеста (`max_rows`, `max_chars`, `headings`) и `depends_on` разбирались,
печатались модели в промпт и никем не читались, а `{ref:}` в никуда становилась «?»
уже в собранном DOCX — на кафедре.

Модель здесь не зовётся: validate — чистая функция от шаблона, значений и манифеста.
"""
import io

import pytest
from docx import Document

from hokoku import Blocks, Code, Diagram, Formula, Image, Markdown, Table, Text, render
from hokoku.manifest import manifest_from_template
from hokoku.validate import validate

from .conftest import ptext


def docx_bytes(build) -> bytes:
    d = Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def built(template, values):
    """Собранный документ по этим значениям.

    Лимит стережёт не запись в списке проблем, а то, что окажется на бумаге: без сборки
    тест доказывал бы форму замечания, а не беду, из-за которой замечание заводили.
    """
    return Document(io.BytesIO(render(template, values).data))


@pytest.fixture
def template() -> bytes:
    def build(d):
        d.add_paragraph("Цель: {{цель:Цель работы}}.")
        d.add_paragraph("{{схема:Блок-схема алгоритма}}")
        d.add_paragraph("{{замеры:Таблица замеров}}")
        d.add_paragraph("{{вывод:Вывод}}")
        d.sections[0].header.paragraphs[0].text = "{{кафедра:Кафедра}}"
    return docx_bytes(build)


@pytest.fixture
def manifest(template):
    m = manifest_from_template(template)
    m.tags["кафедра"].type = "text"      # в колонтитуле строка, а не markdown
    return m


@pytest.fixture
def values():
    """Набор, к которому у validate нет вопросов: типы те, что объявил манифест."""
    return {"цель": Markdown("Изучить работу алгоритма."),
            "схема": Diagram(xml="<mxfile/>"),
            "замеры": Table(rows=[["t", "N"], ["1", "2"]]),
            "вывод": Markdown("Алгоритм работает."),
            "кафедра": Text("ИУ7")}


def codes(problems, key=None) -> list[str]:
    return [p.code for p in problems if key is None or p.key == key]


def one(problems, code) -> dict:
    found = [p for p in problems if p.code == code]
    assert len(found) == 1, f"ждали одну проблему {code}, получили {problems}"
    return found[0]


# ── общее ─────────────────────────────────────────────────────────────────────

def test_чистый_набор_без_проблем(template, values, manifest):
    assert validate(template, values, manifest) == []


def test_проблема_машиночитаема(template, values, manifest):
    """Проблему показывают в интерфейсе и отдают модели на повтор: она обязана быть
    записью с полями, а не строкой прозой."""
    values["замеры"] = Markdown("не таблица")
    p = one(validate(template, values, manifest), "type_mismatch")
    assert p.module == "hokoku" and p.level == "error" and p.key == "замеры"
    assert (p.expected, p.got) == ("table", "markdown")
    assert isinstance(p.message, str) and p.message


def test_все_проблемы_сразу_а_не_первая(template, values, manifest):
    """Список, а не исключение: модель должна получить все замечания за один повтор."""
    values["замеры"] = Markdown("не таблица")
    values["цель"] = Markdown("   ")
    del values["вывод"]
    got = set(codes(validate(template, values, manifest)))
    assert got == {"type_mismatch", "empty_value", "missing_required"}


# ── типы, пустое и пропущенное ────────────────────────────────────────────────

def test_тип_не_тот(template, values, manifest):
    """Модель отдала абзац туда, где ждали схему: без проверки это видно только глазами
    на готовом отчёте."""
    values["схема"] = Markdown("нарисовать не смог")
    p = one(validate(template, values, manifest), "type_mismatch")
    assert (p.key, p.expected, p.got) == ("схема", "diagram", "markdown")


def test_голая_строка_это_text(template, values, manifest):
    """`str`/`int`/`bool` render вставляет как простой текст — значит и тип у них text."""
    values["цель"] = "Изучить работу алгоритма."
    p = one(validate(template, values, manifest), "type_mismatch")
    assert (p.expected, p.got) == ("markdown", "text")


def test_обязательный_тег_без_значения(template, values, manifest):
    del values["цель"]
    p = one(validate(template, values, manifest), "missing_required")
    assert p.key == "цель" and p.level == "error"


def test_необязательный_тег_без_значения_не_ошибка(template, values, manifest):
    """Тег колонтитула необязателен: его вписывает человек, и пустым он остаётся законно."""
    del values["кафедра"]
    problems = validate(template, values, manifest)
    assert codes(problems) == ["tag_unfilled"]
    assert problems[0].level == "info"


def test_пустое_значение_ошибка(template, values, manifest):
    """«Пустое значение — ошибка» действует в render; validate обязан сказать это раньше,
    а не ронять сборку на середине."""
    values["вывод"] = Markdown("   ")
    p = one(validate(template, values, manifest), "empty_value")
    assert p.key == "вывод" and p.level == "error"


def test_значение_без_тега_с_подсказкой(template, values, manifest):
    """Опечатка в ключе (в том числе у модели) молча теряла целое значение."""
    values["цел"] = Markdown("Изучить.")
    p = one(validate(template, values, manifest), "unknown_key")
    assert p.key == "цел" and "цель" in p.message


def test_тег_без_записи_в_манифесте_проверяется_по_умолчанию(template, values, manifest):
    """Записи нет — тег всё равно markdown и обязательный; для колонтитула — необязательный."""
    del manifest.tags["цель"]
    del manifest.tags["кафедра"]
    values["цель"] = Table(rows=[["a"]])
    del values["кафедра"]
    problems = validate(template, values, manifest)
    assert one(problems, "type_mismatch").expected == "markdown"
    assert codes(problems, "кафедра") == ["tag_unfilled"]


# ── ограничения манифеста ─────────────────────────────────────────────────────

def test_лимит_знаков(template, values, manifest):
    manifest.tags["цель"].limits = {"max_chars": 20}
    values["цель"] = Markdown("а" * 500)
    p = one(validate(template, values, manifest), "limit_max_chars")
    assert (p.key, p.expected, p.got) == ("цель", 20, 500)


def test_лимит_знаков_считает_подписи(template, values, manifest):
    """Подпись — такой же написанный текст: пока `_chars` её не считал, `max_chars`
    недосчитывал знаки, а у рисунка и схемы не считалось вовсе ничего."""
    manifest.tags["схема"].limits = {"max_chars": 20}
    values["схема"] = Diagram(xml="<mxfile/>" * 100, caption="а" * 50)
    p = one(validate(template, values, manifest), "limit_max_chars")
    assert (p.key, p.got) == ("схема", 50)          # XML генератора по-прежнему мимо счёта

    manifest.tags["схема"].limits = {}
    manifest.tags["замеры"].limits = {"max_chars": 20}
    values["замеры"] = Table(rows=[["ab"]], caption="в" * 30)
    p = one(validate(template, values, manifest), "limit_max_chars")
    assert (p.key, p.got) == ("замеры", 32)         # подпись плюс ячейки


def test_лимит_строк_таблицы(template, values, manifest):
    """Тот самый случай: 500 строк при `max_rows: 20` уезжали в отчёт без единого слова."""
    manifest.tags["замеры"].limits = {"max_rows": 20}
    values["замеры"] = Table(rows=[[str(i)] for i in range(500)])
    p = one(validate(template, values, manifest), "limit_max_rows")
    assert (p.expected, p.got) == (20, 500)


def test_лимит_строк_видит_таблицу_внутри_markdown(template, values, manifest):
    """Таблица в markdown — такая же таблица: иначе лимит обходится сменой типа."""
    manifest.tags["цель"].limits = {"max_rows": 2}
    rows = "\n".join(f"| {i} | {i} |" for i in range(10))
    values["цель"] = Markdown("| a | b |\n| --- | --- |\n" + rows)
    assert one(validate(template, values, manifest), "limit_max_rows").got == 11


def test_лимит_строк_видит_таблицу_внутри_blocks(template, values, manifest):
    manifest.tags["цель"].type = "blocks"
    manifest.tags["цель"].limits = {"max_rows": 2}
    values["цель"] = Blocks(items=[Text("вот замеры"), Table(rows=[["a"], ["b"], ["c"]])])
    assert one(validate(template, values, manifest), "limit_max_rows").got == 3


def test_лимит_колонок(template, values, manifest):
    manifest.tags["замеры"].limits = {"max_cols": 3}
    values["замеры"] = Table(rows=[["a", "b"], ["1", "2", "3", "4"]])
    p = one(validate(template, values, manifest), "limit_max_cols")
    assert (p.expected, p.got) == (3, 4)


def test_заголовков_не_ставить(template, values, manifest):
    """`headings: false` — раздел отчёта, где заголовок ломает нумерацию оглавления."""
    manifest.tags["цель"].limits = {"headings": False}
    values["цель"] = Markdown("## Цель работы\n\nИзучить алгоритм.")
    p = one(validate(template, values, manifest), "limit_headings")
    assert (p.key, p.expected, p.got) == ("цель", 0, 1)
    assert "Цель работы" in p.message


def test_заголовки_разрешены_по_умолчанию(template, values, manifest):
    values["цель"] = Markdown("## Цель работы\n\nИзучить алгоритм.")
    assert validate(template, values, manifest) == []


def test_заголовок_из_простого_текста_ловится(template, values, manifest):
    """`headings: false` обходился сменой типа: render превращает `Text` в markdown,
    как только в нём встретилась `{ref:}`, а ссылки модель ставить обязана."""
    manifest.tags["цель"].type = "text"
    manifest.tags["цель"].limits = {"headings": False}
    values["цель"] = Text("# Цель работы\n\nКак видно на {ref:схема}, всё сходится.")
    p = one(validate(template, values, manifest), "limit_headings")
    assert (p.key, p.got) == ("цель", 1) and "Цель работы" in p.message
    # и это не буквы «# » в абзаце, а настоящий заголовок, ломающий нумерацию оглавления
    d = built(template, {"цель": values["цель"]})
    assert "Heading 1" in [p.style.name for p in d.paragraphs]


def test_заголовок_из_голой_строки_ловится(template, values, manifest):
    """Голую строку со ссылкой render разбирает как markdown ровно так же."""
    manifest.tags["цель"].type = "text"
    manifest.tags["цель"].limits = {"headings": False}
    values["цель"] = "# Цель работы\n\nСм. {ref:схема}."
    assert one(validate(template, values, manifest), "limit_headings").got == 1
    d = built(template, {"цель": values["цель"]})
    assert "Heading 1" in [p.style.name for p in d.paragraphs]


def test_лимит_строк_обходился_ссылкой_в_простом_тексте(template, values, manifest):
    """Тот же обход снимал и `max_rows`: разметка таблицы в тексте со ссылкой доезжала
    до отчёта настоящей таблицей."""
    manifest.tags["цель"].type = "text"
    manifest.tags["цель"].limits = {"max_rows": 2}
    rows = "\n".join(f"| {i} | {i} |" for i in range(10))
    values["цель"] = Text("См. {ref:схема}.\n\n| a | b |\n| --- | --- |\n" + rows)
    assert one(validate(template, values, manifest), "limit_max_rows").got == 11
    assert [len(t.rows) for t in built(template, {"цель": values["цель"]}).tables] == [11]


def test_текст_без_ссылки_остаётся_текстом(template, values, manifest):
    """Без `{ref:}` render разметку в простом тексте не разбирает — и придираться не к чему:
    иначе лимит начал бы ругаться на решётки, которые в документе так решётками и лягут."""
    manifest.tags["цель"].type = "text"
    manifest.tags["цель"].limits = {"headings": False}
    values["цель"] = Text("# Цель работы")
    assert validate(template, values, manifest) == []
    d = built(template, {"цель": values["цель"]})
    assert [p.style.name for p in d.paragraphs] == ["Normal"] * len(d.paragraphs)


def test_заголовок_внутри_blocks_не_заголовок(template, values, manifest):
    """Подмену типа render делает только для самого значения тега: элемент `Blocks`
    уходит в документ как есть, и решётки остаются решётками."""
    manifest.tags["цель"].type = "blocks"
    manifest.tags["цель"].limits = {"headings": False}
    values["цель"] = Blocks(items=[Text("# Цель работы\n\nСм. {ref:схема}.")])
    assert codes(validate(template, values, manifest), "цель") == []
    d = built(template, {"цель": values["цель"]})
    assert "Heading 1" not in [p.style.name for p in d.paragraphs]


# ── depends_on ────────────────────────────────────────────────────────────────

def test_зависимость_на_несуществующий_тег(template, values, manifest):
    """Опечатка в depends_on делала связь украшением: молчит и не проверяет ничего."""
    manifest.tags["вывод"].depends_on = ["замерры"]
    p = one(validate(template, values, manifest), "depends_on_missing")
    assert (p.key, p.got) == ("вывод", "замерры")


def test_зависимость_на_незаполненный_тег(template, values, manifest):
    """Вывод по замерам, которых нет, модель сочинит — и это худший из исходов."""
    manifest.tags["вывод"].depends_on = ["замеры"]
    manifest.tags["замеры"].required = False
    del values["замеры"]
    p = one(validate(template, values, manifest), "depends_on_unfilled")
    assert (p.key, p.got) == ("вывод", "замеры")


def test_зависимость_незаполненного_необязательного_тега_молчит(template, values, manifest):
    """Сам тег пуст и не обязателен — опираться ему не на что, и говорить не о чем."""
    manifest.tags["вывод"].depends_on = ["замеры"]
    manifest.tags["вывод"].required = False
    manifest.tags["замеры"].required = False
    del values["вывод"]
    del values["замеры"]
    assert codes(validate(template, values, manifest), "вывод") == ["tag_unfilled"]


# ── {ref:} ────────────────────────────────────────────────────────────────────

def test_ссылка_в_никуда(template, values, manifest):
    """`{ref:}` мимо тега становится «?» в готовом документе — заметно уже на кафедре."""
    values["вывод"] = Markdown("Как видно на {ref:схемма}, всё сходится.")
    p = one(validate(template, values, manifest), "unresolved_ref")
    assert (p.key, p.got, p.level) == ("вывод", "схемма", "warning")


def test_ссылка_на_тег_шаблона_законна(template, values, manifest):
    values["вывод"] = Markdown("Как видно на {ref:схема}, всё сходится.")
    assert validate(template, values, manifest) == []


def test_ссылка_на_имя_из_значения_законна(template, values, manifest):
    """Имя ссылки может назначить само значение (`ref=`), а не только ключ тега."""
    values["замеры"] = Table(rows=[["t"], ["1"]], ref="замеры_1")
    values["вывод"] = Markdown("См. {ref:замеры_1}.")
    assert validate(template, values, manifest) == []


def test_ссылка_из_ячейки_таблицы(template, values, manifest):
    values["замеры"] = Table(rows=[["источник"], ["{ref:нетути}"]])
    assert codes(validate(template, values, manifest)) == ["unresolved_ref"]


def test_ссылка_в_подписи_рисунка(template, values, manifest, png):
    """Подпись — то же место, что и текст: `docx_ops._caption_text` разбирает в ней
    ссылки, и «?» в «Рисунок 1 — …» стоит на самом виду."""
    manifest.tags["схема"].type = "image"
    values["схема"] = Image(source=png, caption="ср. {ref:схемма}")
    p = one(validate(template, values, manifest), "unresolved_ref")
    assert (p.key, p.got) == ("схема", "схемма") and "схема" in p.message
    d = built(template, {"схема": values["схема"]})
    assert [t for t in map(ptext, d.paragraphs) if t.startswith("Рисунок")] == ["Рисунок 1 — ср. ?"]


def test_ссылка_в_подписи_таблицы(template, values, manifest, png):
    values["замеры"] = Table(rows=[["t"], ["1"]], caption="см. {ref:нетути}")
    assert one(validate(template, values, manifest), "unresolved_ref").got == "нетути"
    d = built(template, {"замеры": values["замеры"]})
    assert [t for t in map(ptext, d.paragraphs) if t.startswith("Таблица")] == ["Таблица 1 — см. ?"]


def test_ссылка_в_подписи_схемы(template, values, manifest):
    """У схемы подпись идёт тем же путём (`render._emit_image`); сборка тут не зовётся —
    для неё нужен drawio CLI."""
    values["схема"] = Diagram(xml="<mxfile/>", caption="ср. {ref:нетути}")
    assert one(validate(template, values, manifest), "unresolved_ref").got == "нетути"


def test_подпись_на_тег_шаблона_законна(template, values, manifest, png):
    manifest.tags["схема"].type = "image"
    values["схема"] = Image(source=png, caption="ср. с {ref:замеры}")
    assert validate(template, values, manifest) == []


def test_подпись_без_текста_ссылок_не_содержит(template, values, manifest, png):
    """`caption` бывает не строкой: None — номер без текста, False — вовсе без подписи."""
    manifest.tags["схема"].type = "image"
    for cap in (None, False):
        values["схема"] = Image(source=png, caption=cap)
        assert validate(template, values, manifest) == []


def test_ссылка_в_листинге_не_ссылка(template, values, manifest):
    """В листинге разметка не разбирается: `{ref:}` там — просто текст программы."""
    manifest.tags["вывод"].type = "code"
    values["вывод"] = Code("printf(\"{ref:x}\");", lang="c")
    assert validate(template, values, manifest) == []


# ── прочее ────────────────────────────────────────────────────────────────────

def test_формула_и_текст_тоже_знают_свой_тип(template, values, manifest):
    manifest.tags["вывод"].type = "formula"
    values["вывод"] = Formula(latex="E = mc^2")
    assert validate(template, values, manifest) == []


def test_ключи_значений_в_nfc():
    """Word и macOS пишут «й» разложенным (NFD): без нормализации значение не нашло бы
    своего тега и стало бы сразу двумя проблемами — «тега нет» и «значения нет»."""
    tpl = docx_bytes(lambda d: d.add_paragraph("{{случай:Случай}}"))
    m = manifest_from_template(tpl)
    assert validate(tpl, {"случаи\u0306": Markdown("тот самый")}, m) == []


def test_validate_видна_из_пакета():
    """Проверку зовут снаружи (служба, интерфейс, повтор у модели) — она часть шва."""
    import hokoku
    assert hokoku.validate is validate
