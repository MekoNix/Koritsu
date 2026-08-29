from docx import Document

from hokoku import extract_tags
from hokoku.tags import parse_tag, TAG_RE
from .conftest import texts


def test_parse_forms():
    assert parse_tag("{{фио}}") == ("фио", "фио")
    assert parse_tag("{{ global_фио : ФИО студента }}") == ("global_фио", "ФИО студента")
    assert parse_tag("{{a.b-c_1}}") == ("a.b-c_1", "a.b-c_1")
    assert parse_tag("{{с пробелом}}") is None
    assert parse_tag("{{#if x}}") is None and parse_tag("{{x|def}}") is None   # зарезервировано


def test_extract_everywhere(template):
    def build(d):
        d.sections[0].header.paragraphs[0].text = "{{шапка}}"
        d.sections[0].footer.paragraphs[0].text = "стр. {{подвал:Подвал}}"
        p = d.add_paragraph()
        p.add_run("Студент {{").bold = True
        p.add_run("фио")
        p.add_run("}} группа {{группа}}")
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).add_paragraph("{{ячейка}}")
        inner = t.cell(0, 0).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "{{вложенная}}"
        d.add_paragraph("{{фио}} ещё раз")
    path = template(build)
    tags = extract_tags(path)
    assert [(t.key, t.where) for t in tags] == [
        ("фио", "body"), ("группа", "body"), ("ячейка", "table"), ("вложенная", "table"),
        ("шапка", "header"), ("подвал", "footer")]
    assert next(t for t in tags if t.key == "подвал").label == "Подвал"


def test_extract_textbox(template):
    def build(d):
        p = d.add_paragraph()
        r = p.add_run()
        r._r.append(_textbox_xml("Тема: {{тема}}"))
    path = template(build)
    assert [(t.key, t.where) for t in extract_tags(path)] == [("тема", "textbox")]


def _textbox_xml(text):
    from docx.oxml import parse_xml
    return parse_xml(
        '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><v:shape><v:textbox><w:txbxContent>'
        f'<w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>')


def _alt_content_xml(text):
    """Текстовое поле, как его пишет Word: та же копия в mc:Choice и в mc:Fallback."""
    from docx.oxml import parse_xml
    box = f'<w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>'
    return parse_xml(
        '<mc:AlternateContent xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml">'
        f'<mc:Choice Requires="wps"><w:drawing><wps:txbx>{box}</wps:txbx></w:drawing></mc:Choice>'
        f'<mc:Fallback><w:pict><v:shape><v:textbox>{box}'
        '</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent>')


def test_alternate_content_counted_once(template, tmp_path, png):
    """Копия в mc:Fallback пропускается: иначе один тег даёт два рисунка и сбитую нумерацию."""
    from hokoku import Image, render
    path = template(lambda d: d.add_paragraph().add_run()._r.append(_alt_content_xml("{{схема}}")))
    assert [(t.key, t.count) for t in extract_tags(path)] == [("схема", 1)]
    res = render(path, {"схема": Image(png, caption="схема")}, str(tmp_path / "o.docx"))
    assert res.figures == 1 and res.refs == {"схема": 1}


def test_tags_found_in_long_document(template, tmp_path):
    """Абзацы без тегов пропускаются одним XPath — проверяем, что ничего не потерялось:
    тег в конце длинного документа, в таблице, в колонтитуле, в текстовом поле и
    разорванный форматированием на три run'а."""
    from hokoku import Markdown, render
    from .conftest import ptext

    def build(d):
        for i in range(400):
            d.add_paragraph(f"обычный абзац {i}")
        p = d.add_paragraph()
        for part in ("{{раз", "орван", "ный}}"):            # Word рвёт тег на runs
            p.add_run(part)
        d.add_table(rows=1, cols=1).cell(0, 0).text = "{{в_таблице}}"
        d.sections[0].header.paragraphs[0].text = "{{в_шапке}}"
        d.add_paragraph().add_run()._r.append(_textbox_xml("{{в_поле}}"))
        d.add_paragraph("последний {{в_конце}}")

    path = template(build)
    keys = {t.key for t in extract_tags(path)}
    assert keys == {"разорванный", "в_таблице", "в_шапке", "в_поле", "в_конце"}
    out = str(tmp_path / "o.docx")
    res = render(path, {k: Markdown(f"значение {k}") for k in keys}, out)
    assert res.unfilled == [] and res.errors == []
    from docx import Document
    d = Document(out)
    full = "\n".join(ptext(p) for p in d.paragraphs)
    full += "\n" + "\n".join(ptext(p) for p in d.sections[0].header.paragraphs)
    full += "\n" + "\n".join(ptext(p) for p in d.tables[0].cell(0, 0).paragraphs)
    for k in ("разорванный", "в_таблице", "в_шапке", "в_конце"):
        assert f"значение {k}" in full, k
    assert "{{" not in full
