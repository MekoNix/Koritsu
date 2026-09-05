import pytest
from docx import Document
from docx.oxml.ns import qn

from hokoku import Blocks, Code, Image, Markdown, Table, Text, HokokuError, render
from .conftest import ptext, texts
from .test_tags import _textbox_xml


def _runs(path, idx):
    return [(r.text, bool(r.bold), bool(r.italic)) for r in Document(path).paragraphs[idx].runs]


def test_inline_keeps_neighbour_formatting(template, tmp_path):
    def build(d):
        p = d.add_paragraph()
        p.add_run("Студент: ")
        p.add_run("{{фио:ФИО}}").bold = True
        p.add_run(", группа ")
        p.add_run("{{группа}}").italic = True
    out = str(tmp_path / "out.docx")
    res = render(template(build), {"фио": "Иванов И. И.", "группа": Text("ИУ7-31")}, out)
    assert texts(out) == ["Студент: Иванов И. И., группа ИУ7-31"]
    assert _runs(out, 0) == [("Студент: ", False, False), ("Иванов И. И.", True, False),
                             (", группа ", False, False), ("ИУ7-31", False, True)]
    assert res.unfilled == []


def test_tag_split_across_runs_and_two_tags(template, tmp_path):
    def build(d):
        p = d.add_paragraph()
        p.add_run("Преп: {{")
        p.add_run("global_")
        p.add_run("преп}} и {{второй}}!")
    out = str(tmp_path / "out.docx")
    render(template(build), {"global_преп": "Петров", "второй": "второй"}, out)
    assert texts(out) == ["Преп: Петров и второй!"]


def test_unfilled_removed_and_reported(template, tmp_path):
    out = str(tmp_path / "out.docx")
    res = render(template(lambda d: d.add_paragraph("Цель: {{цель}} и {{нет}}.")), {"цель": "x"}, out)
    assert texts(out) == ["Цель: x и ."] and res.unfilled == ["нет"]


def test_same_tag_everywhere(template, tmp_path):
    def build(d):
        d.sections[0].header.paragraphs[0].text = "Тема: {{тема}}"
        d.add_paragraph("{{тема}}")
        d.add_table(rows=1, cols=1).cell(0, 0).text = "{{тема}}"
        d.add_paragraph().add_run()._r.append(_textbox_xml("{{тема}}"))
    out = str(tmp_path / "out.docx")
    render(template(build), {"тема": "Сортировка"}, out)
    d = Document(out)
    assert d.sections[0].header.paragraphs[0].text == "Тема: Сортировка"
    assert ptext(d.paragraphs[0]) == "Сортировка"
    assert d.tables[0].cell(0, 0).text == "Сортировка"
    tx = d.element.body.find(".//" + qn("w:txbxContent"))
    assert "".join(t.text for t in tx.iter(qn("w:t"))) == "Сортировка"


def test_text_newline_becomes_break(template, tmp_path):
    out = str(tmp_path / "out.docx")
    render(template(lambda d: d.add_paragraph("{{a}}")), {"a": Text("строка 1\nстрока 2")}, out)
    p = Document(out).paragraphs[0]
    assert p.text == "строка 1\nстрока 2" and p._p.find(".//" + qn("w:br")) is not None


def test_markdown_blocks(template, tmp_path):
    md = Markdown("""# Заголовок
Текст с **жирным**, *курсивом*, `кодом` и [ссылкой](https://example.com).

1. раз
2. два
   - вложенный
3. три

> цитата

| a | b |
|---|--:|
| 1 | **2** |

```python
print(1)
```
---
конец""")
    out = str(tmp_path / "out.docx")
    res = render(template(lambda d: d.add_paragraph("{{md}}")), {"md": md}, out)
    d = Document(out)
    styles = [(ptext(p), p.style.name) for p in d.paragraphs]
    assert styles[0] == ("Заголовок", "Heading 1")
    body = d.paragraphs[1]
    assert [(r.text, bool(r.bold), bool(r.italic)) for r in body.runs][:4] == [
        ("Текст с ", False, False), ("жирным", True, False), (", ", False, False), ("курсивом", False, True)]
    assert body._p.find(".//" + qn("w:hyperlink")) is not None
    assert "Courier New" in [r.font.name for r in body.runs]
    lists = [p for p in d.paragraphs if p.style.name == "List Paragraph"]
    assert [ptext(p) for p in lists] == ["раз", "два", "вложенный", "три"]
    num_ids = [p._p.find(".//" + qn("w:numId")).get(qn("w:val")) for p in lists]
    assert num_ids[0] == num_ids[1] == num_ids[3] != num_ids[2]      # внешняя нумерация продолжается
    assert [ptext(p) for p in d.paragraphs if p.style.name == "Quote"] == ["цитата"]
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).paragraphs[0].runs[0].bold
    assert d.tables[0].cell(0, 0).paragraphs[0].runs[0].bold        # шапка жирная
    assert res.tables == 1
    assert ("print(1)", "Normal") in styles or ("print(1)", "Code") in styles
    assert styles[-1][0] == "конец"
    # абзац с тегом удалён — первый абзац теперь заголовок
    assert "{{" not in "".join(ptext(p) for p in d.paragraphs)


def test_code_keeps_indent_and_angle_brackets(template, tmp_path):
    out = str(tmp_path / "out.docx")
    render(template(lambda d: d.add_paragraph("{{c}}")),
           {"c": Code("#include <iostream>\nint main() {\n    return 0;\n}\n", "cpp")}, out)
    ps = Document(out).paragraphs
    assert [ptext(p) for p in ps] == ["#include <iostream>", "int main() {", "    return 0;", "}"]
    assert all(r.font.name == "Courier New" for p in ps for r in p.runs)


def test_image_with_caption_numbering(template, tmp_path, png):
    out = str(tmp_path / "out.docx")
    res = render(template(lambda d: (d.add_paragraph("{{a}}"), d.add_paragraph("{{b}}"))),
                 {"a": Image(png, caption="Первая"), "b": Image(png)}, out)
    d = Document(out)
    caps = [ptext(p) for p in d.paragraphs if p.style.name == "Caption"]
    assert caps == ["Рисунок 1 — Первая", "Рисунок 2"] and res.figures == 2
    assert len(d.inline_shapes) == 2


def test_tall_image_fits_page(template, tmp_path, tall_png):
    """Высокая картинка вписывается в страницу целиком: резать растр hokoku не умеет —
    листы даёт генератор схемы многостраничным mxfile."""
    from docx.shared import Emu
    out = str(tmp_path / "out.docx")
    res = render(template(lambda d: d.add_paragraph("{{a}}")),
                 {"a": Image(tall_png, caption="Схема")}, out)
    d = Document(out)
    caps = [ptext(p) for p in d.paragraphs if p.style.name == "Caption"]
    assert caps == ["Рисунок 1 — Схема"] and res.figures == 1
    assert len(d.inline_shapes) == 1
    sec = d.sections[0]
    page_h = sec.page_height - sec.top_margin - sec.bottom_margin
    assert d.inline_shapes[0].height <= page_h


def test_markdown_image_from_images_dir(template, tmp_path, png):
    (tmp_path / "img.png").write_bytes(png)
    out = str(tmp_path / "out.docx")
    render(template(lambda d: d.add_paragraph("{{a}}")),
           {"a": Markdown("текст\n\n![Подпись](img.png)")}, out, images_dir=str(tmp_path))
    assert "Рисунок 1 — Подпись" in texts(out)
    with pytest.raises(HokokuError):
        render(template(lambda d: d.add_paragraph("{{a}}")),
               {"a": Markdown("![x](../../etc/passwd)")}, out, images_dir=str(tmp_path))


def test_broken_image_raises(template, tmp_path):
    with pytest.raises(HokokuError):
        render(template(lambda d: d.add_paragraph("{{a}}")), {"a": Image(b"not a png")}, str(tmp_path / "o.docx"))
    with pytest.raises(HokokuError):
        render(template(lambda d: d.add_paragraph("{{a}}")), {"a": Image("/nonexistent.png")}, str(tmp_path / "o.docx"))


def test_table_value_and_blocks_in_cell(template, tmp_path, png):
    def build(d):
        d.add_paragraph("{{t}}")
        d.add_table(rows=1, cols=2).cell(0, 1).text = "{{cell}}"
    out = str(tmp_path / "out.docx")
    res = render(template(build), {"t": Table([["a", "b"], ["1", "2"]], caption="Сводка"),
                                   "cell": Blocks([Text("ок"), Image(png, caption="в ячейке")])}, out)
    d = Document(out)
    assert "Таблица 1 — Сводка" in texts(out)
    assert d.tables[0].cell(0, 0).text == "a"
    cell = d.tables[1].cell(0, 1)
    assert [ptext(p) for p in cell.paragraphs][:1] == [""] or "ок" in [ptext(p) for p in cell.paragraphs]
    assert "Рисунок 1 — в ячейке" in [ptext(p) for p in cell.paragraphs]
    assert res.tables == 1 and res.figures == 1


def test_table_in_list_item_keeps_indent(template, tmp_path):
    """Тег в пункте списка: и таблица, и подпись над ней стоят под своим пунктом."""
    from docx.oxml import OxmlElement

    def build(d):
        p = d.add_paragraph("{{t}}")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        p._p.get_or_add_pPr().append(ind)
        d.add_paragraph("{{s}}", style="List Number")            # отступ из стиля списка

    out = str(tmp_path / "out.docx")
    render(template(build), {"t": Table([["a", "b"], ["1", "2"]], caption="Сводка"),
                             "s": Table([["a"], ["1"]], caption="Список")}, out)
    body = Document(out).element.body
    tbls = body.findall(qn("w:tbl"))
    inds = [t.find(qn("w:tblPr") + "/" + qn("w:tblInd")).get(qn("w:w")) for t in tbls]
    assert inds == ["720", "360"]                                 # 360 — из стиля List Number
    caps = [p for p in body.findall(qn("w:p"))
            if "".join(t.text or "" for t in p.iter(qn("w:t"))).startswith("Таблица ")]
    assert [p.find(qn("w:pPr") + "/" + qn("w:ind")).get(qn("w:left")) for p in caps] == ["720", "360"]


def test_template_bytes_and_document_input(template, tmp_path):
    path = template(lambda d: d.add_paragraph("{{a}}"))
    out = str(tmp_path / "out.docx")
    render(open(path, "rb").read(), {"a": "x"}, out)
    assert texts(out) == ["x"]
    render(Document(path), {"a": "y"}, out)
    assert texts(out) == ["y"]


def _strip_numbering(src_path, dst_path):
    import io, re, zipfile
    src = zipfile.ZipFile(src_path)
    with zipfile.ZipFile(dst_path, "w", zipfile.ZIP_DEFLATED) as z:
        for i in src.infolist():
            if i.filename == "word/numbering.xml":
                continue
            data = src.read(i.filename)
            if i.filename in ("[Content_Types].xml", "word/_rels/document.xml.rels"):
                data = re.sub(rb"<(Override|Relationship)[^>]*numbering[^>]*/>", b"", data)
            z.writestr(i, data)


def test_template_without_numbering_part(template, tmp_path):
    path = template(lambda d: d.add_paragraph("{{a}} и {{md}}"))
    bare = str(tmp_path / "bare.docx")
    _strip_numbering(path, bare)
    out = str(tmp_path / "out.docx")
    render(bare, {"a": "x", "md": Markdown("1. раз\n2. два\n- три")}, out)
    d = Document(out)
    assert ptext(d.paragraphs[0]) == "x и "
    lists = [p for p in d.paragraphs if p.style.name == "List Paragraph"]
    assert [ptext(p) for p in lists] == ["раз", "два", "три"]
    assert all(p._p.find(".//" + qn("w:numPr")) is not None for p in lists)   # numbering.xml создан
    import zipfile
    assert "word/numbering.xml" in zipfile.ZipFile(out).namelist()


def test_block_value_in_paragraph_with_text_keeps_sentence(template, tmp_path):
    out = str(tmp_path / "out.docx")
    render(template(lambda d: d.add_paragraph("Цель: {{цель}}.")),
           {"цель": Markdown("изучить **сортировки**\n\n- пузырёк\n- быстрая")}, out)
    d = Document(out)
    assert ptext(d.paragraphs[0]) == "Цель: изучить сортировки."
    assert [r.text for r in d.paragraphs[0].runs if r.bold] == ["сортировки"]
    assert [ptext(p) for p in d.paragraphs if p.style.name == "List Paragraph"] == ["пузырёк", "быстрая"]
    out2 = str(tmp_path / "out2.docx")
    render(template(lambda d: d.add_paragraph("Итог: {{b}}!")),
           {"b": Blocks([Text("всё ок"), Code("x = 1")])}, out2)
    ps = Document(out2).paragraphs
    assert ptext(ps[0]) == "Итог: всё ок!" and ptext(ps[1]) == "x = 1"


def test_table_widths_by_content_and_header_shading(template, tmp_path):
    out = str(tmp_path / "out.docx")
    render(template(lambda d: d.add_paragraph("{{t}}")),
           {"t": Table([["n", "Очень длинное название колонки"], ["1", "x"]])}, out)
    t = Document(out).tables[0]
    assert t.cell(0, 0).width < t.cell(0, 1).width
    assert t.cell(0, 0)._tc.find(".//" + qn("w:shd")) is not None
    assert t.cell(1, 0)._tc.find(".//" + qn("w:shd")) is None


def test_block_value_removes_dangling_punctuation(template, tmp_path):
    out = str(tmp_path / "out.docx")
    render(template(lambda d: d.add_paragraph("Цель: {{цель}}.")),
           {"цель": Markdown("## Заголовок\nтекст")}, out)
    ps = Document(out).paragraphs
    assert ptext(ps[0]) == "Цель: " and ptext(ps[1]) == "Заголовок"


def test_nested_quote_is_a_separate_indented_paragraph(template, tmp_path):
    """«> >» — своя цитата с дополнительным отступом; раньше вложенная строка
    приклеивалась к внешней в один абзац и никакого отступа не получала."""
    out = str(tmp_path / "out.docx")
    render(template(lambda d: d.add_paragraph("{{m}}")),
           {"m": Markdown("> внешняя\n> > вложенная\n> > и дальше\n")}, out)
    qs = [p for p in Document(out).paragraphs if p.style.name == "Quote"]
    assert [ptext(p) for p in qs] == ["внешняя", "вложенная и дальше"]
    assert qs[0]._p.find(".//" + qn("w:ind")) is None
    assert int(qs[1]._p.find(".//" + qn("w:ind")).get(qn("w:left"))) == 720


def _grid(path):
    """Ширины колонок из w:tblGrid и из w:tcW первой строки, в dxa."""
    t = Document(path).element.body.find(".//" + qn("w:tbl"))
    grid = [int(g.get(qn("w:w"))) for g in t.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
    tr = t.findall(qn("w:tr"))[0]
    tcw = [int(tc.find(qn("w:tcPr") + "/" + qn("w:tcW")).get(qn("w:w"))) for tc in tr.findall(qn("w:tc"))]
    return grid, tcw


def test_table_grid_matches_cell_widths(template, tmp_path):
    """w:tblGrid переписывается по посчитанным ширинам: LibreOffice раскладывает колонки
    по сетке, а не по w:tcW, и без этого в PDF колонки выходили равными."""
    out = str(tmp_path / "out.docx")
    render(template(lambda d: d.add_paragraph("{{t}}")),
           {"t": Table([["№", "Наименование параметра", "Ед."],
                        ["1", "Очень длинное наименование, которое надо перенести", "шт."]])}, out)
    grid, tcw = _grid(out)
    assert grid == pytest.approx(tcw, abs=2)
    assert grid[1] > 5 * grid[0] and grid[1] > 5 * grid[2]


def test_short_columns_are_not_squeezed_below_content(template, tmp_path):
    """Ужимается только избыток над самым длинным словом: узкие колонки держат ширину,
    рвётся длинный текст, которому перенос не страшен."""
    out = str(tmp_path / "out.docx")
    long = "Очень длинное описание параметра, которое всё равно переносится по словам"
    render(template(lambda d: d.add_paragraph("{{t}}")),
           {"t": Table([["№", "Идентификатор", "Описание"], ["1", "коэффициент_альфа", long]])}, out)
    t = Document(out).tables[0]
    assert t.cell(0, 0).width.cm > 1.1                      # было 0.79 — продавлено пропорцией
    assert t.cell(0, 1).width.cm > 6.0                      # было 4.4 — слово не влезало
    assert sum(c.width.cm for c in t.rows[0].cells) == pytest.approx(15.24, abs=0.1)


def test_table_in_list_item_keeps_indent(template, tmp_path):
    """Тег в пункте списка: таблица встаёт под пунктом (w:tblInd), а не у левого поля."""
    out = str(tmp_path / "out.docx")
    render(template(lambda d: d.add_paragraph("{{t}}", style="List Bullet")),
           {"t": Table([["a", "b"], ["1", "2"]])}, out)
    tbl = Document(out).element.body.find(".//" + qn("w:tbl"))
    ind = tbl.find(qn("w:tblPr") + "/" + qn("w:tblInd"))
    assert ind is not None and int(ind.get(qn("w:w"))) > 0


def _table_xml(*cells, grid=("1701", "5000", "1943")):
    """Таблица как её пишет Word: колонки заданы сеткой, w:tcW у ячеек нет."""
    from docx.oxml import parse_xml
    tcs = "".join(f"<w:tc><w:p><w:r><w:t>{c}</w:t></w:r></w:p></w:tc>" for c in cells)
    return parse_xml(
        '<w:tbl xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr><w:tblGrid>'
        + "".join(f'<w:gridCol w:w="{w}"/>' for w in grid)
        + f"</w:tblGrid><w:tr>{tcs}</w:tr></w:tbl>")


def test_image_in_cell_without_tcw_uses_grid(template, tmp_path, png):
    """Титульник из Word: у ячеек нет w:tcW — ширину берём из w:tblGrid, иначе логотип
    рисовался в полстраницы и вылезал за таблицу."""
    out = str(tmp_path / "out.docx")
    render(template(lambda d: d.element.body.insert(0, _table_xml("{{лого}}", "МГТУ", "2026"))),
           {"лого": Image(png, caption=False, width_cm=8)}, out)
    assert Document(out).inline_shapes[0].width.cm == pytest.approx(1701 / 567 - 0.5, abs=0.05)


def test_table_outside_list_has_no_indent(template, tmp_path):
    """Обычный абзац — отступа у таблицы нет."""
    out = str(tmp_path / "out.docx")
    render(template(lambda d: d.add_paragraph("{{t}}")), {"t": Table([["a", "b"]])}, out)
    tbl = Document(out).element.body.find(".//" + qn("w:tbl"))
    assert tbl.find(qn("w:tblPr") + "/" + qn("w:tblInd")) is None


# ── разметка чужого бланка ────────────────────────────────────────────────────

def test_комментарий_вырезается_из_документа(template, tmp_path):
    """Комментарий docxtpl печатался в готовом отчёте обычным текстом: его пишут
    для того, кто заполняет бланк, а не для читателя."""
    def build(d):
        d.add_paragraph("{# 4–6 предложений: итог квартала #}")
        d.add_paragraph("Цель: {{цель}}")
    out = str(tmp_path / "out.docx")
    render(template(build), {"цель": Text("Разобрать сортировки")}, out)
    assert texts(out) == ["Цель: Разобрать сортировки"]


def test_комментарий_в_строке_с_тегом(template, tmp_path):
    def build(d):
        d.add_paragraph("Цель: {# одним абзацем #}{{цель}}")
    out = str(tmp_path / "out.docx")
    render(template(build), {"цель": Text("Разобрать сортировки")}, out)
    assert texts(out) == ["Цель: Разобрать сортировки"]


def test_комментарий_разорванный_между_runs(template, tmp_path):
    """Word рвёт текст на runs где угодно — по проверке правописания в том числе."""
    def build(d):
        p = d.add_paragraph()
        p.add_run("{# одним ")
        p.add_run("абзацем #}")
        p.add_run("Цель: {{цель}}")
    out = str(tmp_path / "out.docx")
    render(template(build), {"цель": Text("Разобрать")}, out)
    assert texts(out) == ["Цель: Разобрать"]


def test_комментарий_отдельным_абзацем_в_ячейке(template, tmp_path):
    """Единственный абзац ячейки не удаляется: ячейка без абзаца — сломанный DOCX."""
    def build(d):
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "{# только пояснение #}"
        d.add_paragraph("{{цель}}")
    out = str(tmp_path / "out.docx")
    render(template(build), {"цель": Text("Разобрать")}, out)
    doc = Document(out)
    assert ptext(doc.tables[0].cell(0, 0).paragraphs[0]) == ""


def test_конструкция_остаётся_текстом_и_не_ломает_сборку(template, tmp_path):
    """Цикл Jinja пока не поддержан: сборка идёт, конструкция видна как есть."""
    def build(d):
        d.add_paragraph("{% for k in kpis %}")
        d.add_paragraph("{{цель}}")
        d.add_paragraph("{% endfor %}")
    out = str(tmp_path / "out.docx")
    res = render(template(build), {"цель": Text("Разобрать")}, out)
    assert texts(out) == ["{% for k in kpis %}", "Разобрать", "{% endfor %}"]
    assert res.errors == []
