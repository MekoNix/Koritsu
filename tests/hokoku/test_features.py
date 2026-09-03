"""Возможности второй волны: поля, ссылки, подсветка, форматы, безопасность путей."""
import io
import shutil

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from hokoku import (Blocks, Code, Diagram, HokokuError, Image, Markdown, Table, Text, Toc,
                    extract_tags, render)
from .conftest import ptext, texts


def _fields(doc, kind):
    return [f.get(qn("w:instr")) for f in doc.element.body.iter(qn("w:fldSimple")) if kind in f.get(qn("w:instr"))]


def test_tag_inside_hyperlink_and_field(template, tmp_path):
    def build(d):
        p = d.add_paragraph("см. ")
        rid = d.part.relate_to("https://x.y", RT.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "{{ссылка}}"; r.append(t); h.append(r)
        p._p.append(h)
        p2 = d.add_paragraph("Дата: ")
        f = OxmlElement("w:fldSimple"); f.set(qn("w:instr"), " DATE ")
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "{{дата}}"; r.append(t); f.append(r)
        p2._p.append(f)
    path = template(build)
    assert [t.key for t in extract_tags(path)] == ["ссылка", "дата"]
    out = str(tmp_path / "o.docx")
    res = render(path, {"ссылка": "документация", "дата": "01.09"}, out)
    assert res.unfilled == [] and texts(out) == ["см. документация", "Дата: 01.09"]


def test_numbers_bool_and_nfc(template, tmp_path):
    import unicodedata
    out = str(tmp_path / "o.docx")
    render(template(lambda d: d.add_paragraph("n={{n}} f={{f}} b={{b}} й={{й}}")),
           {"n": 5, "f": 2.5, "b": True, unicodedata.normalize("NFD", "й"): "ok"}, out)
    assert texts(out) == ["n=5 f=2.5 b=да й=ok"]


def test_output_bytes_and_filelike(template, tmp_path):
    path = template(lambda d: d.add_paragraph("{{a}}"))
    res = render(path, {"a": "x"})
    assert res.output is None and res.data[:2] == b"PK"
    assert ptext(Document(io.BytesIO(res.data)).paragraphs[0]) == "x"
    buf = io.BytesIO()
    render(path, {"a": "y"}, buf)
    assert ptext(Document(io.BytesIO(buf.getvalue())).paragraphs[0]) == "y"


def test_document_input_not_mutated(template, tmp_path):
    d = Document(template(lambda d: d.add_paragraph("{{a}}")))
    render(d, {"a": "x"}, str(tmp_path / "o.docx"))
    assert d.paragraphs[0].text == "{{a}}"


def test_svg_image(template, tmp_path):
    pytest.importorskip("cairosvg")
    svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="100" height="50"><rect width="100" height="50" fill="red"/></svg>'
    out = str(tmp_path / "o.docx")
    res = render(template(lambda d: d.add_paragraph("{{i}}")), {"i": Image(svg, caption="svg")}, out)
    assert res.figures == 1 and len(Document(out).inline_shapes) == 1


def test_captions_are_seq_fields_with_bookmarks_and_refs(template, tmp_path, png):
    def build(d):
        d.add_paragraph("Схема на {ref:схема}, таблица {ref:таб}, см. также {ref:нет}.")
        d.add_paragraph("{{схема}}")
        d.add_paragraph("{{таб}}")
        d.add_paragraph("{{текст}}")
    out = str(tmp_path / "o.docx")
    res = render(template(build), {
        "схема": Image(png, caption="Схема"),
        "таб": Table([["a"], ["1"]], caption="Т"),
        "текст": Markdown("см. рисунок {ref:схема} и таблицу {ref:таб}"),
    }, out)
    d = Document(out)
    # {ref:} в тексте самого шаблона — тоже поле REF; несуществующее имя видно программно
    assert res.refs == {"схема": 1, "таб": 1} and res.unresolved_refs == ["нет"]
    assert "Схема на 1, таблица 1, см. также ?." in texts(out)
    seq = _fields(d, "SEQ")
    assert any("SEQ Рисунок" in f for f in seq) and any("SEQ Таблица" in f for f in seq)
    names = [b.get(qn("w:name")) for b in d.element.body.iter(qn("w:bookmarkStart"))]
    assert "_Ref_схема" in names and "_Ref_таб" in names
    assert "Рисунок 1 — Схема" in texts(out) and "Таблица 1 — Т" in texts(out)
    assert "см. рисунок 1 и таблицу 1" in texts(out)          # REF-поля с кэшем номеров
    refs = _fields(d, "REF")
    assert any("_Ref_схема" in f for f in refs)


def test_diagram_pages_are_sheets_of_one_figure(template, tmp_path, png, monkeypatch):
    """Многостраничный mxfile — один рисунок в несколько листов: номер и закладка общие."""
    calls = []

    def fake(xml, page=None, **kw):
        calls.append(page)
        return png

    import sys
    monkeypatch.setattr(sys.modules["hokoku.render"], "drawio_to_png", fake)   # hokoku.render — функция
    xml = "<mxfile>" + "".join(f'<diagram name="Л{i}"/>' for i in range(1, 4)) + "</mxfile>"
    out = str(tmp_path / "o.docx")
    res = render(template(lambda d: d.add_paragraph("{{схема}}")),
                 {"схема": Diagram(xml, caption="Алгоритм")}, out)
    assert calls == [1, 2, 3]                                   # запрошены все листы
    assert res.figures == 1 and res.refs == {"схема": 1}        # номер один на всю схему
    caps = [t for t in texts(out) if t.startswith("Рисунок")]
    assert caps == ["Рисунок 1 — Алгоритм (лист 1 из 3)",
                    "Рисунок 1 — Алгоритм (лист 2 из 3)",
                    "Рисунок 1 — Алгоритм (лист 3 из 3)"]
    d = Document(out)
    seq = _fields(d, "SEQ Рисунок")
    assert len(seq) == 3 and "\\c" not in seq[0] and all("\\c" in f for f in seq[1:])
    names = [b.get(qn("w:name")) for b in d.element.body.iter(qn("w:bookmarkStart"))]
    assert names.count("_Ref_схема") == 1                       # закладка одна, ссылка не двоится
    assert len(d.inline_shapes) == 3


def test_diagram_single_page_and_bad_page(template, tmp_path, png, monkeypatch):
    import sys
    monkeypatch.setattr(sys.modules["hokoku.render"], "drawio_to_png", lambda xml, page=None, **kw: png)
    xml = '<mxfile><diagram name="A"/><diagram name="B"/></mxfile>'
    out = str(tmp_path / "o.docx")
    res = render(template(lambda d: d.add_paragraph("{{схема}}")),
                 {"схема": Diagram(xml, caption="Один лист", page=2)}, out)
    assert res.figures == 1
    assert [t for t in texts(out) if t.startswith("Рисунок")] == ["Рисунок 1 — Один лист"]
    from hokoku.images import count_pages, drawio_to_png
    assert count_pages(xml) == 2 and count_pages("<mxfile><diagram/></mxfile>") == 1
    with pytest.raises(ValueError, match="запрошена 5"):        # CLI молча отдал бы последнюю
        drawio_to_png(xml, 5)


def test_captions_without_fields_style_override(template, tmp_path, png):
    out = str(tmp_path / "o.docx")
    render(template(lambda d: d.add_paragraph("{{a}}")), {"a": Image(png, caption="X")}, out,
           style={"captions": {"fields": False, "figure": "Рис. {n} – {caption}"}})
    d = Document(out)
    assert _fields(d, "SEQ") == [] and "Рис. 1 – X" in texts(out)
    with pytest.raises(ValueError):
        render(template(lambda d: d.add_paragraph("{{a}}")), {"a": "x"}, out, style={"nope": {}})


def test_intro_paragraph_keeps_with_picture(template, tmp_path, png):
    def build(d):
        d.add_paragraph("Блок-схема алгоритма:")
        d.add_paragraph("{{схема}}")
        d.add_paragraph("Просто текст")
        d.add_paragraph("{{ещё}}")
    out = str(tmp_path / "o.docx")
    render(template(build), {"схема": Image(png), "ещё": Image(png)}, out)
    ps = Document(out).paragraphs
    assert ptext(ps[0]) == "Блок-схема алгоритма:" and ps[0]._p.find(".//" + qn("w:keepNext")) is not None
    after = next(p for p in ps if ptext(p) == "Просто текст")
    assert after._p.find(".//" + qn("w:keepNext")) is None
    pic = next(p for p in ps if p._p.find(".//" + qn("w:drawing")) is not None)
    assert pic._p.find(".//" + qn("w:keepNext")) is not None             # картинка прилипает к подписи


def test_code_highlight_and_line_numbers(template, tmp_path):
    pytest.importorskip("pygments")
    out = str(tmp_path / "o.docx")
    render(template(lambda d: d.add_paragraph("{{c}}")),
           {"c": Code("def f(x):\n    return x  # c\n", "python", line_numbers=True)}, out)
    ps = Document(out).paragraphs
    assert [ptext(p) for p in ps] == ["1  def f(x):", "2      return x  # c"]
    colors = {r.font.color.rgb for p in ps for r in p.runs if r.font.color and r.font.color.rgb}
    assert len(colors) >= 2                                           # серые номера + подсветка
    render(template(lambda d: d.add_paragraph("{{c}}")), {"c": Code("x", "python", highlight=False)}, out)
    assert all(r.font.color.rgb is None for p in Document(out).paragraphs for r in p.runs)


def test_markdown_tasks_nested_quote_image_attrs(template, tmp_path, png):
    (tmp_path / "i.png").write_bytes(png)
    out = str(tmp_path / "o.docx")
    render(template(lambda d: d.add_paragraph("{{m}}")),
           {"m": Markdown("- [ ] сделать\n- [x] готово\n\n> > глубоко\n\n![Подпись](i.png){width=3cm align=left}")},
           out, images_dir=str(tmp_path))
    d = Document(out)
    t = texts(out)
    assert "☐ сделать" in t and "☑ готово" in t and "глубоко" in t and "> глубоко" not in t
    shp = d.inline_shapes[0]
    assert abs(shp.width.cm - 3.0) < 0.05
    pic = next(p for p in d.paragraphs if p._p.find(".//" + qn("w:drawing")) is not None)
    assert pic._p.find(".//" + qn("w:jc")).get(qn("w:val")) == "left"


def test_value_with_braces_is_not_expanded(template, tmp_path):
    out = str(tmp_path / "o.docx")
    res = render(template(lambda d: d.add_paragraph("{{a}} {{b}}")), {"a": "x{{b}}y", "b": "B"}, out)
    assert texts(out) == ["x{{b}}y B"] and res.unfilled == []


def test_strict_paths(template, tmp_path, png):
    (tmp_path / "ok.png").write_bytes(png)
    out = str(tmp_path / "o.docx")
    render(template(lambda d: d.add_paragraph("{{i}}")), {"i": Image("ok.png")}, out,
           images_dir=str(tmp_path), strict_paths=True)
    with pytest.raises(HokokuError):
        render(template(lambda d: d.add_paragraph("{{i}}")), {"i": Image("../../etc/passwd")}, out,
               images_dir=str(tmp_path), strict_paths=True)
    with pytest.raises(HokokuError):
        render(template(lambda d: d.add_paragraph("{{i}}")), {"i": Image("ok.png")}, out, strict_paths=True)


def test_extract_tags_places_and_count(template):
    def build(d):
        d.sections[0].header.paragraphs[0].text = "{{фио}}"
        d.add_paragraph("{{фио:ФИО}} и {{фио}}")
    tags = extract_tags(template(build))
    assert len(tags) == 1 and tags[0].count == 3 and tags[0].label == "ФИО"
    assert set(tags[0].places) == {"body", "header"}


def test_numbering_reuses_abstract_num(template, tmp_path):
    out = str(tmp_path / "o.docx")
    render(template(lambda d: d.add_paragraph("{{m}}")),
           {"m": Markdown("\n\n".join(f"- a{i}\n- b{i}\n\nтекст {i}" for i in range(20)) + "\n\n1. x\n2. y")}, out)
    base = len(Document().part.numbering_part.element.findall(qn("w:abstractNum")))
    numbering = Document(out).part.numbering_part.element
    assert len(numbering.findall(qn("w:abstractNum"))) - base == 2   # один на вид списка (ul, ol)
    assert len(numbering.findall(qn("w:num"))) >= 21                  # но свой w:num на каждый список


@pytest.mark.skipif(shutil.which("drawio") is None, reason="нет drawio CLI")
def test_diagram_from_drawio_xml(template, tmp_path):
    xml = ('<mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/>'
           '<mxCell id="2" value="A" style="rounded=1" vertex="1" parent="1">'
           '<mxGeometry x="20" y="20" width="80" height="40" as="geometry"/></mxCell></root></mxGraphModel>')
    out = str(tmp_path / "o.docx")
    res = render(template(lambda d: d.add_paragraph("{{d}}")), {"d": Diagram(xml, caption="Схема")}, out)
    assert res.figures == 1 and "Рисунок 1 — Схема" in texts(out)


def test_diagram_without_cli_raises(template, tmp_path, monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda name: None)
    with pytest.raises(HokokuError):
        render(template(lambda d: d.add_paragraph("{{d}}")), {"d": Diagram("<x/>")}, str(tmp_path / "o.docx"))


def test_formula_value_and_markdown_math(template, tmp_path):
    M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    from hokoku import Formula
    out = str(tmp_path / "o.docx")
    res = render(template(lambda d: (d.add_paragraph("{{f}}"), d.add_paragraph("{{m}}"))), {
        "f": Formula(r"E = m c^2", ref="энергия"),
        "m": Markdown("Энергия $E$ по формуле {ref:энергия}; и ещё:\n\n$$ \\frac{a}{b} = \\sqrt{x} $$\n\nконец"),
    }, out)
    d = Document(out)
    body = d.element.body
    assert len(list(body.iter(M + "oMath"))) == 3                    # формула, $E$, $$…$$
    assert res.formulas == 2 and res.refs["энергия"] == 1
    assert any("SEQ Формула" in f for f in _fields(d, "SEQ"))
    t = texts(out)
    assert any(x.strip().startswith("(") and x.strip().endswith(")") for x in t)   # «(1)» у формулы
    assert "Энергия  по формуле 1; и ещё:" in t
    assert d.settings.element.find(qn("w:updateFields")) is not None


def test_broken_formula_is_an_error_not_silence(template, tmp_path):
    """Обрывок формулы раньше собирался молча — в отчёт уходили пустые скобки «(1)»,
    и человек узнавал об этом, только открыв документ."""
    from hokoku import Formula
    out = str(tmp_path / "o.docx")
    for latex in (r"\frac{1}{", r"\sqrt", r"x^{", r"\begin{matrix} a"):
        with pytest.raises(HokokuError, match="формула не разобрана"):
            render(template(lambda d: d.add_paragraph("{{f}}")), {"f": Formula(latex)}, out)
    res = render(template(lambda d: (d.add_paragraph("{{f}}"), d.add_paragraph("{{ок}}"))),
                 {"f": Formula(r"\frac{1}{"), "ок": Markdown("остальное на месте")},
                 out, on_error="skip")
    assert [e["key"] for e in res.errors] == ["f"] and "остальное на месте" in "\n".join(texts(out))
    render(template(lambda d: d.add_paragraph("{{f}}")),      # неизвестная команда — по-прежнему текстом
           {"f": Formula(r"\unknowncmd{x}")}, out)


def test_toc(template, tmp_path):
    from hokoku import Toc
    out = str(tmp_path / "o.docx")
    render(template(lambda d: (d.add_paragraph("{{toc}}"), d.add_heading("Раздел", 1))),
           {"toc": Toc(levels=2, title="Содержание")}, out)
    d = Document(out)
    instr = "".join(t.text for t in d.element.body.iter(qn("w:instrText")))
    assert 'TOC \\o "1-2"' in instr
    assert "Содержание" in texts(out)
    assert d.settings.element.find(qn("w:updateFields")).get(qn("w:val")) == "true"
    # результат поля пуст: LibreOffice поле TOC не разворачивает и печатает кэш
    # как обычный текст — служебная подсказка уезжала в сданный PDF
    toc_p = next(p for p in d.element.body.iter(qn("w:p")) if p.find(qn("w:r") + "/" + qn("w:instrText")) is not None)
    assert "".join(t.text or "" for t in toc_p.iter(qn("w:t"))).strip() == ""


def _pdf_text(path: str) -> str | None:
    """Текст PDF: pdftotext, иначе pymupdf; None — нечем прочитать."""
    import subprocess
    if shutil.which("pdftotext"):
        return subprocess.run(["pdftotext", path, "-"], capture_output=True, text=True).stdout
    try:
        import pymupdf
    except ImportError:
        return None
    with pymupdf.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


def test_toc_no_placeholder_in_pdf(template, tmp_path):
    """В PDF на месте оглавления не должно быть служебной подсказки об обновлении полей."""
    from hokoku import Toc, docx_to_pdf
    from hokoku.pdf import libreoffice_available
    if not libreoffice_available():
        pytest.skip("нет LibreOffice")
    out = str(tmp_path / "o.docx")
    render(template(lambda d: (d.add_paragraph("{{toc}}"), d.add_heading("Раздел", 1),
                               d.add_paragraph("текст"))),
           {"toc": Toc(levels=2, title="Содержание")}, out)
    pdf = docx_to_pdf(out, str(tmp_path / "o.pdf"))
    text = _pdf_text(pdf)
    if text is None:
        pytest.skip("нечем прочитать PDF (нет pdftotext и pymupdf)")
    assert "Содержание" in text and "обновится" not in text and "F9" not in text


def test_drawio_timeout_from_render(template, tmp_path, png, monkeypatch):
    """render(drawio_timeout=) доходит до drawio CLI; без него — умолчание images (120 с)."""
    import importlib
    R = importlib.import_module("hokoku.render")     # hokoku.render — это функция, не модуль
    seen = []

    def fake(xml, page=None, **kw):
        seen.append(kw.get("timeout"))
        return png

    monkeypatch.setattr(R, "drawio_to_png", fake)
    tpl = template(lambda d: d.add_paragraph("{{d}}"))
    render(tpl, {"d": Diagram("<mxfile><diagram/></mxfile>")}, str(tmp_path / "a.docx"), drawio_timeout=5)
    render(tpl, {"d": Diagram("<mxfile><diagram/></mxfile>")}, str(tmp_path / "b.docx"))
    assert seen == [5, None]


def test_pdf_timeout_and_expired(tmp_path, monkeypatch):
    """docx_to_pdf(timeout=) уходит в LibreOffice, а его срыв — HokokuError, не TimeoutExpired."""
    import os
    import subprocess

    from hokoku import docx_to_pdf
    from hokoku import pdf as P
    seen = []

    def fake_run(cmd, **kw):
        seen.append(kw.get("timeout"))
        open(os.path.join(cmd[cmd.index("--outdir") + 1], "in.pdf"), "wb").close()
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(P.shutil, "which", lambda name: "/bin/true")
    monkeypatch.setattr(P.subprocess, "run", fake_run)
    src = str(tmp_path / "in.docx")
    open(src, "wb").close()
    docx_to_pdf(src, str(tmp_path / "out.pdf"), timeout=7)
    assert seen == [7]

    def raiser(cmd, **kw):
        raise subprocess.TimeoutExpired(cmd, kw.get("timeout"))

    monkeypatch.setattr(P.subprocess, "run", raiser)
    with pytest.raises(HokokuError, match="не уложился"):
        docx_to_pdf(src, str(tmp_path / "out.pdf"), timeout=7)


def test_caption_false_and_headers_are_not_numbered(template, tmp_path, png):
    """Логотип — не «Рисунок 1»: caption=False убирает подпись, колонтитул не нумеруется."""
    def build(d):
        d.sections[0].header.paragraphs[0].text = "{{лого}}"
        d.add_paragraph("{{без_подписи}}")
        d.add_paragraph("{{с_подписью}}")
        d.add_paragraph("{{таблица}}")
    out = str(tmp_path / "o.docx")
    res = render(template(build), {
        "лого": Image(png),
        "без_подписи": Image(png, caption=False),
        "с_подписью": Image(png, caption="Схема"),
        "таблица": Table([["a"], ["1"]], caption=False),
    }, out)
    d = Document(out)
    caps = [t for t in texts(out) if t.startswith(("Рисунок", "Таблица"))]
    assert caps == ["Рисунок 1 — Схема"]                  # ровно одна подпись на весь документ
    assert res.figures == 1 and res.tables == 0           # номера не потрачены впустую
    assert len(d.inline_shapes) == 2 and len(d.tables) == 1
    hdr = d.sections[0].header
    assert not any("Рисунок" in (t.text or "") for t in hdr._element.iter(qn("w:t")))


def test_on_error_skip_collects_and_keeps_rest(template, tmp_path, png):
    """Одна битая картинка не должна уносить весь отчёт: собрать что можно, беды — списком."""
    def build(d):
        d.add_paragraph("{{a}}")
        d.add_paragraph("{{битая}}")
        d.add_paragraph("{{чужой_тип}}")
        d.add_paragraph("{{b}}")
    vals = {"a": Markdown("первый"), "битая": Image("/нет/такой/картинки.png"),
            "чужой_тип": object(), "b": Markdown("второй")}
    out = str(tmp_path / "o.docx")

    with pytest.raises(HokokuError):                        # по умолчанию — как раньше
        render(template(build), vals, str(tmp_path / "raise.docx"))

    res = render(template(build), vals, out, on_error="skip")
    full = "\n".join(texts(out))
    assert "первый" in full and "второй" in full and "{{" not in full
    assert [e["key"] for e in res.errors] == ["битая", "чужой_тип"]
    assert "не найдена" in res.errors[0]["message"]
    assert res.unfilled == [] and res.figures == 0

    with pytest.raises(ValueError):
        render(template(build), {}, out, on_error="ignore")


def test_two_block_tags_in_one_paragraph_keep_own_refs(template, tmp_path, png):
    """Два блочных тега в одном абзаце: номер и закладка у каждого свои."""
    out = str(tmp_path / "o.docx")
    res = render(template(lambda d: d.add_paragraph("{{первая}} {{вторая}}")),
                 {"первая": Image(png, caption="A"),
                  "вторая": Image(png, caption="B")}, out)
    assert res.refs == {"первая": 1, "вторая": 2}


def test_ref_inside_value_table_gets_number(template, tmp_path, png):
    """«см. {ref:рис}» в ячейке таблицы-значения: ячейки пишет ops.add_table напрямую,
    и поле REF раньше оставалось с кэшем «?», не попадая даже в unresolved_refs."""
    def build(d):
        d.add_paragraph("{{рис}}")
        d.add_paragraph("{{таб}}")
    out = str(tmp_path / "o.docx")
    res = render(template(build), {
        "рис": Image(png, caption="Схема"),
        "таб": Table([["что"], ["см {ref:рис}"]], caption="Т"),
        "нетакой": None,
    }, out)
    cell = Document(out).tables[0].rows[1].cells[0]
    assert ptext(cell.paragraphs[0]) == "см 1"                      # кэш номера, не «?»
    assert res.unresolved_refs == []

    res2 = render(template(build), {"таб": Table([["см {ref:нету}"]], caption="Т")},
                  str(tmp_path / "o2.docx"))
    assert res2.unresolved_refs == ["нету"]                         # несуществующая — видна


def test_ref_inside_caption_gets_number(template, tmp_path, png):
    """«ср. {ref:схема}» в самой подписи: подпись пишется docx_ops.make_run, и ссылка
    оставалась литералом — ни поля REF, ни записи в unresolved_refs."""
    def build(d):
        d.add_paragraph("{{схема}}")
        d.add_paragraph("{{график}}")
        d.add_paragraph("{{таб}}")
    out = str(tmp_path / "o.docx")
    res = render(template(build), {
        "схема":  Image(png, caption="исходная"),
        "график": Image(png, caption="ср. {ref:схема}"),
        "таб":    Table([["a"]], caption="к {ref:нету}"),
    }, out)
    caps = [ptext(p) for p in Document(out).paragraphs if p.style.name == "Caption"]
    assert "Рисунок 2 — ср. 1" in caps and "Таблица 1 — к ?" in caps
    assert res.unresolved_refs == ["нету"]


def test_ref_inside_blocks_gets_number(template, tmp_path, png):
    """Ссылка внутри Blocks — во всех местах, где render пишет видимый текст: абзац Text,
    голая строка, абзац и пункт списка markdown, ячейка таблицы, подпись. Раньше Text
    и строка уезжали в документ буквой «{ref:схема}» — и молча, мимо unresolved_refs."""
    def build(d):
        d.add_paragraph("{{схема}}")
        d.add_paragraph("{{блок}}")
    out = str(tmp_path / "o.docx")
    res = render(template(build), {
        "схема": Image(png, caption="исходная"),
        "блок": Blocks([
            Text("текстом {ref:схема}"),
            "строкой {ref:схема}",
            Markdown("абзацем {ref:схема}\n\n- пунктом {ref:схема}\n"),
            Table([["в ячейке {ref:схема}"]], caption="в подписи {ref:схема}"),
        ]),
    }, out)
    doc = Document(out)
    got = [ptext(p) for p in doc.paragraphs]
    assert "текстом 1" in got and "строкой 1" in got            # кэш номера, не литерал
    assert "абзацем 1" in got and "пунктом 1" in got
    assert "Таблица 1 — в подписи 1" in got
    assert ptext(doc.tables[0].rows[0].cells[0].paragraphs[0]) == "в ячейке 1"
    assert len(_fields(doc, "REF")) == 6 and res.unresolved_refs == []


def test_unresolved_ref_inside_blocks_is_visible(template, tmp_path, png):
    """Ссылка в никуда из Blocks: в документе «?», и имя названо в unresolved_refs —
    молчание здесь хуже «?», потому что беду не видно и программно."""
    def build(d):
        d.add_paragraph("{{схема}}")
        d.add_paragraph("{{блок}}")
    out = str(tmp_path / "o.docx")
    res = render(template(build), {
        "схема": Image(png, caption="исходная"),
        "блок": Blocks([Text("текстом {ref:нетА}"), "строкой {ref:нетБ}",
                        Markdown("абзацем {ref:нетВ}")]),
    }, out)
    assert res.unresolved_refs == ["нетА", "нетБ", "нетВ"]
    got = [ptext(p) for p in Document(out).paragraphs]
    assert "текстом ?" in got and "строкой ?" in got and "абзацем ?" in got


def test_ref_in_blocks_first_item_written_into_tag_paragraph(template, tmp_path, png):
    """Первый абзац Blocks вписывается на место тега («Цель: {{цель}}.»), а не встаёт
    следующим — ссылка в нём проходит тот же путь."""
    def build(d):
        d.add_paragraph("{{схема}}")
        d.add_paragraph("Цель: {{цель}}.")
    out = str(tmp_path / "o.docx")
    res = render(template(build), {
        "схема": Image(png, caption="исходная"),
        "цель": Blocks([Text("повторить {ref:схема}"), Text("и ещё {ref:нету}")]),
    }, out)
    assert "Цель: повторить 1." in [ptext(p) for p in Document(out).paragraphs]
    assert res.unresolved_refs == ["нету"]


def test_ref_in_toc_title_is_field(template, tmp_path, png):
    """Заголовок оглавления — отдельный абзац перед полем TOC, поле REF в нём законно.
    Раньше ссылка оставалась в нём литералом и мимо unresolved_refs."""
    def build(d):
        d.add_paragraph("{{схема}}")
        d.add_paragraph("{{оглав}}")
    out = str(tmp_path / "o.docx")
    res = render(template(build), {
        "схема": Image(png, caption="исходная"),
        "оглав": Toc(title="Оглавление (см. {ref:схема})"),
    }, out)
    doc = Document(out)
    assert "Оглавление (см. 1)" in [ptext(p) for p in doc.paragraphs]
    assert any("_Ref_схема" in f for f in _fields(doc, "REF"))
    assert res.unresolved_refs == []

    res2 = render(template(build), {"оглав": Toc(title="Оглавление (см. {ref:нету})")},
                  str(tmp_path / "o2.docx"))
    assert res2.unresolved_refs == ["нету"]
    assert "Оглавление (см. ?)" in [ptext(p) for p in Document(str(tmp_path / "o2.docx")).paragraphs]


def test_ref_in_code_stays_literal(template, tmp_path, png):
    """`{ref:x}` в листинге — текст программы: render его не трогает намеренно,
    и в unresolved_refs он не попадает (иначе фигурные скобки в коде станут бедой)."""
    def build(d):
        d.add_paragraph("{{схема}}")
        d.add_paragraph("{{код}}")
    out = str(tmp_path / "o.docx")
    res = render(template(build), {"схема": Image(png, caption="исходная"),
                                   "код": Blocks([Code("print('{ref:схема}')")])}, out)
    assert "print('{ref:схема}')" in [ptext(p) for p in Document(out).paragraphs]
    assert res.unresolved_refs == []


def test_skip_survives_alien_exception_and_keeps_blocks_tail(template, tmp_path):
    """skip обязан собрать остальное: и когда python-docx не знает формат картинки,
    и когда битый элемент стоит в середине Blocks."""
    import io as _io
    from PIL import Image as PIL
    buf = _io.BytesIO()
    PIL.new("RGB", (60, 40), "red").save(buf, format="WEBP")        # python-docx WEBP не знает
    out = str(tmp_path / "o.docx")
    res = render(template(lambda d: (d.add_paragraph("{{a}}"), d.add_paragraph("{{b}}"))),
                 {"a": Image(buf.getvalue()),
                  "b": Blocks([Text("до"), Image("/нет/файла.png"), Text("после")])},
                 out, on_error="skip")
    full = "\n".join(texts(out))
    assert "до" in full and "после" in full                         # хвост Blocks на месте
    assert [e["key"] for e in res.errors] == ["a", "b"]
    assert "UnrecognizedImageError" in res.errors[0]["message"]

    with pytest.raises(Exception):                                  # без skip — как раньше
        render(template(lambda d: d.add_paragraph("{{a}}")), {"a": Image(buf.getvalue())},
               str(tmp_path / "raise.docx"))


def test_unknown_value_key_is_reported(template, tmp_path):
    """Опечатка в ключе значения (в том числе у модели) раньше исчезала бесследно:
    в unfilled её нет — там только объявленные теги, в errors тоже."""
    out = str(tmp_path / "o.docx")
    res = render(template(lambda d: d.add_paragraph("{{цель}} {{задачи}}")),
                 {"цель": Text("есть"), "цель_рабты": Markdown("опечатка"), "лишний": Text("x")}, out)
    assert res.unfilled == ["задачи"]                       # тег без значения
    assert res.unknown_keys == ["лишний", "цель_рабты"]     # значение без тега


def test_table_build_is_linear(template, tmp_path):
    """Наполнение таблицы росло квадратично: tbl.cell() каждый раз строил всю матрицу
    ячеек заново, и 400 строк собирались 42 с."""
    import time
    out = str(tmp_path / "o.docx")
    times = {}
    for n in (40, 160):
        rows = [["№", "Параметр", "Значение", "Единица"]] + [[str(i), f"п{i}", "1.0", "с"]
                                                             for i in range(n)]
        t = time.time()
        render(template(lambda d: d.add_paragraph("{{т}}")), {"т": Table(rows)}, out)
        times[n] = time.time() - t
    assert times[160] < times[40] * 8                       # квадрат дал бы ×16 и хуже
    assert times[160] < 3.0


def test_bookmark_ids_are_per_document(template, tmp_path, png):
    """Номер закладки был общим на процесс: вывод не воспроизводился побайтно, а два
    одновременных рендера могли разорвать пару bookmarkStart/bookmarkEnd."""
    def build(d):
        d.add_paragraph("{{схема}}")
        d.add_paragraph("{{ссылка}}")
    vals = {"схема": Image(png, caption="С"), "ссылка": Markdown("см. {ref:схема}")}
    first = render(template(build), vals, None).data
    second = render(template(build), vals, None).data
    import io
    import zipfile
    def doc_xml(data):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return z.read("word/document.xml")
    assert doc_xml(first) == doc_xml(second)                # одинаковый вход — одинаковый выход


def test_broken_formula_does_not_eat_its_number(template, tmp_path):
    """При on_error="skip" битая формула забирала номер себе: следующая получала «3»,
    хотя в документе она вторая, и {ref:} указывал не на ту формулу."""
    from hokoku import Formula
    out = str(tmp_path / "o.docx")
    res = render(template(lambda d: [d.add_paragraph("{{%s}}" % k) for k in ("f1", "f2", "f3")]),
                 {"f1": Formula("a^2"), "f2": Formula(r"\\frac{1}{"), "f3": Formula("b^2")},
                 out, on_error="skip")
    assert res.formulas == 2 and res.refs == {"f1": 1, "f3": 2}
    d = Document(out)
    nums = ["".join(t.text or "" for t in f.iter(qn("w:t")))
            for f in d.element.body.iter(qn("w:fldSimple")) if "SEQ" in (f.get(qn("w:instr")) or "")]
    assert nums == ["1", "2"]                       # номера в документе и в refs совпадают

    with pytest.raises(HokokuError, match="пустая"):
        render(template(lambda d: d.add_paragraph("{{f}}")), {"f": Formula("   ")},
               str(tmp_path / "empty.docx"))


def test_empty_value_is_an_error(template, tmp_path):
    """Пустое значение раньше съедало тег молча: ни в unfilled (там только теги вовсе без
    значения), ни в errors — «модель ничего не вернула» выглядело как «тег заполнен»."""
    from hokoku import Table
    out = str(tmp_path / "o.docx")
    build = lambda d: [d.add_paragraph("{{%s}}" % k) for k in ("a", "b", "c", "d", "есть")]
    res = render(template(build), {"a": "", "b": Text("  "), "c": Markdown(""), "d": Table([]),
                                   "есть": Markdown("текст")}, out, on_error="skip")
    assert [e["key"] for e in res.errors] == ["a", "b", "c", "d"]
    assert res.unfilled == [] and "текст" in "\n".join(texts(out))
    assert "{{" not in "\n".join(texts(out))              # пустые теги всё равно убраны

    with pytest.raises(HokokuError, match="значение пустое"):
        render(template(build), {"a": ""}, str(tmp_path / "raise.docx"))


def test_docx_bytes_to_pdf(tmp_path, monkeypatch):
    """Байты внутрь, байты наружу: DOCX уже лежит в памяти (RenderResult.data), и путь
    через диск — это три лишних действия и три места, где остаётся мусор."""
    import os
    import subprocess

    from hokoku import docx_bytes_to_pdf
    from hokoku import pdf as P
    seen = []

    def fake_run(cmd, **kw):
        docx = cmd[-1]
        seen.append((os.path.basename(docx), open(docx, "rb").read(), kw.get("timeout")))
        outdir = cmd[cmd.index("--outdir") + 1]
        with open(os.path.join(outdir, "report.pdf"), "wb") as f:
            f.write(b"%PDF-1.7 fake")
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(P.shutil, "which", lambda name: "/bin/true")
    monkeypatch.setattr(P.subprocess, "run", fake_run)
    assert docx_bytes_to_pdf(b"docx-bytes", timeout=7) == b"%PDF-1.7 fake"
    assert seen == [("report.docx", b"docx-bytes", 7)]

    def raiser(cmd, **kw):
        raise subprocess.TimeoutExpired(cmd, kw.get("timeout"))

    monkeypatch.setattr(P.subprocess, "run", raiser)
    with pytest.raises(HokokuError, match="не уложился"):
        docx_bytes_to_pdf(b"docx-bytes", timeout=7)
    with pytest.raises(HokokuError, match="нужны байты"):
        docx_bytes_to_pdf(str(tmp_path / "o.docx"))


def test_docx_bytes_to_pdf_real(template, tmp_path):
    """Настоящий прогон через LibreOffice: PDF получается из байтов, мусора не остаётся."""
    import glob
    import os
    import tempfile

    from hokoku import docx_bytes_to_pdf
    from hokoku.pdf import libreoffice_available
    if not libreoffice_available():
        pytest.skip("нет LibreOffice")
    trash = os.path.join(tempfile.gettempdir(), "hokoku_*")
    before = set(glob.glob(trash))
    res = render(template(lambda d: d.add_paragraph("{{a}}")), {"a": "текст"}, None)
    pdf = docx_bytes_to_pdf(res.data, timeout=120)
    assert pdf.startswith(b"%PDF")
    assert set(glob.glob(trash)) == before        # временный каталог убран, DOCX на диск не лёг
