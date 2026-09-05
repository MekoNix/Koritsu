"""
Живой режим: работа это упорядоченный список именованных блоков.

Проверяется не «функции не падают», а четыре обещания живого режима, каждое из
которых стоит дефекта:

  1. операции чистые — прежний список цел, значит отмена есть каждый ход;
  2. документ это сборка списка, и собирает её тот же `render` — подписи, счётчики
     SEQ и поля REF работают, а не «должны бы»;
  3. `validate_work` ловит дубли ключей и ссылки в никуда **до** сборки, потому что
     после неё видно только глазами и только «?» посреди отчёта;
  4. схемы инструментов годятся строгому режиму поставщика: у каждого поля есть тип,
     `additionalProperties` закрыт, объединение помечено `type`/`const`.
"""
import io

import pytest
from docx import Document
from docx.oxml.ns import qn

import hokoku
from hokoku import Code, Diagram, Image, Markdown, PageBreak, Table, Text, Toc, live
from .conftest import ptext


def _texts(data: bytes) -> list[str]:
    return [ptext(p) for p in Document(io.BytesIO(data)).paragraphs]


def _fields(data: bytes, kind: str) -> list[str]:
    doc = Document(io.BytesIO(data))
    return [f.get(qn("w:instr")) for f in doc.element.body.iter(qn("w:fldSimple"))
            if kind in (f.get(qn("w:instr")) or "")]


def собрать(*значения) -> hokoku.Work:
    """Список из значений подряд, ключи — сгенерированные (b-01, b-02, …)."""
    work = hokoku.Work()
    for v in значения:
        work = live.insert(work, live.block(live.new_key(work), v))
    return work


# ── чистота операций и отмена ────────────────────────────────────────────────

def test_вставка_не_трогает_прежний_список():
    """Отмена нужна каждый ход: агент переписал абзац, вышло хуже — вернуть нечем,
    если прежний список изменили на месте."""
    было = собрать(Text("первый"))
    стало = live.insert(было, live.block("b-99", Text("второй")))
    assert было.keys() == ["b-01"] and стало.keys() == ["b-01", "b-99"]
    assert было.blocks[0] is стало.blocks[0]          # общий блок не копируется зря


def test_каждая_операция_возвращает_новый_список():
    было = собрать(Text("раз"), Text("два"), Text("три"))
    снимок = было.keys()
    хвост = [
        live.replace(было, "b-02", Text("другое")),
        live.remove(было, "b-02"),
        live.move(было, "b-01", after="b-03"),
        live.rename(было, "b-02", "итог"),
        live.insert(было, live.block("новый", Text("ещё")), before="b-01"),
    ]
    assert было.keys() == снимок                       # ни одна не тронула исходный
    assert [w.keys() for w in хвост] == [
        ["b-01", "b-02", "b-03"],
        ["b-01", "b-03"],
        ["b-02", "b-03", "b-01"],
        ["b-01", "итог", "b-03"],
        ["новый", "b-01", "b-02", "b-03"],
    ]
    assert хвост[0].get("b-02").text == "другое" and было.get("b-02").text == "два"


def test_место_называется_ключом_соседа_а_не_числом():
    """Индекс сдвинется при первой вставке выше — и сдвинется молча."""
    w = собрать(Text("а"), Text("б"))
    w = live.insert(w, live.block("вставка", Text("между")), after="b-01")
    assert w.keys() == ["b-01", "вставка", "b-02"]
    with pytest.raises(live.LiveError) as e:
        live.insert(w, live.block("x", Text("y")), after="b-01", before="b-02")
    assert e.value.code == "bad_place"
    with pytest.raises(live.LiveError) as e:
        live.move(w, "b-01", after="b-77")
    assert e.value.code == "unknown_block"


def test_ключи_уникальны_и_проверяются_после_нормализации():
    """Пункт status.md: два ключа, различающиеся только формой NFC, схлопнутся в один."""
    import unicodedata
    w = собрать(Text("а"), Text("б"))
    assert live.new_key(w) == "b-03"
    w = live.insert(w, live.block("b-03", Text("в")))
    assert live.new_key(w) == "b-04"                   # занятый номер генератор пропускает

    ключ = "тезисы"
    w2 = live.insert(hokoku.Work(), live.block(unicodedata.normalize("NFD", "тёзисы"), Text("а")))
    assert w2.keys() == [unicodedata.normalize("NFC", "тёзисы")]
    with pytest.raises(live.LiveError) as e:
        live.insert(w2, live.block(unicodedata.normalize("NFD", "тёзисы"), Text("б")))
    assert e.value.code == "duplicate_key"
    assert live.check_key(unicodedata.normalize("NFD", ключ)) == ключ


@pytest.mark.parametrize("плохой", ["", "два слова", "с:двоеточием", "с|чертой",
                                    "с#решёткой", "со/слэшем", "с{скобкой}"])
def test_ключ_обязан_годиться_тегом(плохой):
    """Ключ блока становится тегом синтетического шаблона — значит он обязан пережить
    поездку через `{{ключ}}`, иначе блок молча выпадет из документа."""
    with pytest.raises(live.LiveError) as e:
        live.check_key(плохой)
    assert e.value.code == "bad_key"


def test_переименование_чинит_ссылки_и_ref_значения():
    """Ключ это адрес: голое переименование порвало бы `{ref:}` во всех соседях."""
    w = собрать(Table([["a"], ["1"]], caption="Замеры"),
                Markdown("см. таблицу {ref:b-01} и ещё раз {ref:b-01}"),
                Image(b"", caption="к {ref:b-01}"))
    w = live.rename(w, "b-01", "замеры")
    assert w.keys() == ["замеры", "b-02", "b-03"]
    assert w.get("b-02").text == "см. таблицу {ref:замеры} и ещё раз {ref:замеры}"
    assert w.get("b-03").caption == "к {ref:замеры}"
    # `ref=`, назначенный самим значением, — тот же адрес: если он равен старому ключу,
    # переименование обязано его догнать
    w2 = live.rename(собрать(Image(b"", caption="рис", ref="b-01"),
                             Text("см. {ref:b-01}")), "b-01", "b-08")
    assert w2.get("b-08").value.ref == "b-08" and w2.get("b-02").text == "см. {ref:b-08}"
    # а собственное имя, не совпадающее с ключом, переименование ключа не трогает: адрес
    # там «схема», и ссылки на него ничего не рвут
    w3 = live.rename(собрать(Image(b"", caption="рис", ref="схема"),
                             Text("см. {ref:схема}")), "b-01", "b-08")
    assert w3.get("b-08").ref == "схема" and w3.get("b-02").text == "см. {ref:схема}"


def test_метку_меняют_свободно_а_ключ_остаётся_адресом():
    w = live.insert(hokoku.Work(), live.block("b-01", Text("текст"), label="Черновик"))
    w2 = live.replace(w, "b-01", label="Постановка задачи")
    assert w2.get("b-01").label == "Постановка задачи" and w2.get("b-01").key == "b-01"
    assert w.get("b-01").label == "Черновик"


# ── сборка ────────────────────────────────────────────────────────────────────

def test_сборка_восьми_блоков_всех_видов_держит_порядок(png):
    """Один документ из всех видов сразу: порядок в XML обязан совпасть с порядком
    списка, а номера — считаться по нему же."""
    w = собрать(
        Toc(levels=2, title="Содержание"),
        live.heading(1, "Введение"),
        Markdown("Опора — таблица {ref:b-04} и рисунок {ref:b-05}."),
        Table([["величина", "значение"], ["ток", "5"]], caption="Замеры"),
        Image(png, caption="Стенд"),
        Code("print(1)", lang="python"),
        PageBreak(),
        live.heading(2, "Итог"),
    )
    assert [b.kind for b in w] == ["toc", "heading", "markdown", "table", "image",
                                   "code", "page_break", "heading"]
    assert hokoku.validate_work(w) == []
    res = hokoku.render_work(w)
    строки = _texts(res.data)

    # порядок: содержание → заголовок → текст → подпись таблицы → рисунок → листинг → итог
    # (ячейки самой таблицы в `paragraphs` не попадают — они внутри w:tbl)
    видимые = [t for t in строки if t.strip()]
    assert видимые == ["Содержание", "Введение",
                       "Опора — таблица 1 и рисунок 1.",
                       "Таблица 1 — Замеры",
                       "Рисунок 1 — Стенд", "print(1)", "Итог"]
    ячейки = [c.text for t in Document(io.BytesIO(res.data)).tables for r in t.rows
              for c in r.cells]
    assert ячейки == ["величина", "значение", "ток", "5"]
    assert res.figures == 1 and res.tables == 1 and res.refs == {"b-04": 1, "b-05": 1}
    assert res.unresolved_refs == [] and res.unfilled == [] and res.unknown_keys == []
    assert hokoku.assemble(w)[:2] == b"PK"


def test_номера_поля_seq_а_ссылки_поля_ref(png):
    """Номер — поле SEQ, ссылка — поле REF с закладкой: вставка в середину сдвигает
    номера у всех ниже, и делает это Word, а не мы."""
    w = собрать(Markdown("рисунок {ref:b-02}, таблица {ref:b-03}"),
                Image(png, caption="Схема"),
                Table([["a"], ["1"]], caption="Данные"))
    data = hokoku.assemble(w)
    assert any("SEQ Рисунок" in f for f in _fields(data, "SEQ"))
    assert any("SEQ Таблица" in f for f in _fields(data, "SEQ"))
    assert any("_Ref_b-02" in f for f in _fields(data, "REF"))
    имена = [b.get(qn("w:name")) for b in Document(io.BytesIO(data)).element.body
             .iter(qn("w:bookmarkStart"))]
    assert "_Ref_b-02" in имена and "_Ref_b-03" in имена
    assert "рисунок 1, таблица 1" in _texts(data)

    # вставили рисунок выше — прежний стал вторым, и ссылка на него тоже
    w2 = live.insert(w, live.block("первый", Image(png, caption="Общий вид")), after="b-01")
    res = hokoku.render_work(w2)
    assert res.refs == {"первый": 1, "b-02": 2, "b-03": 1}
    assert "рисунок 2, таблица 1" in _texts(res.data)


def test_заголовки_ставятся_стилем_и_видны_оглавлению():
    """Заголовок — markdown из одной строки; стиль Heading N важен не видом, а тем,
    что по нему заголовок находит поле TOC и «Обновить оглавление» в Word."""
    w = собрать(Toc(levels=3), live.heading(1, "Введение"), Text("текст"),
                live.heading(2, "Постановка"))
    assert live.outline(w) == [{"key": "b-02", "level": 1, "text": "Введение"},
                               {"key": "b-04", "level": 2, "text": "Постановка"}]
    data = hokoku.assemble(w)
    док = Document(io.BytesIO(data))
    стили = [p.style.name for p in док.paragraphs if p.text.strip()]
    assert "Heading 1" in стили and "Heading 2" in стили
    # поле TOC — сложное (w:instrText), а не fldSimple; Word заполнит его при открытии
    инструкции = [t.text for t in док.element.body.iter(qn("w:instrText"))]
    assert any(f' TOC \\o "1-3"' in i for i in инструкции)


def test_сборка_ложится_на_тот_же_путь_что_значения_тегов(png):
    """Второго рисовальщика не заводится: `work_template` плюс `work_values` — это
    обычные шаблон и значения, и `render` о живом режиме не знает ничего."""
    w = собрать(live.heading(1, "Раздел"), Image(png, caption="Схема"))
    tpl = hokoku.work_template(w)
    assert [t.key for t in hokoku.extract_tags(tpl)] == ["b-01", "b-02"]
    свой = hokoku.render(tpl, hokoku.work_values(w), None)
    assert _texts(свой.data) == _texts(hokoku.assemble(w))
    assert свой.figures == 1

    # тот же шаблон и те же значения годятся build_report как обычное задание
    из_задания = hokoku.build_report(
        {"wire_version": hokoku.WIRE_VERSION,
         "template": {"artifact": "tpl"},
         "values": hokoku.values_to_json(hokoku.work_values(w), artifact_of=lambda d: "png")},
        resolve_artifact=lambda a: tpl if a == "tpl" else png,
        workdir=None, store_artifact=lambda n, d, k: "out")
    assert из_задания["ok"] and из_задания["counts"]["figures"] == 1


def test_оформление_образца_доезжает_обоими_путями(tmp_path, png):
    """Профиль правит поля и стили документа, а слова подписей живут в styles.yaml —
    забыть второй путь значит выдать кафедральный отчёт с нашими подписями."""
    from hokoku.sample import StyleProfile
    from hokoku.template import A4_GOST, GOST_BODY
    профиль = StyleProfile(page=A4_GOST, body=GOST_BODY,
                           captions={"figure": "Рис. {n}. {caption}"})
    w = собрать(Image(png, caption="Стенд"))
    assert "Рис. 1. Стенд" in _texts(hokoku.assemble(w, profile=профиль))
    # явная перегрузка сильнее образца
    assert "Fig 1 Стенд" in _texts(hokoku.assemble(
        w, profile=профиль, style={"captions": {"figure": "Fig {n} {caption}"}}))


# ── валидатор на списке ───────────────────────────────────────────────────────

def test_валидатор_ловит_дубли_ключей():
    """Два одинаковых ключа не спорят, а молча схлопываются: второй затрёт первый."""
    w = hokoku.Work((live.block("раздел", Text("первый")),
                     live.block("раздел", Text("второй"))))
    беды = [p.to_dict() for p in hokoku.validate_work(w)]
    assert [p["code"] for p in беды] == ["duplicate_key"]
    assert беды[0]["level"] == "error" and беды[0]["key"] == "раздел"
    with pytest.raises(live.LiveError):
        hokoku.assemble(w)                             # собрать такой список нельзя


def test_валидатор_ловит_ссылки_в_никуда(png):
    """`{ref:}` мимо цели render превращает в «?» прямо в тексте — видно только глазами
    и только на готовом документе."""
    w = собрать(Image(png, caption="Схема"), Markdown("см. {ref:b-01} и {ref:b-09}"))
    беды = hokoku.validate_work(w)
    assert [(p.code, p.key, p.got) for p in беды] == [("unresolved_ref", "b-02", "b-09")]
    # Ошибка, а не предупреждение: цель ссылки в
    # живом списке видна точно, а «?» в готовом документе хуже отказа собрать его.
    assert беды[0].level == "error" and "блока" in беды[0].message
    assert hokoku.unresolved_refs(w) == ["b-09"]
    assert "«?»" in беды[0].message

    # убрали рисунок — ссылка на него повисла, и это видно до сборки
    без = live.remove(w, "b-01")
    assert hokoku.unresolved_refs(без) == ["b-01", "b-09"]


def test_ссылка_на_ненумерованный_блок_это_ошибка_а_не_предупреждение():
    """Номер получают только рисунок, схема, таблица с подписью и нумерованная формула.
    Ссылка на заголовок или на абзац текста даст «?», хотя блок с таким ключом есть.

    Уровень — `error`: «?» в документе хуже отказа,
    и прогон с такой ссылкой до архива доходить не должен. Предупреждением она
    молча уезжала бы в готовую работу.
    """
    w = собрать(live.heading(1, "Введение"), Markdown("см. {ref:b-01}"))
    assert hokoku.unresolved_refs(w) == ["b-01"]
    беды = hokoku.validate_work(w)
    assert [(p.code, p.level, p.got) for p in беды] == [("unresolved_ref", "error", "b-01")]
    assert "?" in _texts(hokoku.assemble(w))[-1]


def test_шаблонный_режим_ссылку_в_никуда_ошибкой_не_считает():
    """Сито одно на оба входа, а уровень разный по делу: в шаблонном режиме
    значения цели во время проверки может не быть вовсе, и «проверить нечем» —
    это предупреждение, а не отказ собирать документ."""
    from hokoku.validate import _refs
    предупреждение = _refs({"тег": Markdown("см. {ref:нету}")}, {})
    assert [(p.code, p.level) for p in предупреждение] == [("unresolved_ref", "warning")]


def test_валидатор_ловит_пустое_и_неизвестный_вид():
    """Пустое значение render считает ошибкой, а «модель ничего не вернула» выглядит
    ровно как «блок написан»."""
    w = hokoku.Work((hokoku.Block("b-01", "text", Text("   ")),
                     hokoku.Block("b-02", "table", Text("не таблица")),
                     hokoku.Block("b-03", "blocks", hokoku.Blocks([Text("а")]))))
    коды = [(p.code, p.key) for p in hokoku.validate_work(w)]
    assert ("empty_value", "b-01") in коды
    assert ("type_mismatch", "b-02") in коды
    assert ("unknown_kind", "b-03") in коды        # blocks в живом списке не бывает
    with pytest.raises(live.LiveError):
        live.kind_of(hokoku.Blocks([]))


def test_валидатор_держит_потолки_службы():
    """Счёт тот же, что у `build_report`: две проверки «слишком много» разошлись бы."""
    w = собрать(Table([["a"] * 3 for _ in range(5)]))
    assert hokoku.validate_work(w) == []
    беды = hokoku.validate_work(w, limits={"max_table_rows": 2})
    assert [p.code for p in беды] == ["limit_exceeded"] and "max_table_rows" in беды[0].message
    много = hokoku.Work(tuple(live.block(f"b-{i}", Text("x")) for i in range(5)))
    assert [p.code for p in hokoku.validate_work(много, limits={"max_values": 3})] \
        == ["limit_exceeded"]


# ── инструменты агента ────────────────────────────────────────────────────────

def _проверить_схему(schema, путь="корень"):
    """Наше подмножество JSON Schema: тип у каждого поля, объект закрыт, `anyOf`
    помечено. Строгий режим поставщика проверит `llm.jsonschema` на стороне
    оркестратора — сюда его не тащим, `hokoku` соседей не импортирует."""
    assert isinstance(schema, dict), путь
    for запрет in ("$ref", "oneOf", "allOf", "not", "$defs", "definitions"):
        assert запрет not in schema, f"{путь}: {запрет}"
    if "anyOf" in schema:
        assert schema["anyOf"], путь
        for i, ветка in enumerate(schema["anyOf"]):
            # помеченное объединение: у каждой ветки дискриминатор type с const
            assert ветка["properties"]["type"].get("const"), f"{путь}.anyOf[{i}]"
            _проверить_схему(ветка, f"{путь}.anyOf[{i}]")
        return
    assert "type" in schema or "enum" in schema or "const" in schema, путь
    if schema.get("type") == "object":
        assert schema.get("additionalProperties") is False, путь
        assert isinstance(schema.get("required"), list), путь
        for имя, поле in (schema.get("properties") or {}).items():
            assert set(schema["required"]) <= set(schema["properties"]), путь
            _проверить_схему(поле, f"{путь}.{имя}")
    if isinstance(schema.get("items"), dict):
        _проверить_схему(schema["items"], f"{путь}.items")


def test_схемы_инструментов_валидны_и_строги():
    инструменты = hokoku.live_tools()
    assert [t.name for t in инструменты] == list(live.TOOL_NAMES)
    for t in инструменты:
        assert t.description.strip() and isinstance(t.schema, dict), t.name
        _проверить_схему(t.schema, t.name)
    # значение блока — то же помеченное объединение, что у `set_tag`, но без `blocks`
    варианты = [в["properties"]["type"]["const"] for в in hokoku.any_block_value_schema()["anyOf"]]
    assert "blocks" not in варианты and set(варианты) == set(hokoku.VALUE_TYPES) - {"blocks"}
    assert set(варианты) | {"heading"} == set(live.KINDS)


def test_инструменты_чистые_и_зовутся_по_имени(png):
    """Оркестратор зовёт по имени и получает новый список: хранение и версии — его дело."""
    w = hokoku.Work()
    w, ответ = hokoku.call_tool(w, "insert_block", {
        "value": {"type": "markdown", "text": "## Постановка задачи"},
        "label": "Постановка"})
    assert ответ == {"key": "b-01", "kind": "heading", "blocks": 1}

    было = w
    w, ответ = hokoku.call_tool(w, "insert_block", {
        "after": "b-01", "value": {"type": "table", "rows": [["a"], ["1"]],
                                   "caption": "Замеры"}})
    assert ответ["key"] == "b-02" and ответ["kind"] == "table"
    assert было.keys() == ["b-01"]                     # прежний список цел

    w, ответ = hokoku.call_tool(w, "insert_block", {
        "before": "b-01", "value": {"type": "toc", "levels": 2}})
    assert w.keys() == ["b-03", "b-01", "b-02"]

    w, ответ = hokoku.call_tool(w, "move_block", {"key": "b-02", "before": "b-01"})
    assert ответ["order"] == ["b-03", "b-02", "b-01"]

    w, ответ = hokoku.call_tool(w, "replace_block", {
        "key": "b-01", "value": {"type": "markdown", "text": "# Постановка задачи"}})
    assert ответ == {"key": "b-01", "kind": "heading"}

    w, ответ = hokoku.call_tool(w, "list_blocks")
    assert [b["key"] for b in ответ["blocks"]] == ["b-03", "b-02", "b-01"]
    assert ответ["blocks"][2] == {"key": "b-01", "kind": "heading", "label": "Постановка",
                                  "preview": "# Постановка задачи", "chars": 19, "level": 1}
    assert ответ["blocks"][1]["caption"] == "Замеры"


def test_удаление_докладывает_повисшие_ссылки(png):
    w = собрать(Image(png, caption="Схема"), Markdown("см. {ref:b-01}"))
    w2, ответ = hokoku.call_tool(w, "remove_block", {"key": "b-01"})
    assert ответ == {"removed": "b-01", "blocks": 1, "unresolved_refs": ["b-01"]}
    assert w.keys() == ["b-01", "b-02"]


def test_беда_инструмента_приходит_с_кодом_и_подсказкой():
    """Модель, получившая внятный отказ, чинится следующим ходом; упавший прогон стоит
    всех уже потраченных денег и не даёт ничего."""
    w = собрать(Text("а"))
    for имя, args, код in (
        ("insert_blok", {}, "unknown_tool"),
        ("remove_block", {"key": "b-77"}, "unknown_block"),
        ("remove_block", {"kye": "b-01"}, "unknown_argument"),
        ("insert_block", {"value": {"type": "текст", "text": "а"}}, "bad_value"),
        ("insert_block", {"value": {"type": "text", "txet": "а"}}, "bad_value"),
        ("move_block", {"key": "b-01", "after": "b-01"}, "bad_place"),
    ):
        with pytest.raises(live.LiveError) as e:
            hokoku.call_tool(w, имя, args)
        assert e.value.code == код, (имя, args)
        assert e.value.payload["error"] == код and e.value.payload["message"]


def test_картинка_приходит_артефактом_а_не_путём(png):
    """Путей нет ни у одного аргумента: имя файла подделывается содержимым, а
    идентификатор — нет."""
    w, ответ = hokoku.call_tool(
        hokoku.Work(), "insert_block",
        {"value": {"type": "image", "artifact": "png-1", "caption": "Стенд"}},
        resolve_artifact=lambda a: png)
    assert ответ["kind"] == "image" and w.get("b-01").caption == "Стенд"
    with pytest.raises(live.LiveError) as e:
        hokoku.call_tool(hokoku.Work(), "insert_block",
                         {"value": {"type": "image", "artifact": "нет"}})
    assert e.value.code == "bad_value"


# ── связный текст одним проходом ──────────────────────────────────────────────

def test_места_под_текст_знают_соседей(png):
    """Скелет петлёй, весь связный текст — одним проходом по готовому списку,
    видя соседей."""
    w = собрать(live.heading(1, "Введение"),
                live.draft("зачем работа"),
                Table([["a"], ["1"]], caption="Замеры"),
                Markdown("уже написано"),
                Text(""))
    места = hokoku.text_slots(w)
    assert [м["key"] for м in места] == ["b-02", "b-05"]
    assert места[0]["hint"] == "зачем работа"
    assert "heading b-01" in места[0]["before"] and "Введение" in места[0]["before"]
    assert "table b-03" in места[0]["after"]
    assert места[1]["after"] == ""                      # последний блок, соседа нет


def test_схема_текстового_прохода_строгая_и_с_соседями():
    w = собрать(live.heading(1, "Введение"), live.draft("зачем работа"))
    схема = hokoku.texts_schema(w)
    _проверить_схему(схема, "texts_schema")
    assert схема["required"] == ["b-02"] and схема["properties"]["b-02"]["type"] == "string"
    assert "перед — heading b-01" in схема["properties"]["b-02"]["description"]
    with pytest.raises(live.LiveError) as e:
        hokoku.texts_schema(собрать(Text("всё написано")))
    assert e.value.code == "nothing_to_write"


def test_текст_вписывается_а_нетекстовые_блоки_не_трогаются(png):
    w = собрать(live.heading(1, "Введение"),
                live.draft("зачем работа"),
                Table([["a"], ["1"]], caption="Замеры"),
                Image(png, caption="Стенд"),
                Text(""))
    w2 = hokoku.fill_texts(w, {"b-02": "Работа о том-то.", "b-05": "Итог такой."})
    assert w2.get("b-02").text == "Работа о том-то." and w2.get("b-02").kind == "markdown"
    assert w2.get("b-05").kind == "text"
    for ключ in ("b-01", "b-03", "b-04"):
        assert w2.get(ключ) is w.get(ключ)              # нетекстовые блоки те же самые
    assert hokoku.text_slots(w2) == [] and hokoku.validate_work(w2) == []
    assert w.get("b-02").text.startswith(live.DRAFT_MARK)   # прежний список цел


def test_текст_в_нетекстовый_блок_это_отказ_а_не_молчание(png):
    """Молча пропущенный ключ хуже отказа: текст, за который заплачено, просто исчез бы,
    а обнаружилось бы это на готовом отчёте по пустому разделу."""
    w = собрать(live.draft(), Table([["a"], ["1"]], caption="Замеры"))
    with pytest.raises(live.LiveError) as e:
        hokoku.fill_texts(w, {"b-02": "текст", "b-77": "ещё", "b-01": "  "})
    assert e.value.code == "bad_texts"
    сообщение = e.value.payload["message"]
    assert "b-02 (table)" in сообщение and "b-77" in сообщение and "b-01" in сообщение
    assert w.get("b-02").kind == "table"                # список не тронут


def test_черновик_собирается_и_показывает_что_в_нём_будет():
    """Буквально пустым место под текст оставить нельзя: пустое значение render считает
    ошибкой, и список не собрался бы даже для показа человеку."""
    w = собрать(live.heading(1, "Введение"), live.draft("зачем работа"))
    assert hokoku.validate_work(w) == []
    assert "черновик: зачем работа" in _texts(hokoku.assemble(w))


# ── описание списка ───────────────────────────────────────────────────────────

def test_описание_списка_не_отдаёт_содержимого_целиком():
    """Модели документ не нужен, а отдать ей документ значит отдать весь отчёт в окно
    на каждом ходу."""
    длинный = "слово " * 200
    w = собрать(Markdown(длинный))
    (запись,) = hokoku.list_blocks(w)
    assert запись["chars"] == len(длинный)
    assert len(запись["preview"]) == live.PREVIEW_CHARS and запись["preview"].endswith("…")


def test_ключ_адрес_а_метка_для_человека():
    """B и C зовут блоки по ключу, а человеку показывают метку."""
    w = live.insert(hokoku.Work(), live.block("b-01", Text("текст"), label="Постановка задачи"))
    (запись,) = hokoku.list_blocks(w)
    assert запись["key"] == "b-01" and запись["label"] == "Постановка задачи"
    assert w.get("b-01").ref == "b-01"                  # ссылаются по ключу, не по метке


def test_подпись_и_ref_живут_на_значении_а_блок_их_только_читает(png):
    """Два места для подписи разошлись бы молча, а в документ уехало бы одно."""
    b = live.block("b-01", Image(png, caption="Стенд", ref="стенд"))
    assert b.caption == "Стенд" and b.ref == "стенд"
    assert live.block("b-02", Text("а")).caption is None
    assert live.block("b-02", Text("а")).ref == "b-02"   # по умолчанию имя ссылки — ключ
    assert "caption" not in {f.name for f in __import__("dataclasses").fields(hokoku.Block)}


def test_схема_рисуется_и_нумеруется_вместе_с_рисунками(png, monkeypatch):
    import sys
    monkeypatch.setattr(sys.modules["hokoku.render"], "drawio_to_png",
                        lambda xml, page=None, **kw: png)
    w = собрать(Image(png, caption="Стенд"),
                Diagram("<mxfile><diagram/></mxfile>", caption="Алгоритм"),
                Markdown("см. {ref:b-02}"))
    res = hokoku.render_work(w)
    assert res.refs == {"b-01": 1, "b-02": 2} and res.figures == 2
    assert "см. 2" in _texts(res.data)


# ── листинги ──────────────────────────────────────────────────────────────────

def test_листинг_нумеруется_и_ссылается():
    """`Code(caption=)` → «Листинг N — …»: подпись над кодом, номер полем SEQ, ссылка
    полем REF. Без подписи (умолчание) листинг номера не получает."""
    w = собрать(Markdown("алгоритм в {ref:b-02}, разбор в {ref:b-03}"),
                Code("def f(): pass", lang="python", caption="Точка входа"),
                Code("x = 1", lang="python", caption="Разбор строки"),
                Code("# просто фрагмент", lang="python"))
    assert set(live.ref_targets(w)) == {"b-02", "b-03"}
    assert hokoku.validate_work(w) == []
    res = hokoku.render_work(w)
    assert res.listings == 2 and res.refs == {"b-02": 1, "b-03": 2}
    строки = _texts(res.data)
    assert "Листинг 1 — Точка входа" in строки and "Листинг 2 — Разбор строки" in строки
    assert "алгоритм в 1, разбор в 2" in строки
    assert any("SEQ Листинг" in f for f in _fields(res.data, "SEQ"))
    имена = [b.get(qn("w:name")) for b in Document(io.BytesIO(res.data)).element.body
             .iter(qn("w:bookmarkStart"))]
    assert имена.count("_Ref_b-02") == 1 and "_Ref_b-04" not in имена

    # вставили листинг выше — номера сдвинулись у всех ниже
    w2 = live.insert(w, live.block("нулевой", Code("import os", caption="Импорты")),
                     after="b-01")
    assert hokoku.render_work(w2).refs == {"нулевой": 1, "b-02": 2, "b-03": 3}


def test_фрагмент_внутри_markdown_номера_не_получает():
    """```-вставка — часть текста, а не листинг: «Листинг 4» над каждым трёхстрочным
    примером сбило бы нумерацию настоящих листингов и все ссылки ниже."""
    w = собрать(Markdown("Так:\n\n```python\nx = 1\n```\n\nи всё."),
                Code("y = 2", caption="Настоящий листинг"))
    res = hokoku.render_work(w)
    assert res.listings == 1 and res.refs == {"b-02": 1}
    assert len([t for t in _texts(res.data) if t.startswith("Листинг")]) == 1


def test_ссылка_в_подписи_листинга_разбирается(png):
    """«Листинг 1 — см. {ref:схема}» — такая же подпись, как у рисунка, и «?» в ней
    стоит на самом виду."""
    w = собрать(Image(png, caption="Схема"), Code("x = 1", caption="к {ref:b-01}"))
    assert hokoku.validate_work(w) == []
    assert "Листинг 1 — к 1" in _texts(hokoku.assemble(w))
    плохо = собрать(Code("x = 1", caption="к {ref:нету}"))
    assert hokoku.unresolved_refs(плохо) == ["нету"]


def test_код_программы_ссылок_не_содержит():
    """`{ref:x}` в тексте программы — текст программы: render его не трогает, и
    переименование блока тоже."""
    w = собрать(Table([["a"], ["1"]], caption="Т"),
                Code('print("{ref:b-01}")', caption="Печать"))
    assert hokoku.unresolved_refs(w) == []
    w2 = live.rename(w, "b-01", "таблица")
    assert w2.get("b-02").value.text == 'print("{ref:b-01}")'
    assert 'print("{ref:b-01}")' in _texts(hokoku.assemble(w2))
