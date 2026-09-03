"""
Документ без шаблона: пустой DOCX со своими стилями, полями и нумерацией, и запасные
стили для чужого шаблона, в котором их нет. Проверяем не структуру в памяти, а
собранный файл: цена ошибки здесь — отчёт, который открывается не так, как выглядел.
"""
import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from hokoku import (Image, Markdown, Text, blank_document, check_template,
                    document_bytes, ensure_style, ensure_styles, render)
from hokoku.template import BodyText, PageSetup, STYLE_SPECS
from .conftest import ptext


def _strip_styles(doc, keep=("Normal", "Default Paragraph Font")):
    """Документ без встроенных стилей — так выглядит шаблон, собранный не в Word
    (экспорт из Google Docs, генератор чужой службы)."""
    root = doc.styles.element
    for st in list(root.findall(qn("w:style"))):
        name = st.find(qn("w:name")).get(qn("w:val"))
        if name not in keep:
            root.remove(st)
    doc._hokoku_style_ids = None            # кэш имён стилей живёт на документе
    return doc


def _as_bytes(doc) -> bytes:
    return document_bytes(doc)


# ── документ с нуля ───────────────────────────────────────────────────────────

def test_blank_document_is_empty_and_gost():
    d = blank_document()
    assert d.paragraphs == []               # ни одного лишнего абзаца в готовом отчёте
    s = d.sections[0]
    assert round(s.page_width.cm) == 21 and round(s.page_height.cm, 1) == 29.7
    assert round(s.left_margin.cm, 1) == 3.0 and round(s.right_margin.cm, 1) == 1.5
    assert d.styles["Normal"].font.name == "Times New Roman"
    assert d.styles["Normal"].font.size.pt == 14
    assert d.styles["Normal"].paragraph_format.line_spacing == 1.5


def test_blank_document_has_all_fallback_styles():
    d = blank_document()
    names = {s.name for s in d.styles}
    assert set(STYLE_SPECS) <= names
    assert d.styles["Heading 1"].font.size.pt == 16          # styles.yaml, раздел headings
    ppr = d.styles["Heading 1"].element.find(qn("w:pPr"))
    assert ppr.find(qn("w:outlineLvl")).get(qn("w:val")) == "0"   # без этого нет оглавления
    # заводской «Heading 1» python-docx синий и Cambria — он переписан, а не оставлен
    assert d.styles["Heading 1"].font.color.rgb is None
    assert d.styles["Code"].element.get(qn("w:customStyle")) == "1"


def test_blank_document_page_numbers():
    d = blank_document()
    footer = d.sections[0].footer
    instr = [f.get(qn("w:instr")) for f in footer.paragraphs[0]._p.iter(qn("w:fldSimple"))]
    assert instr == [" PAGE "]               # номер полем, а не текстом: текст не пересчитается
    assert not blank_document(page_numbers=False).sections[0].footer.paragraphs[0]._p \
        .findall(qn("w:fldSimple"))


def test_blank_document_custom_page_and_body():
    d = blank_document(page=PageSetup(width_cm=29.7, height_cm=21.0, left_cm=2, right_cm=2,
                                      top_cm=1.5, bottom_cm=1.5),
                       body=BodyText(font="Arial", size_pt=12, line_spacing=1.0,
                                     first_line_cm=0, align="left"))
    assert round(d.sections[0].page_width.cm, 1) == 29.7
    assert d.styles["Normal"].font.name == "Arial" and d.styles["Normal"].font.size.pt == 12


def test_blank_document_renders_and_reopens(tmp_path, png):
    """Главное требование задачи: сочинённый документ идёт обычным путём render
    и открывается — стилями, а не прямым форматированием."""
    d = blank_document()
    d.add_paragraph("Тема: {{тема}}")
    d.add_paragraph("{{тело}}")
    d.add_paragraph("{{рис}}")
    out = str(tmp_path / "out.docx")
    res = render(_as_bytes(d), {
        "тема": Text("Сортировка"),
        "тело": Markdown("# Раздел\n\nтекст\n\n> цитата\n\n```python\nx = 1\n```"),
        "рис": Image(png, caption="Схема"),
    }, out)
    assert res.unfilled == [] and res.errors == [] and res.figures == 1
    o = Document(out)                        # файл действительно открывается
    styles = [(ptext(p), p.style.name) for p in o.paragraphs]
    assert ("Раздел", "Heading 1") in styles
    assert ("цитата", "Quote") in styles
    assert ("x = 1", "Code") in styles
    assert [n for t, n in styles if t.startswith("Рисунок 1")] == ["Caption"]
    assert len(o.inline_shapes) == 1


def test_blank_document_goes_through_build_report(tmp_path):
    """`build_report` — вход для недоверенного: сочинённый шаблон обязан пройти и его
    (validate_docx в том числе), иначе kadai упрётся в проверку собственного документа."""
    from hokoku import build_report
    d = blank_document()
    d.add_paragraph("{{текст}}")
    data = _as_bytes(d)
    res = build_report({"template": {"artifact": "tpl"},
                        "values": {"текст": {"type": "text", "text": "готово"}}},
                       resolve_artifact=lambda a: data, workdir=str(tmp_path))
    assert res["ok"], res
    docx = str(tmp_path / res["outputs"]["docx"]["file"])
    assert ptext(Document(docx).paragraphs[0]) == "готово"


# ── запасные стили для чужого шаблона ─────────────────────────────────────────

def test_missing_styles_are_created_not_faked(tmp_path):
    d = _strip_styles(Document())
    d.add_paragraph("{{тело}}")
    out = str(tmp_path / "out.docx")
    render(_as_bytes(d), {"тело": Markdown("# Раздел\n\n> цитата\n\n```\nкод\n```")}, out)
    o = Document(out)
    got = {ptext(p): p.style.name for p in o.paragraphs}
    assert got["Раздел"] == "Heading 1" and got["цитата"] == "Quote" and got["код"] == "Code"
    assert o.styles["Heading 1"].element.find(qn("w:pPr")).find(qn("w:outlineLvl")) is not None


def test_ensure_style_is_idempotent_and_cached():
    d = _strip_styles(Document())
    first = ensure_style(d, "Caption")
    assert first and ensure_style(d, "Caption") == first
    # второй стиль с тем же именем — верный признак, что кэш style_ids не обновили
    assert [s.name for s in d.styles].count("Caption") == 1


def test_ensure_style_refuses_unknown_name():
    """Сочинять произвольные стили нельзя: вид «List Paragraph» задаёт шаблон,
    и подделка была бы хуже прямого форматирования."""
    assert ensure_style(Document(), "List Paragraph") is None
    assert ensure_style(Document(), "Каких-то нет") is None


def test_ensure_styles_does_not_touch_existing():
    d = Document()
    d.styles["Heading 1"].font.size = Pt(11)
    assert "Heading 1" not in ensure_styles(d)
    assert d.styles["Heading 1"].font.size.pt == 11       # чужое оформление не переписываем


def test_own_style_of_foreign_template_survives_render(tmp_path):
    """Шаблон, в котором свой «Heading 1» и нет «Code»: недостающий стиль мы создаём,
    существующий — не трогаем. Обратное означало бы, что hokoku переписывает
    кафедральное оформление под своё, и заметили бы это только на защите."""
    d = Document()
    d.styles["Heading 1"].font.size = Pt(11)
    d.styles["Heading 1"].font.name = "Courier New"
    for st in list(d.styles.element.findall(qn("w:style"))):
        if st.find(qn("w:name")).get(qn("w:val")) == "Code":
            d.styles.element.remove(st)
    d._hokoku_style_ids = None
    d.add_paragraph("{{т}}")
    out = str(tmp_path / "out.docx")
    render(_as_bytes(d), {"т": Markdown("# Раздел\n\n```\nкод\n```")}, out)
    o = Document(out)
    assert o.styles["Heading 1"].font.size.pt == 11
    assert o.styles["Heading 1"].font.name == "Courier New"
    got = {ptext(p): p.style.name for p in o.paragraphs}
    assert got["Раздел"] == "Heading 1" and got["код"] == "Code"


def test_created_styles_keep_the_spacing_of_the_old_fallback():
    """Запасной стиль ставит те же интервалы, что прямое форматирование до него
    (`set_spacing(before=240, after=120)` у заголовка, 60/200 у подписи): переход
    на стили — про способ, а не про новый вид документа."""
    d = _strip_styles(Document())
    ensure_styles(d)
    for name, before, after in (("Heading 1", "240", "120"), ("Caption", "60", "200")):
        sp = d.styles[name].element.find(qn("w:pPr")).find(qn("w:spacing"))
        assert (sp.get(qn("w:before")), sp.get(qn("w:after"))) == (before, after)


def test_heading_size_override_reaches_created_style(tmp_path):
    """render(style={"headings": …}) должен действовать и там, где стиль сочиняем мы:
    иначе перегрузка молча ничего не меняет."""
    d = _strip_styles(Document())
    d.add_paragraph("{{т}}")
    out = str(tmp_path / "out.docx")
    render(_as_bytes(d), {"т": Markdown("# Раздел")}, out, style={"headings": {1: 22}})
    assert Document(out).styles["Heading 1"].font.size.pt == 22


# ── check_template ────────────────────────────────────────────────────────────

def test_check_template_clean_on_blank():
    assert check_template(blank_document()) == []


def test_check_template_reports_missing_styles():
    problems = check_template(_strip_styles(Document()))
    warned = {p.key for p in problems if p.level == "warning"}
    assert warned == {"Heading 1", "Heading 2", "Heading 3", "Caption"}
    assert all(p.module == "template" and p.code == "style_missing" for p in problems)
    assert "подставит свой" in [p for p in problems if p.key == "Caption"][0].message


def test_check_template_code_style_is_only_info():
    """У заводского документа Word нет стиля Code — это не повод пугать человека."""
    problems = check_template(Document())
    assert [p.level for p in problems] == ["info"]
    assert problems[0].key == "Code"


def test_check_template_margins_wider_than_page():
    d = blank_document()
    d.sections[0].left_margin = Cm(20)
    problems = [p for p in check_template(d) if p.level == "error"]
    assert problems and problems[0].code == "no_text_area"
    assert "за край" in problems[0].message


def test_check_template_narrow_text_area():
    d = blank_document()
    d.sections[0].left_margin = Cm(12)
    codes = [p.code for p in check_template(d)]
    assert "narrow_text_area" in codes


def test_check_template_validates_untrusted_docx():
    """Шаблон — недоверенный файл: проверка не должна быть отдельной дверью мимо safety."""
    from hokoku import DocxValidationError
    with pytest.raises(DocxValidationError):
        check_template("не docx".encode())
