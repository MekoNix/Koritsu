"""
Манифест шаблона: заготовка по тегам, JSON, сверка с шаблоном, промпт.

Ни один тест здесь не зовёт модель: манифест — это то, что уходит модели вместо DOCX,
и проверяется он сам по себе. Промпт проверяется на устойчивость (тот же манифест —
тот же текст), потому что на этом держится кэш префикса.
"""
import io

import pytest
from docx import Document

from hokoku import VALUE_TYPES
from hokoku.manifest import (DEFAULT_TYPE, MANIFEST_TYPES, ManifestError, TagSpec,
                             check_manifest, manifest_from_json, manifest_from_template,
                             manifest_prompt, manifest_schema, manifest_to_json,
                             suggest_type)


def docx_bytes(build) -> bytes:
    d = Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def template() -> bytes:
    def build(d):
        d.add_paragraph("Цель: {{цель:Цель работы}}.")
        d.add_paragraph("{{схема:Блок-схема алгоритма}}")
        d.add_paragraph("{{замеры:Таблица замеров}}")
        d.sections[0].header.paragraphs[0].text = "{{кафедра:Кафедра}}"
    return docx_bytes(build)


def codes(warnings, key=None) -> list[str]:
    return [w.code for w in warnings if key is None or w.key == key]


# ── заготовка ─────────────────────────────────────────────────────────────────

def test_заготовка_ставит_markdown_и_догадку_по_метке(template):
    """По умолчанию всё markdown; тип по метке — предложение, и оно помечено как догадка."""
    m = manifest_from_template(template)
    assert list(m.tags) == ["цель", "схема", "замеры", "кафедра"]     # порядок документа
    assert (m.tags["цель"].type, m.tags["цель"].guessed) == (DEFAULT_TYPE, False)
    assert (m.tags["схема"].type, m.tags["схема"].guessed) == ("diagram", True)
    assert (m.tags["замеры"].type, m.tags["замеры"].guessed) == ("table", True)
    assert m.tags["цель"].label == "Цель работы"


def test_заготовка_знает_про_колонтитул(template):
    """В колонтитуле нет ни нумерации, ни подписи, и заполняет его человек."""
    tag = manifest_from_template(template).tags["кафедра"]
    assert (tag.required, tag.numbered) == (False, False)


def test_шаблон_опознаётся_по_sha256(template):
    import hashlib
    m = manifest_from_template(template)
    assert m.template_sha256 == hashlib.sha256(template).hexdigest()
    assert (m.manifest_version, m.wire_version) == (1, 1)


@pytest.mark.parametrize("label,expected", [
    ("Блок-схема", "diagram"), ("Рисунок установки", "image"), ("Листинг функции", "code"),
    ("Формула Герона", "formula"), ("Таблица замеров", "table"), ("Оглавление", "toc"),
    ("Цель работы", None), ("", None)])
def test_догадка_по_метке(label, expected):
    assert suggest_type(label) == expected


def test_типы_манифеста_те_же_что_у_значений():
    """Разъехавшиеся списки типов — тег, тип которого нельзя выразить значением."""
    assert set(MANIFEST_TYPES) == set(VALUE_TYPES) - {"page_break"}


# ── JSON ──────────────────────────────────────────────────────────────────────

def test_круговой_прогон_json(template):
    m = manifest_from_template(template, system_prompt="Отчёт по ГОСТ 7.32.")
    m.tags["цель"].prompt = "Сформулируй цель одним абзацем."
    m.tags["цель"].limits = {"max_chars": 600, "headings": False}
    m.tags["цель"].depends_on = ["задание"]
    m.tags["схема"].tool = "make_flowchart"
    d = manifest_to_json(m)
    assert manifest_to_json(manifest_from_json(d)) == d


def test_умолчания_в_json_не_выписываются(template):
    """Манифест объявляет исключения: диф правки должен показывать правку."""
    d = manifest_to_json(manifest_from_template(template))
    assert d["tags"]["цель"] == {"type": "markdown", "label": "Цель работы"}


def test_неизвестное_поле_названо():
    with pytest.raises(ManifestError, match=r'promt.*похоже на "prompt"'):
        manifest_from_json({"tags": {"цель": {"promt": "текст"}}})


def test_неизвестный_ключ_манифеста():
    with pytest.raises(ManifestError, match="lanugage"):
        manifest_from_json({"lanugage": "ru"})


def test_неизвестный_тип_тега():
    with pytest.raises(ManifestError, match="picture"):
        manifest_from_json({"tags": {"схема": {"type": "picture"}}})


def test_неизвестное_ограничение():
    with pytest.raises(ManifestError, match=r'max_char.*похоже на "max_chars"'):
        manifest_from_json({"tags": {"цель": {"limits": {"max_char": 600}}}})


def test_манифест_из_будущего_отказ():
    with pytest.raises(ManifestError, match="знает до"):
        manifest_from_json({"wire_version": 99})


@pytest.mark.parametrize("bad", [
    {"tags": {"цель": {"required": "да"}}},
    {"tags": {"цель": {"source_hint": "робот"}}},
    {"tags": {"цель": {"depends_on": "задание"}}},
    {"tags": {"цель": {"limits": {"max_chars": 0}}}},
    {"manifest_version": 0},
    {"tags": ["цель"]}])
def test_отказы_разбора(bad):
    with pytest.raises(ManifestError):
        manifest_from_json(bad)


def test_ключи_в_nfc():
    """Word и macOS пишут «й» разложенным; ключи в NFC с обеих сторон."""
    nfd = "случаи\u0306"
    assert list(manifest_from_json({"tags": {nfd: {}}}).tags) == ["случай"]
    with pytest.raises(ManifestError, match="повторяется"):
        manifest_from_json({"tags": {nfd: {}, "случай": {}}})


# ── сверка с шаблоном ─────────────────────────────────────────────────────────

def test_тег_без_записи_и_запись_без_тега(template):
    m = manifest_from_template(template)
    del m.tags["схема"]
    m.tags["выводы"] = TagSpec(label="Выводы", prompt="Три пункта.")
    w = check_manifest(m, template)
    assert "tag_without_entry" in codes(w, "схема")
    assert "entry_without_tag" in codes(w, "выводы")


def test_похоже_на_переименование(template):
    """Молча переносить промпт нельзя: два тега, поменянных местами в Word, дали бы
    перепутанные значения. Только подсказка человеку."""
    m = manifest_from_template(template)
    m.tags["замеры_"] = m.tags.pop("замеры")
    w = [x for x in check_manifest(m, template) if x.key == "замеры_"]
    assert w and "переименовали" in w[0].message


def test_метка_изменилась_и_тип_догадка(template):
    m = manifest_from_template(template)
    m.tags["цель"].label = "Задача работы"
    w = check_manifest(m, template)
    assert "label_changed" in codes(w, "цель")
    assert "type_guessed" in codes(w, "схема")


def test_нумерация_в_колонтитуле(template):
    m = manifest_from_template(template)
    m.tags["кафедра"].numbered = True
    assert "numbered_in_header" in codes(check_manifest(m, template), "кафедра")


def test_подтверждённый_манифест_молчит(template):
    m = manifest_from_template(template)
    for tag in m.tags.values():
        tag.guessed, tag.prompt = False, "задание"
    assert check_manifest(m, template) == []


# ── обновление под новый DOCX ─────────────────────────────────────────────────

def test_обновление_бережёт_промпт_и_метит_исчезнувшие(template):
    """Тег убрали — запись не удаляется: промпт и значения переживут правку шаблона."""
    m = manifest_from_template(template)
    m.tags["схема"].prompt = "Блок-схема сортировки."
    m.tags["схема"].guessed = False
    tpl2 = docx_bytes(lambda d: (d.add_paragraph("{{цель:Цель работы}}"),
                                  d.add_paragraph("{{выводы:Выводы}}")))
    m2 = manifest_from_template(tpl2, base=m)
    assert m2.tags["схема"].missing is True
    assert m2.tags["схема"].prompt == "Блок-схема сортировки."
    assert m2.tags["выводы"].type == DEFAULT_TYPE and m2.tags["выводы"].prompt == ""
    assert m2.manifest_version == 2
    assert m2.tags["цель"].missing is False


def test_обновление_обновляет_метку_не_трогая_промпт(template):
    m = manifest_from_template(template)
    m.tags["цель"].prompt = "Одним абзацем."
    tpl2 = docx_bytes(lambda d: d.add_paragraph("{{цель:Задача работы}}"))
    tag = manifest_from_template(tpl2, base=m).tags["цель"]
    assert (tag.label, tag.prompt) == ("Задача работы", "Одним абзацем.")


def test_вернувшийся_тег_замечен(template):
    m = manifest_from_template(template)
    m.tags["цель"].missing = True
    assert "tag_returned" in codes(check_manifest(m, template), "цель")


# ── промпт ────────────────────────────────────────────────────────────────────

@pytest.fixture
def filled(template):
    m = manifest_from_template(template, system_prompt="Отчёт по ГОСТ 7.32. Деловой стиль.")
    m.tags["цель"].prompt = "Сформулируй цель одним абзацем."
    m.tags["цель"].example = "Изучить алгоритмы сортировки."
    m.tags["цель"].limits = {"max_chars": 600, "headings": False}
    m.tags["цель"].depends_on = ["задание"]
    m.tags["замеры"].limits = {"max_rows": 50, "max_cols": 8}
    return m


def test_промпт_говорит_про_каждый_тег(filled):
    text = manifest_prompt(filled)
    assert text.startswith("Отчёт по ГОСТ 7.32. Деловой стиль.")
    for key in filled.tags:
        assert f"[{key}]" in text
    assert "задание: Сформулируй цель одним абзацем." in text
    assert "пример: Изучить алгоритмы сортировки." in text
    assert "не длиннее 600 знаков, заголовков не ставить" in text
    assert "опирается на: задание" in text
    assert "строк не больше 50, колонок не больше 8" in text
    assert "XML не сочинять" in text                 # схему рисует инструмент, не модель


def test_в_промпте_нет_служебного(filled):
    text = manifest_prompt(filled)
    assert filled.template_sha256 not in text
    assert "wire_version" not in text and "source_hint" not in text and "sha256" not in text


def test_промпт_устойчив(filled):
    """Тот же манифест — тот же текст знак в знак: иначе кэш префикса промахивается."""
    assert manifest_prompt(filled) == manifest_prompt(filled)


def test_промпт_сужается_до_одного_тега(filled):
    text = manifest_prompt(filled, keys=["цель"])
    assert "[цель]" in text and "[схема]" not in text


def test_исчезнувший_тег_в_промпт_не_идёт(filled):
    filled.tags["схема"].missing = True
    assert "[схема]" not in manifest_prompt(filled)


def test_известное_исчезновение_не_кричит(template):
    """Уже помеченное «нет в шаблоне» — состояние, а не новость на каждой сверке."""
    m = manifest_from_template(template)
    tpl2 = docx_bytes(lambda d: d.add_paragraph("{{цель:Цель работы}}"))
    m2 = manifest_from_template(tpl2, base=m)
    levels = {w.level for w in check_manifest(m2, tpl2) if w.code == "entry_without_tag"}
    assert levels == {"info"}


# ── схема всего отчёта ────────────────────────────────────────────────────────

def test_схема_отчёта_склеена_из_схем_типов(filled):
    """Уровень «весь отчёт одним вызовом»: объект «ключ тега → схема его типа»."""
    from hokoku import value_schema
    s = manifest_schema(filled)
    assert s["type"] == "object" and s["additionalProperties"] is False
    assert list(s["properties"]) == list(filled.tags)
    assert s["properties"]["замеры"] == value_schema("table", for_model=True)
    assert "$schema" not in s                      # в корне запроса поставщик ждёт объект
    assert set(s["required"]) == {"цель", "схема", "замеры"}     # кафедра необязательна


def test_схему_отчёта_принимает_валидатор_llm(filled):
    """Схема, которую наш же валидатор не берёт, — это 400 от поставщика на живом вызове."""
    from llm.jsonschema import check_schema
    check_schema(manifest_schema(filled))


def test_ответ_по_схеме_отчёта_разбирается(filled):
    """Свойство, ради которого схема и собирается: что модель по ней вернула, то
    `values_from_json` разбирает без единой ошибки."""
    from llm.jsonschema import validate as schema_validate
    from hokoku import values_from_json
    answer = {
        "цель": {"type": "markdown", "text": "Изучить алгоритмы сортировки."},
        "схема": {"type": "diagram", "artifact": "sort-flow", "caption": "Схема",
                  "width_cm": None, "align": "center", "ref": None, "page": None},
        "замеры": {"type": "table", "rows": [["n", "t"], ["10", "1"]], "header": True,
                   "caption": None, "col_widths_cm": None, "align": None, "ref": None},
        "кафедра": {"type": "markdown", "text": "ИУ7"}}
    assert schema_validate(answer, manifest_schema(filled)) == []
    values, errors = values_from_json(answer, resolve_artifact=lambda a: b"<mxfile/>")
    assert errors == []
    assert set(values) == set(answer)


def test_исчезнувший_тег_в_схему_не_идёт_и_keys_сужает(filled):
    filled.tags["схема"].missing = True
    assert "схема" not in manifest_schema(filled)["properties"]
    assert list(manifest_schema(filled, keys=["цель"])["properties"]) == ["цель"]
